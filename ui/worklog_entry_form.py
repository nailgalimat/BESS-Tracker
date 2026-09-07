"""
ui/worklog_entry_form.py
--------------------------
"Field Log" — split into two pages (like Work Log / Work Report):

  FieldLogEntryPage    — a roomy form to create a new field-log entry with photos
  FieldLogRecordsPage  — filter bar + timeline of existing logs, export, and the
                         monthly technical-works report

Shared widgets (photo thumbnails, log cards, the edit dialog with the
structured spare-parts / stock panel) live in this module and are used by both
pages plus the main window.
"""

import os
import shutil

from PyQt5.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QFormLayout, QSplitter,
    QLabel, QComboBox, QDateEdit, QLineEdit, QTextEdit,
    QPushButton, QMessageBox, QGroupBox, QScrollArea, QFrame,
    QFileDialog, QDialog, QDialogButtonBox, QSizePolicy,
    QGridLayout, QAbstractScrollArea, QMenu, QAction,
    QTableWidget, QTableWidgetItem, QHeaderView, QAbstractItemView,
    QDoubleSpinBox, QCheckBox, QSpinBox,
)
from PyQt5.QtCore import Qt, QDate, pyqtSignal, QSize
from PyQt5.QtGui import QFont, QColor, QPixmap, QIcon

from services.project_service import (
    get_all_projects, get_zones_for_project,
    get_blocks_for_zone, get_containers_for_block,
)
from services.worklog_entry_service import (
    save_worklog_entry, get_worklog_entries, get_worklog_entry,
    update_worklog_entry, delete_worklog_entry, get_tags,
    CATEGORIES, CATEGORY_LABELS, CATEGORY_COLORS,
)
from services.image_service import (
    save_image, get_images_for_log, delete_image, delete_images_for_log,
)


# ── Small helpers ─────────────────────────────────────────────────────────────

def _category_combo(current: str = "maintenance") -> QComboBox:
    cb = QComboBox()
    for key in CATEGORIES:
        cb.addItem(CATEGORY_LABELS[key], userData=key)
    idx = CATEGORIES.index(current) if current in CATEGORIES else 0
    cb.setCurrentIndex(idx)
    return cb


# Common BESS faults offered as suggestions in the (editable) fault field.
# Editable, so an engineer can type any custom text too — all optional.
COMMON_FAULTS = [
    "", "BSC-PCS Communication", "LC-PCS Communication", "LC-BSC Communication",
    "PCS Fault", "Converter Unit Fault", "Islanding Protection",
    "DC-DC Converter Fault", "DC Over-voltage", "DC Under-voltage",
    "Battery Over-temperature", "Battery Under-temperature",
    "Cell Over-voltage", "Cell Under-voltage", "Cell Imbalance",
    "Insulation Impedance Low", "Dry Node Input Fault", "CMU Fault",
    "Fan Fault", "SCU-DSP Communication", "System Not Ready",
    "Grid Outage", "Auxiliary Power Loss", "Fire / Smoke Alarm",
]


def _fault_combo(current: str = "") -> QComboBox:
    """Editable combo: pick a common fault OR type a custom one. Optional."""
    cb = QComboBox()
    cb.setEditable(True)
    cb.addItems(COMMON_FAULTS)
    cb.setCurrentText(current or "")
    cb.lineEdit().setPlaceholderText("Fault / problem (optional)")
    return cb


def _status_combo(current: str = "") -> QComboBox:
    """Optional status: unspecified / open / done."""
    cb = QComboBox()
    cb.addItem("— (unspecified)", userData="")
    cb.addItem("Open",  userData="open")
    cb.addItem("Done",  userData="done")
    idx = {"": 0, "open": 1, "done": 2}.get((current or "").lower(), 0)
    cb.setCurrentIndex(idx)
    return cb


def _entry_photo_prefix(entry: dict) -> str:
    """Build a filename prefix from an entry: '<date>_<location>' for photo export."""
    date = entry.get('log_date', '') or ''
    if entry.get('block_number') is not None:
        z = entry.get('zone_number', '')
        b = entry.get('block_number', '')
        ci = entry.get('container_index', '')
        loc = f"Z{z}-B{b}-C{ci}"
    else:
        loc = entry.get('site_location', '') or ''
    return f"{date}_{loc}".strip('_') or "photo"


def _load_pixmap_safe(path: str, w: int, h: int) -> QPixmap:
    """Load and scale a pixmap; return a placeholder on failure."""
    if path and os.path.isfile(path):
        px = QPixmap(path)
        if not px.isNull():
            return px.scaled(w, h, Qt.KeepAspectRatio, Qt.SmoothTransformation)
    # Grey placeholder
    px = QPixmap(w, h)
    px.fill(QColor("#E0E0E0"))
    return px


# ── Photo thumbnail (clickable) ───────────────────────────────────────────────

class PhotoThumb(QLabel):
    """80×60 px clickable thumbnail.  Emits clicked(file_path)."""
    clicked = pyqtSignal(str)

    def __init__(
        self,
        file_path:     str,
        image_id:      str = "",
        upload_status: str = "local",
        parent=None,
    ):
        super().__init__(parent)
        self._path          = file_path
        self._image_id      = image_id
        self._upload_status = upload_status
        self.setFixedSize(82, 62)
        self.setAlignment(Qt.AlignCenter)
        self.setCursor(Qt.PointingHandCursor)

        has_local = bool(file_path and os.path.isfile(file_path))

        if has_local:
            self.setStyleSheet(
                "border: 1px solid #BDBDBD; border-radius: 4px; background: #F5F5F5;"
            )
            self.setPixmap(_load_pixmap_safe(file_path, 80, 60))
            self.setToolTip(os.path.basename(file_path))
        elif upload_status == "remote":
            self.setStyleSheet(
                "border: 1px solid #90CAF9; border-radius: 4px;"
                " background: #E3F2FD; color: #1565C0; font-size: 20px;"
            )
            self.setText("☁")
            self.setToolTip("Photo stored on server — click to download")
        else:
            self.setStyleSheet(
                "border: 1px dashed #BDBDBD; border-radius: 4px;"
                " background: #F5F5F5; color: #9E9E9E; font-size: 18px;"
            )
            self.setText("?")
            self.setToolTip("Photo file not found locally")

    def mousePressEvent(self, e):
        if e.button() == Qt.LeftButton:
            self.clicked.emit(self._path)


# ── Full-size image viewer dialog ─────────────────────────────────────────────

class ImageViewerDialog(QDialog):
    def __init__(self, file_path: str, parent=None):
        super().__init__(parent)
        self._path = file_path
        self.setWindowTitle(os.path.basename(file_path))
        self.setMinimumSize(640, 480)
        lay = QVBoxLayout(self)
        lbl = QLabel()
        lbl.setAlignment(Qt.AlignCenter)
        px = _load_pixmap_safe(file_path, 900, 700)
        lbl.setPixmap(px)
        lay.addWidget(lbl)

        btn_row = QHBoxLayout()
        save_btn = QPushButton("💾 Save As…")
        save_btn.clicked.connect(self._save_as)
        btn_row.addWidget(save_btn)
        btn_row.addStretch()
        close_btn = QPushButton("Close")
        close_btn.clicked.connect(self.accept)
        btn_row.addWidget(close_btn)
        lay.addLayout(btn_row)

    def _save_as(self):
        if not (self._path and os.path.isfile(self._path)):
            QMessageBox.warning(self, "Save", "Photo file is not available locally.")
            return
        ext = os.path.splitext(self._path)[1] or ".jpg"
        suggested = os.path.basename(self._path) or f"photo{ext}"
        dest, _ = QFileDialog.getSaveFileName(
            self, "Save Photo As", suggested,
            "Images (*.jpg *.jpeg *.png *.webp *.bmp *.tiff);;All files (*)")
        if not dest:
            return
        try:
            shutil.copy2(self._path, dest)
        except Exception as ex:
            QMessageBox.critical(self, "Save Error", str(ex))
            return
        if QMessageBox.question(
            self, "Saved", f"Saved to:\n{dest}\n\nOpen containing folder?",
            QMessageBox.Yes | QMessageBox.No
        ) == QMessageBox.Yes:
            try:
                os.startfile(os.path.dirname(dest))
            except Exception:
                pass


# ── Photo picker widget (used inside the entry form) ─────────────────────────

class PhotoPickerWidget(QWidget):
    """Row of thumbnails for staged (not-yet-saved) photos."""
    photos_changed = pyqtSignal()

    def __init__(self, parent=None):
        super().__init__(parent)
        self._paths: list = []
        self._build()

    def _build(self):
        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)
        outer.setSpacing(4)

        add_btn = QPushButton("📷  Add Photos")
        add_btn.setStyleSheet(
            "QPushButton{border:1px dashed #90A4AE;border-radius:6px;"
            "padding:6px 12px;background:#FAFAFA;color:#37474F;}"
            "QPushButton:hover{background:#E3F2FD;border-color:#1976D2;}"
        )
        add_btn.clicked.connect(self._add_photos)
        outer.addWidget(add_btn)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFixedHeight(100)
        scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        scroll.setVerticalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        scroll.setSizeAdjustPolicy(QAbstractScrollArea.AdjustToContents)
        scroll.setStyleSheet("QScrollArea{border:none;background:transparent;}")

        self._thumb_container = QWidget()
        self._thumb_layout = QHBoxLayout(self._thumb_container)
        self._thumb_layout.setContentsMargins(0, 0, 0, 0)
        self._thumb_layout.setSpacing(6)
        self._thumb_layout.addStretch()
        scroll.setWidget(self._thumb_container)
        outer.addWidget(scroll)

    def _add_photos(self):
        paths, _ = QFileDialog.getOpenFileNames(
            self, "Select Photos", "",
            "Images (*.jpg *.jpeg *.png *.webp *.heic *.bmp *.tiff)"
        )
        for p in paths:
            if p not in self._paths:
                self._paths.append(p)
                self._add_thumb(p)
        if paths:
            self.photos_changed.emit()

    def _add_thumb(self, path: str):
        cell = QWidget()
        cl = QVBoxLayout(cell)
        cl.setContentsMargins(0, 0, 0, 0)
        cl.setSpacing(2)

        thumb = PhotoThumb(path)
        thumb.clicked.connect(
            lambda p: ImageViewerDialog(p, self).exec_()
        )
        cl.addWidget(thumb)

        remove_btn = QPushButton("✕")
        remove_btn.setFixedSize(82, 18)
        remove_btn.setStyleSheet(
            "QPushButton{border:none;background:#EF9A9A;color:#B71C1C;"
            "border-radius:3px;font-size:10px;}"
            "QPushButton:hover{background:#EF5350;color:white;}"
        )
        remove_btn.clicked.connect(lambda _, p=path: self._remove_photo(p))
        cl.addWidget(remove_btn)

        count = self._thumb_layout.count()
        self._thumb_layout.insertWidget(count - 1, cell)

    def _remove_photo(self, path: str):
        if path in self._paths:
            self._paths.remove(path)
        self._rebuild_thumbs()
        self.photos_changed.emit()

    def _rebuild_thumbs(self):
        while self._thumb_layout.count():
            item = self._thumb_layout.takeAt(0)
            if item.widget():
                item.widget().deleteLater()
        self._thumb_layout.addStretch()
        for p in self._paths:
            self._add_thumb(p)

    def get_paths(self) -> list:
        return list(self._paths)

    def clear(self):
        self._paths.clear()
        self._rebuild_thumbs()


# ── Structured spare-parts / stock panel ──────────────────────────────────────

class SparePartsPanel(QGroupBox):
    """
    Manage structured spare parts for a saved entry and deduct them from stock.
    Stock is only touched when the user explicitly presses "Deduct from stock"
    (separate confirmation) — never automatically.
    """

    def __init__(self, entry_id: str, parent=None):
        super().__init__("🔧  Spare Parts from Stock", parent)
        self._entry_id = entry_id
        self._build()
        self.refresh()

    def _build(self):
        lay = QVBoxLayout(self)
        lay.setSpacing(6)

        self._wh_label = QLabel("Warehouse: —")
        self._wh_label.setStyleSheet("color:#1565C0;font-size:11px;")
        lay.addWidget(self._wh_label)

        self._table = QTableWidget(0, 5)
        self._table.setHorizontalHeaderLabels(
            ["Material", "Description", "Qty", "Unit", "Status"])
        self._table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self._table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self._table.setSelectionMode(QAbstractItemView.SingleSelection)
        self._table.verticalHeader().setVisible(False)
        self._table.setMinimumHeight(120)
        hh = self._table.horizontalHeader()
        hh.setSectionResizeMode(0, QHeaderView.ResizeToContents)
        hh.setSectionResizeMode(1, QHeaderView.Stretch)
        hh.setSectionResizeMode(2, QHeaderView.ResizeToContents)
        hh.setSectionResizeMode(3, QHeaderView.ResizeToContents)
        hh.setSectionResizeMode(4, QHeaderView.ResizeToContents)
        lay.addWidget(self._table)

        # Add-row controls
        add_row = QHBoxLayout()
        self._mat_combo = QComboBox()
        self._mat_combo.setEditable(True)
        self._mat_combo.setMinimumWidth(220)
        self._mat_combo.lineEdit().setPlaceholderText("Material (number or pick from list)")
        self._load_materials()
        add_row.addWidget(self._mat_combo, 3)

        self._qty_spin = QDoubleSpinBox()
        self._qty_spin.setRange(0.0, 1_000_000.0)
        self._qty_spin.setDecimals(2)
        self._qty_spin.setValue(1.0)
        self._qty_spin.setFixedWidth(90)
        add_row.addWidget(self._qty_spin)

        add_btn = QPushButton("➕ Add")
        add_btn.clicked.connect(self._add_part)
        add_row.addWidget(add_btn)
        lay.addLayout(add_row)

        # Action buttons
        act_row = QHBoxLayout()
        self._deduct_btn = QPushButton("📤 Deduct from stock")
        self._deduct_btn.setStyleSheet(
            "QPushButton{background:#EF6C00;color:white;font-weight:bold;"
            "border-radius:4px;padding:5px 12px;}"
            "QPushButton:hover{background:#E65100;}"
            "QPushButton:disabled{background:#E0E0E0;color:#9E9E9E;}"
        )
        self._deduct_btn.clicked.connect(self._deduct_all)
        act_row.addWidget(self._deduct_btn)

        self._row_btn = QPushButton("Row action ▾")
        self._row_btn.clicked.connect(self._row_menu)
        act_row.addWidget(self._row_btn)
        act_row.addStretch()
        lay.addLayout(act_row)

        hint = QLabel("Deduction records an OUT stock transaction (ref = LOG-…). "
                      "Nothing is deducted automatically.")
        hint.setWordWrap(True)
        hint.setStyleSheet("color:#9E9E9E;font-size:10px;")
        lay.addWidget(hint)

    def _load_materials(self):
        self._mat_combo.clear()
        self._mat_combo.addItem("", userData="")
        try:
            from services.material_service import get_all_materials
            for m in get_all_materials():
                num = m["material_number"]
                desc = m.get("description") or ""
                label = f"{num} — {desc}" if desc else num
                self._mat_combo.addItem(label, userData=num)
        except Exception:
            pass
        self._mat_combo.setCurrentIndex(0)

    def _selected_material_number(self) -> str:
        # userData if a catalog item is chosen; otherwise the typed text (may be
        # "NUMBER — desc" or a bare number)
        data = self._mat_combo.currentData()
        if data:
            return data
        txt = self._mat_combo.currentText().strip()
        if " — " in txt:
            txt = txt.split(" — ", 1)[0].strip()
        return txt

    def refresh(self):
        from services.worklog_parts_service import get_parts_for_entry
        parts = get_parts_for_entry(self._entry_id)
        self._parts = parts
        self._table.setRowCount(0)

        wh_name = parts[0]["target_warehouse_name"] if parts else None
        if wh_name is None:
            try:
                from services.worklog_parts_service import resolve_warehouse_for_entry
                wh_name = resolve_warehouse_for_entry(self._entry_id)["name"]
            except Exception:
                wh_name = "—"
        self._wh_label.setText(f"Deduct from warehouse: {wh_name}")

        pending = 0
        for p in parts:
            r = self._table.rowCount()
            self._table.insertRow(r)
            item_mat = QTableWidgetItem(p["material_number"])
            item_mat.setData(Qt.UserRole, p["id"])
            self._table.setItem(r, 0, item_mat)
            self._table.setItem(r, 1, QTableWidgetItem(p.get("description") or ""))
            self._table.setItem(r, 2, QTableWidgetItem(f"{p['quantity']:g}"))
            self._table.setItem(r, 3, QTableWidgetItem(p.get("unit") or ""))
            if p["deducted"]:
                st = QTableWidgetItem("✔ Deducted")
                st.setForeground(QColor("#2E7D32"))
            else:
                avail = p.get("available", 0)
                short = avail < p["quantity"]
                st = QTableWidgetItem(
                    f"Pending (in stock: {avail:g})" + (" ⚠" if short else ""))
                st.setForeground(QColor("#C62828") if short else QColor("#EF6C00"))
                pending += 1
            self._table.setItem(r, 4, st)

        self._deduct_btn.setText(f"📤 Deduct from stock ({pending})")
        self._deduct_btn.setEnabled(pending > 0)

    def _add_part(self):
        mat = self._selected_material_number()
        if not mat:
            QMessageBox.warning(self, "Material", "Select or type a material number.")
            return
        qty = self._qty_spin.value()
        if qty <= 0:
            QMessageBox.warning(self, "Quantity", "Quantity must be greater than 0.")
            return
        try:
            from services.worklog_parts_service import add_part
            add_part(self._entry_id, mat, qty)
        except Exception as ex:
            QMessageBox.critical(self, "Error", str(ex))
            return
        self._mat_combo.setCurrentIndex(0)
        self._qty_spin.setValue(1.0)
        self.refresh()

    def _current_part(self):
        r = self._table.currentRow()
        if r < 0:
            return None
        item = self._table.item(r, 0)
        if not item:
            return None
        pid = item.data(Qt.UserRole)
        return next((p for p in self._parts if p["id"] == pid), None)

    def _row_menu(self):
        p = self._current_part()
        if not p:
            QMessageBox.information(self, "Row", "Select a row in the table first.")
            return
        menu = QMenu(self)
        if p["deducted"]:
            menu.addAction("↩ Reverse deduction (return to stock)", self._reverse_current)
        else:
            menu.addAction("📤 Deduct this part", self._deduct_current)
            menu.addAction("🗑 Remove part", self._delete_current)
        menu.exec_(self._row_btn.mapToGlobal(self._row_btn.rect().bottomLeft()))

    def _deduct_current(self):
        p = self._current_part()
        if not p:
            return
        if QMessageBox.question(
            self, "Deduct from stock",
            f"Deduct {p['quantity']:g} x {p['material_number']} from stock?",
            QMessageBox.Yes | QMessageBox.No, QMessageBox.No
        ) != QMessageBox.Yes:
            return
        from services.worklog_parts_service import deduct_part
        res = deduct_part(p["id"])
        if not res["ok"]:
            QMessageBox.warning(self, "Deduction", res["message"])
        self.refresh()

    def _deduct_all(self):
        from services.worklog_parts_service import get_parts_for_entry, deduct_all_pending
        parts = get_parts_for_entry(self._entry_id)
        pending = [p for p in parts if not p["deducted"]]
        if not pending:
            return
        wh = pending[0]["target_warehouse_name"]
        lines = "\n".join(f"  • {p['quantity']:g} x {p['material_number']}"
                          f"  (in stock: {p.get('available',0):g})"
                          for p in pending)
        short = [p for p in pending if p.get("available", 0) < p["quantity"]]
        warn = ("\n\n⚠ Some parts exceed the stock on hand — "
                "they will go to 0." if short else "")
        if QMessageBox.question(
            self, "Deduct from stock",
            f"Deduct from warehouse '{wh}':\n\n{lines}{warn}\n\nContinue?",
            QMessageBox.Yes | QMessageBox.No, QMessageBox.No
        ) != QMessageBox.Yes:
            return
        res = deduct_all_pending(self._entry_id)
        if res["failed"]:
            QMessageBox.warning(
                self, "Partially deducted",
                f"Deducted: {res['deducted']}\nErrors:\n" + "\n".join(res["failed"]))
        else:
            QMessageBox.information(self, "Done",
                                    f"Parts deducted: {res['deducted']}.")
        self.refresh()

    def _reverse_current(self):
        p = self._current_part()
        if not p:
            return
        if QMessageBox.question(
            self, "Reverse deduction",
            f"Return {p['quantity']:g} x {p['material_number']} back to stock?",
            QMessageBox.Yes | QMessageBox.No, QMessageBox.No
        ) != QMessageBox.Yes:
            return
        from services.worklog_parts_service import reverse_part
        res = reverse_part(p["id"])
        if not res["ok"]:
            QMessageBox.warning(self, "Reverse", res["message"])
        self.refresh()

    def _delete_current(self):
        p = self._current_part()
        if not p:
            return
        from services.worklog_parts_service import delete_part
        try:
            delete_part(p["id"])
        except Exception as ex:
            QMessageBox.warning(self, "Remove", str(ex))
        self.refresh()


# ── Timeline card (one log entry) ────────────────────────────────────────────

class LogCard(QFrame):
    """Displays one work_log_entry in the timeline."""
    edit_requested   = pyqtSignal(str)
    delete_requested = pyqtSignal(str)

    def __init__(self, entry: dict, images: list, parent=None):
        super().__init__(parent)
        self._entry  = entry
        self._images = images
        self._build()

    def _build(self):
        entry = self._entry
        cat   = entry.get('category', 'other')
        bg, fg = CATEGORY_COLORS.get(cat, ("#F5F5F5", "#424242"))

        self.setFrameShape(QFrame.StyledPanel)
        self.setStyleSheet(f"""
            LogCard {{
                background: {bg};
                border: 1px solid {fg}40;
                border-radius: 8px;
                margin: 2px 4px;
            }}
        """)

        outer = QVBoxLayout(self)
        outer.setContentsMargins(10, 8, 10, 8)
        outer.setSpacing(4)

        # Header row
        header = QHBoxLayout()
        cat_lbl = QLabel(CATEGORY_LABELS.get(cat, cat.title()))
        cat_lbl.setStyleSheet(
            f"background:{fg}; color:white; border-radius:10px;"
            f" padding:1px 8px; font-size:11px; font-weight:bold;"
        )
        cat_lbl.setFixedHeight(22)
        header.addWidget(cat_lbl)

        loc = entry.get('site_location') or ''
        if entry.get('block_number'):
            z = entry.get('zone_number', '')
            b = entry.get('block_number', '')
            ci = entry.get('container_index', '')
            loc = f"Z{z} · B{b} · C{ci}"
            if entry.get('container_type'):
                loc += f"  ({entry['container_type']})"

        loc_lbl = QLabel(loc)
        loc_lbl.setStyleSheet(f"color:{fg}; font-size:12px;")
        header.addWidget(loc_lbl)
        header.addStretch()

        date_lbl = QLabel(entry.get('log_date', ''))
        date_lbl.setStyleSheet("color:#757575; font-size:11px;")
        header.addWidget(date_lbl)
        outer.addLayout(header)

        # Description
        desc = entry.get('description', '')
        lines = desc.split('\n')
        preview = '\n'.join(lines[:3])
        if len(lines) > 3:
            preview += ' …'
        desc_lbl = QLabel(preview)
        desc_lbl.setWordWrap(True)
        desc_lbl.setStyleSheet("color:#212121; font-size:12px; padding:2px 0;")
        outer.addWidget(desc_lbl)

        # Optional technical-work meta chips
        meta_bits = []
        if entry.get('fault_name'):
            meta_bits.append(("⚠ " + entry['fault_name'], "#C62828"))
        _st = (entry.get('status') or '').lower()
        if _st == 'open':
            meta_bits.append(("● Open", "#EF6C00"))
        elif _st == 'done':
            meta_bits.append(("✔ Done", "#2E7D32"))
        if entry.get('sap_ticket'):
            meta_bits.append(("SAP: " + entry['sap_ticket'], "#455A64"))
        if entry.get('spare_parts'):
            meta_bits.append(("🔧 " + entry['spare_parts'], "#455A64"))
        # Structured stock parts badge
        try:
            from services.worklog_parts_service import get_parts_count
            pc = get_parts_count(entry['id'])
            if pc["total"]:
                if pc["pending"]:
                    meta_bits.append((f"📦 stock: {pc['deducted']}/{pc['total']} "
                                      f"(pending {pc['pending']})", "#EF6C00"))
                else:
                    meta_bits.append((f"📦 stock: {pc['total']} deducted", "#2E7D32"))
        except Exception:
            pass
        if meta_bits:
            meta_row = QHBoxLayout()
            meta_row.setSpacing(6)
            for _text, _color in meta_bits:
                chip = QLabel(_text)
                chip.setStyleSheet(
                    f"color:{_color}; font-size:11px; font-weight:bold;"
                    f" background:{_color}18; border-radius:8px; padding:1px 7px;")
                meta_row.addWidget(chip)
            meta_row.addStretch()
            outer.addLayout(meta_row)

        # Photo thumbnails
        if self._images:
            photo_row = QHBoxLayout()
            for img in self._images[:6]:
                us         = img.get('upload_status', 'local')
                thumb_path = img.get('thumbnail_path') or img.get('file_path', '')
                th = PhotoThumb(
                    thumb_path,
                    image_id      = img.get('id', ''),
                    upload_status = us,
                )
                th.setToolTip(img.get('filename', ''))
                full_path   = img.get('file_path', '')
                image_id    = img.get('id', '')
                work_log_id = entry.get('id', '')
                th.clicked.connect(
                    lambda _=None, fp=full_path, iid=image_id, wid=work_log_id, s=us:
                        self._open_photo(fp, iid, wid, s)
                )
                photo_row.addWidget(th)
            if len(self._images) > 6:
                more = QLabel(f"+{len(self._images) - 6} more")
                more.setStyleSheet("color:#757575; font-size:11px;")
                photo_row.addWidget(more)
            photo_row.addStretch()
            outer.addLayout(photo_row)

        # Footer
        footer = QHBoxLayout()
        serial = entry.get('equipment_serial') or entry.get('serial_number') or ''
        if serial:
            ser_lbl = QLabel(f"S/N: {serial}")
            ser_lbl.setStyleSheet("color:#757575; font-size:10px;")
            footer.addWidget(ser_lbl)
        photo_count = len(self._images)
        if photo_count:
            pc_lbl = QLabel(f"📷 {photo_count}")
            pc_lbl.setStyleSheet("color:#757575; font-size:10px;")
            footer.addWidget(pc_lbl)
        footer.addStretch()

        edit_btn = QPushButton("✏️ Edit")
        edit_btn.setFixedHeight(24)
        edit_btn.setStyleSheet(
            "QPushButton{border:1px solid #90A4AE;border-radius:4px;"
            "padding:0 8px;font-size:11px;background:white;}"
            "QPushButton:hover{background:#E3F2FD;}"
        )
        edit_btn.clicked.connect(lambda: self.edit_requested.emit(self._entry['id']))
        footer.addWidget(edit_btn)

        del_btn = QPushButton("🗑 Delete")
        del_btn.setFixedHeight(24)
        del_btn.setStyleSheet(
            "QPushButton{border:1px solid #EF9A9A;border-radius:4px;"
            "padding:0 8px;font-size:11px;background:white;}"
            "QPushButton:hover{background:#FFEBEE;}"
        )
        del_btn.clicked.connect(lambda: self.delete_requested.emit(self._entry['id']))
        footer.addWidget(del_btn)
        outer.addLayout(footer)

    def _open_photo(self, file_path: str, image_id: str,
                    work_log_id: str, upload_status: str):
        if file_path and os.path.isfile(file_path):
            ImageViewerDialog(file_path, self).exec_()
            return
        if image_id:
            try:
                from services.image_service import download_remote_image
                path = download_remote_image(image_id, work_log_id)
                if path:
                    ImageViewerDialog(path, self).exec_()
                    return
            except Exception:
                pass
            QMessageBox.information(
                self, "Photo on server",
                "This photo is stored on the server but could not be downloaded.\n"
                "Make sure sync is configured and the server is reachable, then try again."
            )
        else:
            QMessageBox.warning(self, "Photo unavailable", "Photo file not found.")


# ── Edit / detail dialog ──────────────────────────────────────────────────────

class EditWorklogEntryDialog(QDialog):
    """Edit an existing work_log_entry (all fields + photos + structured parts)."""

    def __init__(self, entry_id: str, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Edit Field Log Entry")
        self.setMinimumSize(600, 680)
        self._entry_id     = entry_id
        self._images_to_delete: list = []
        self._new_photos:       list = []

        entry = get_worklog_entry(entry_id)
        if not entry:
            QMessageBox.critical(self, "Error", "Entry not found.")
            self.reject()
            return

        self._entry = entry
        self._build_ui(entry)

    def _build_ui(self, entry: dict):
        outer = QVBoxLayout(self)
        outer.setContentsMargins(6, 6, 6, 6)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setStyleSheet("QScrollArea{border:none;}")
        content = QWidget()
        lay = QVBoxLayout(content)
        lay.setSpacing(8)

        form = QFormLayout()
        form.setSpacing(6)

        self._date_edit = QDateEdit()
        self._date_edit.setCalendarPopup(True)
        self._date_edit.setDisplayFormat("yyyy-MM-dd")
        d = QDate.fromString(entry.get('log_date', ''), 'yyyy-MM-dd')
        self._date_edit.setDate(d if d.isValid() else QDate.currentDate())
        form.addRow("Date:", self._date_edit)

        self._proj_combo = QComboBox()
        self._proj_combo.addItem("— Unassigned (Mobile) —", userData=None)
        projects = get_all_projects()
        for p in projects:
            self._proj_combo.addItem(p.name, userData=p.id)
        cur_pid = entry.get('project_id')
        if cur_pid is not None:
            for i in range(self._proj_combo.count()):
                if self._proj_combo.itemData(i) == cur_pid:
                    self._proj_combo.setCurrentIndex(i)
                    break
        self._proj_combo.currentIndexChanged.connect(self._on_project_changed)
        form.addRow("Project:", self._proj_combo)

        self._cont_combo = QComboBox()
        self._cont_combo.addItem("— None —", userData=None)
        form.addRow("Container:", self._cont_combo)
        self._populate_containers(cur_pid, entry.get('container_id'))

        self._cat_combo = _category_combo(entry.get('category', 'maintenance'))
        form.addRow("Category:", self._cat_combo)

        self._fault_combo = _fault_combo(entry.get('fault_name', ''))
        form.addRow("Fault:", self._fault_combo)

        self._desc_input = QTextEdit()
        self._desc_input.setMinimumHeight(100)
        self._desc_input.setPlainText(entry.get('description', ''))
        form.addRow("Description:", self._desc_input)

        self._status_combo = _status_combo(entry.get('status', ''))
        form.addRow("Status:", self._status_combo)

        self._sap_input = QLineEdit(entry.get('sap_ticket', ''))
        self._sap_input.setPlaceholderText("SAP ticket / request no. (optional)")
        form.addRow("SAP Ticket:", self._sap_input)

        self._parts_input = QLineEdit(entry.get('spare_parts', ''))
        self._parts_input.setPlaceholderText("Free text (short). Structured tracking below.")
        form.addRow("Spare Parts:", self._parts_input)

        self._serial_input = QLineEdit(entry.get('equipment_serial', ''))
        form.addRow("Equipment S/N:", self._serial_input)

        self._loc_input = QLineEdit(entry.get('site_location', ''))
        form.addRow("Site/Location:", self._loc_input)

        tags = get_tags(entry['id'])
        self._tags_input = QLineEdit(', '.join(tags))
        self._tags_input.setPlaceholderText("comma-separated tags...")
        form.addRow("Tags:", self._tags_input)

        lay.addLayout(form)

        # Structured spare parts / stock panel
        self._parts_panel = SparePartsPanel(entry['id'])
        lay.addWidget(self._parts_panel)

        # Existing photos
        images = get_images_for_log(entry['id'])
        if images:
            photos_group = QGroupBox("Existing Photos")
            pgv = QVBoxLayout(photos_group)

            hdr = QHBoxLayout()
            save_all_btn = QPushButton("💾 Save all photos…")
            save_all_btn.setToolTip("Copy all photos of this entry to a folder on your computer")
            save_all_btn.clicked.connect(self._save_all_photos)
            hdr.addWidget(save_all_btn)
            hdr.addStretch()
            pgv.addLayout(hdr)

            pg_lay = QHBoxLayout()
            entry_id_for_dl = entry['id']
            for img in images:
                cell = QWidget()
                cl = QVBoxLayout(cell)
                cl.setContentsMargins(0, 0, 0, 0)
                cl.setSpacing(2)
                us         = img.get('upload_status', 'local')
                thumb_path = img.get('thumbnail_path') or img.get('file_path', '')
                th = PhotoThumb(
                    thumb_path,
                    image_id      = img.get('id', ''),
                    upload_status = us,
                )
                full_path   = img.get('file_path', '')
                image_id_dl = img.get('id', '')
                th.clicked.connect(
                    lambda _=None, fp=full_path, iid=image_id_dl, wid=entry_id_for_dl, s=us:
                        self._open_photo_dl(fp, iid, wid, s)
                )
                cl.addWidget(th)
                del_btn = QPushButton("✕")
                del_btn.setFixedSize(82, 18)
                del_btn.setStyleSheet(
                    "QPushButton{border:none;background:#EF9A9A;color:#B71C1C;"
                    "border-radius:3px;font-size:10px;}"
                    "QPushButton:hover{background:#EF5350;color:white;}"
                )
                img_id = img['id']
                del_btn.clicked.connect(
                    lambda _, iid=img_id, w=cell: self._mark_delete(iid, w)
                )
                cl.addWidget(del_btn)
                pg_lay.addWidget(cell)
            pg_lay.addStretch()
            pgv.addLayout(pg_lay)
            lay.addWidget(photos_group)

        # Add new photos
        self._photo_picker = PhotoPickerWidget()
        add_group = QGroupBox("Add New Photos")
        ag_lay = QVBoxLayout(add_group)
        ag_lay.addWidget(self._photo_picker)
        lay.addWidget(add_group)

        scroll.setWidget(content)
        outer.addWidget(scroll)

        buttons = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        buttons.accepted.connect(self._save)
        buttons.rejected.connect(self.reject)
        outer.addWidget(buttons)

    def _on_project_changed(self):
        pid = self._proj_combo.currentData()
        self._populate_containers(pid, None)

    def _populate_containers(self, project_id, current_container_id):
        self._cont_combo.blockSignals(True)
        self._cont_combo.clear()
        self._cont_combo.addItem("— None —", userData=None)
        if project_id is not None:
            try:
                zones = get_zones_for_project(project_id)
                for zone in zones:
                    blocks = get_blocks_for_zone(project_id, zone)
                    for block in blocks:
                        containers = get_containers_for_block(project_id, zone, block)
                        for c in containers:
                            label = f"Z{zone} B{block} · {c.container_type or '?'} #{c.container_index}"
                            self._cont_combo.addItem(label, userData=c.id)
            except Exception:
                pass
        if current_container_id is not None:
            for i in range(self._cont_combo.count()):
                if self._cont_combo.itemData(i) == current_container_id:
                    self._cont_combo.setCurrentIndex(i)
                    break
        self._cont_combo.blockSignals(False)

    def _open_photo_dl(self, file_path: str, image_id: str,
                       work_log_id: str, upload_status: str):
        if file_path and os.path.isfile(file_path):
            ImageViewerDialog(file_path, self).exec_()
            return
        if image_id:
            try:
                from services.image_service import download_remote_image
                path = download_remote_image(image_id, work_log_id)
                if path:
                    ImageViewerDialog(path, self).exec_()
                    return
            except Exception:
                pass
            QMessageBox.information(
                self, "Photo on server",
                "This photo is stored on the server but could not be downloaded.\n"
                "Make sure sync is configured and the server is reachable."
            )
        else:
            QMessageBox.warning(self, "Photo unavailable", "Photo file not found.")

    def _mark_delete(self, image_id: str, cell_widget: QWidget):
        self._images_to_delete.append(image_id)
        cell_widget.setVisible(False)

    def _save_all_photos(self):
        """Copy every photo of this entry to a folder the user chooses."""
        folder = QFileDialog.getExistingDirectory(
            self, "Choose a folder to save photos into")
        if not folder:
            return
        prefix = _entry_photo_prefix(self._entry)
        try:
            from services.image_service import export_entry_images
            res = export_entry_images(self._entry_id, folder, name_prefix=prefix)
        except Exception as ex:
            QMessageBox.critical(self, "Save Error", str(ex))
            return
        msg = f"Saved {res['saved']} of {res['total']} photo(s) to:\n{folder}"
        if res["failed"]:
            msg += "\n\nCould not save:\n" + "\n".join(str(f) for f in res["failed"])
            msg += ("\n\n(Photos stored only on the server need sync to be "
                    "configured and the server reachable.)")
        if QMessageBox.question(
            self, "Photos Saved", msg + "\n\nOpen folder?",
            QMessageBox.Yes | QMessageBox.No
        ) == QMessageBox.Yes:
            try:
                os.startfile(folder)
            except Exception:
                pass

    def _save(self):
        for iid in self._images_to_delete:
            try:
                delete_image(iid)
            except Exception:
                pass

        for p in self._photo_picker.get_paths():
            try:
                save_image(self._entry_id, p)
            except Exception as ex:
                QMessageBox.warning(
                    self, "Photo Error",
                    f"Could not save {os.path.basename(p)}:\n{ex}"
                )

        tags = [t.strip() for t in self._tags_input.text().split(',') if t.strip()]

        update_worklog_entry(
            self._entry_id,
            tags        = tags,
            log_date    = self._date_edit.date().toString('yyyy-MM-dd'),
            category    = self._cat_combo.currentData(),
            description = self._desc_input.toPlainText().strip(),
            fault_name  = self._fault_combo.currentText().strip(),
            status      = self._status_combo.currentData(),
            sap_ticket  = self._sap_input.text().strip(),
            spare_parts = self._parts_input.text().strip(),
            equipment_serial = self._serial_input.text().strip(),
            site_location    = self._loc_input.text().strip(),
            project_id       = self._proj_combo.currentData(),
            container_id     = self._cont_combo.currentData(),
        )
        self.accept()


# ── Monthly report dialog ─────────────────────────────────────────────────────

class MonthlyReportDialog(QDialog):
    """Options for the monthly technical-works report (Phase 4)."""

    MONTHS = ["January", "February", "March", "April", "May", "June",
              "July", "August", "September", "October", "November", "December"]

    def __init__(self, projects, default_project_id=None, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Monthly Technical Works Report")
        self.setMinimumWidth(420)
        self._build(projects, default_project_id)

    def _build(self, projects, default_project_id):
        lay = QVBoxLayout(self)
        form = QFormLayout()
        form.setSpacing(8)

        self._proj = QComboBox()
        self._proj.addItem("— All projects —", userData=None)
        for p in projects:
            self._proj.addItem(p.name, userData=p.id)
        if default_project_id is not None:
            for i in range(self._proj.count()):
                if self._proj.itemData(i) == default_project_id:
                    self._proj.setCurrentIndex(i)
                    break
        form.addRow("Project:", self._proj)

        now = QDate.currentDate()
        self._month = QComboBox()
        for i, m in enumerate(self.MONTHS, 1):
            self._month.addItem(m, userData=i)
        self._month.setCurrentIndex(now.month() - 1)
        form.addRow("Month:", self._month)

        self._year = QSpinBox()
        self._year.setRange(2020, 2100)
        self._year.setValue(now.year())
        form.addRow("Year:", self._year)

        self._fmt = QComboBox()
        self._fmt.addItem("PDF", userData="pdf")
        self._fmt.addItem("Word (.docx)", userData="docx")
        form.addRow("Format:", self._fmt)

        self._photos = QCheckBox("Include photo appendix")
        self._photos.setChecked(True)
        form.addRow("", self._photos)

        lay.addLayout(form)

        note = QLabel("Photos increase the file size. Uncheck for a compact report.")
        note.setWordWrap(True)
        note.setStyleSheet("color:#9E9E9E;font-size:10px;")
        lay.addWidget(note)

        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Ok).setText("Generate")
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        lay.addWidget(buttons)

    def values(self) -> dict:
        return {
            "project_id":   self._proj.currentData(),
            "project_name": self._proj.currentText() if self._proj.currentData() else "",
            "month":        self._month.currentData(),
            "year":         self._year.value(),
            "format":       self._fmt.currentData(),
            "include_photos": self._photos.isChecked(),
        }


# ── PAGE 1: Entry form ────────────────────────────────────────────────────────

class FieldLogEntryPage(QWidget):
    """Roomy form to create a new field-log entry with photos."""

    entry_saved = pyqtSignal()

    def __init__(self, parent=None):
        super().__init__(parent)
        self._containers = []
        self._build_ui()
        self._load_projects()

    def _build_ui(self):
        root = QVBoxLayout(self)
        root.setSpacing(6)

        title = QLabel("📸  Field Log — New Entry")
        title.setFont(QFont("Segoe UI", 14, QFont.Bold))
        root.addWidget(title)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setStyleSheet("QScrollArea{border:none;}")
        content = QWidget()
        content.setMaximumWidth(940)
        cl = QVBoxLayout(content)
        cl.setSpacing(8)

        # Two columns: Location | Details
        cols = QHBoxLayout()
        cols.setSpacing(10)

        # Location group
        loc_group = QGroupBox("Location")
        lf = QFormLayout(loc_group)
        lf.setSpacing(6)

        self.proj_combo = QComboBox()
        self.proj_combo.currentIndexChanged.connect(self._on_project_changed)
        lf.addRow("Project *:", self.proj_combo)

        self.date_edit = QDateEdit()
        self.date_edit.setDate(QDate.currentDate())
        self.date_edit.setCalendarPopup(True)
        self.date_edit.setDisplayFormat("yyyy-MM-dd")
        lf.addRow("Date *:", self.date_edit)

        self.zone_combo = QComboBox()
        self.zone_combo.currentIndexChanged.connect(self._on_zone_changed)
        lf.addRow("Zone:", self.zone_combo)

        self.block_combo = QComboBox()
        self.block_combo.currentIndexChanged.connect(self._on_block_changed)
        lf.addRow("Block:", self.block_combo)

        self.container_combo = QComboBox()
        self.container_combo.currentIndexChanged.connect(self._on_container_changed)
        lf.addRow("Container:", self.container_combo)

        self.serial_label = QLabel("—")
        self.serial_label.setStyleSheet("color:#1565C0;font-weight:bold;")
        lf.addRow("Serial #:", self.serial_label)

        self.loc_input = QLineEdit()
        self.loc_input.setPlaceholderText("e.g. Inverter room, rooftop…")
        lf.addRow("Site/Area:", self.loc_input)

        cols.addWidget(loc_group, 1)

        # Details group
        det_group = QGroupBox("Details")
        df = QFormLayout(det_group)
        df.setSpacing(6)

        self.cat_combo = _category_combo()
        df.addRow("Category *:", self.cat_combo)

        self.fault_combo = _fault_combo()
        df.addRow("Fault:", self.fault_combo)

        self.desc_input = QTextEdit()
        self.desc_input.setMinimumHeight(90)
        self.desc_input.setPlaceholderText("Describe the work or observation…")
        df.addRow("Description *:", self.desc_input)

        self.status_combo = _status_combo()
        df.addRow("Status:", self.status_combo)

        self.sap_input = QLineEdit()
        self.sap_input.setPlaceholderText("SAP ticket / request no. (optional)")
        df.addRow("SAP Ticket:", self.sap_input)

        self.parts_input = QLineEdit()
        self.parts_input.setPlaceholderText("Short note (structured tracking when editing)")
        df.addRow("Spare Parts:", self.parts_input)

        self.tags_input = QLineEdit()
        self.tags_input.setPlaceholderText("tag1, tag2 (optional)")
        df.addRow("Tags:", self.tags_input)

        self.serial_input = QLineEdit()
        self.serial_input.setPlaceholderText("Equipment serial (if known)")
        df.addRow("Equipment S/N:", self.serial_input)

        cols.addWidget(det_group, 1)
        cl.addLayout(cols)

        # Photos (full width)
        photo_group = QGroupBox("Photos (optional)")
        pg_lay = QVBoxLayout(photo_group)
        self.photo_picker = PhotoPickerWidget()
        pg_lay.addWidget(self.photo_picker)
        cl.addWidget(photo_group)

        # Buttons
        btn_row = QHBoxLayout()
        clear_btn = QPushButton("🔄 Clear")
        clear_btn.clicked.connect(self._clear_form)
        btn_row.addWidget(clear_btn)
        btn_row.addStretch()

        save_btn = QPushButton("✅  Save Entry")
        save_btn.setFixedHeight(40)
        save_btn.setMinimumWidth(200)
        save_btn.setStyleSheet(
            "QPushButton{background:#4CAF50;color:white;font-weight:bold;"
            "border-radius:4px;padding:0 16px;}"
            "QPushButton:hover{background:#388E3C;}"
        )
        save_btn.clicked.connect(self._save_entry)
        btn_row.addWidget(save_btn)
        cl.addLayout(btn_row)

        cl.addStretch()
        scroll.setWidget(content)
        root.addWidget(scroll)

        tip = QLabel("Records and the monthly report are on the 'Field Log — Records' page.")
        tip.setStyleSheet("color:#9E9E9E;font-size:11px;")
        root.addWidget(tip)

    # ── cascade ────────────────────────────────────────────────────────────
    def refresh_projects(self):
        self._load_projects()

    def _load_projects(self):
        projects = get_all_projects()
        self.proj_combo.blockSignals(True)
        self.proj_combo.clear()
        self.proj_combo.addItem("— Select —", userData=None)
        for p in projects:
            self.proj_combo.addItem(p.name, userData=p.id)
        self.proj_combo.blockSignals(False)
        self._on_project_changed()

    def _on_project_changed(self):
        self.zone_combo.blockSignals(True)
        self.zone_combo.clear()
        self.zone_combo.addItem("— any —", userData=None)
        pid = self.proj_combo.currentData()
        if pid:
            for z in get_zones_for_project(pid):
                self.zone_combo.addItem(f"Zone {z}", userData=z)
        self.zone_combo.blockSignals(False)
        self._on_zone_changed()

    def _on_zone_changed(self):
        self.block_combo.blockSignals(True)
        self.block_combo.clear()
        self.block_combo.addItem("— any —", userData=None)
        pid  = self.proj_combo.currentData()
        zone = self.zone_combo.currentData()
        if pid and zone:
            for b in get_blocks_for_zone(pid, zone):
                self.block_combo.addItem(f"Block {b}", userData=b)
        self.block_combo.blockSignals(False)
        self._on_block_changed()

    def _on_block_changed(self):
        self.container_combo.clear()
        self.container_combo.addItem("— any —", userData=None)
        self._containers = []
        pid   = self.proj_combo.currentData()
        zone  = self.zone_combo.currentData()
        block = self.block_combo.currentData()
        if pid and zone and block:
            self._containers = get_containers_for_block(pid, zone, block)
            for c in self._containers:
                self.container_combo.addItem(
                    f"C{c.container_index} ({c.container_type})", userData=c.id
                )
        self._on_container_changed()

    def _on_container_changed(self):
        cid = self.container_combo.currentData()
        if cid is None:
            self.serial_label.setText("—")
            return
        c = next((x for x in self._containers if x.id == cid), None)
        if c:
            self.serial_label.setText(c.serial_number or "N/A")
            if not self.serial_input.text():
                self.serial_input.setText(c.serial_number or "")
            if not self.loc_input.text():
                proj = self.proj_combo.currentText()
                zone = self.zone_combo.currentData() or ''
                block = self.block_combo.currentData() or ''
                ci    = c.container_index
                self.loc_input.setText(f"{proj} / Z{zone} / B{block} / C{ci}")

    # ── save ───────────────────────────────────────────────────────────────
    def _save_entry(self):
        pid = self.proj_combo.currentData()
        if pid is None:
            QMessageBox.warning(self, "Required", "Please select a project.")
            return
        desc = self.desc_input.toPlainText().strip()
        if not desc:
            QMessageBox.warning(self, "Required", "Description is required.")
            return

        cid    = self.container_combo.currentData()
        cat    = self.cat_combo.currentData()
        serial = self.serial_input.text().strip()
        loc    = self.loc_input.text().strip()

        if not loc and cid is not None:
            c = next((x for x in self._containers if x.id == cid), None)
            if c:
                z = self.zone_combo.currentData() or ''
                b = self.block_combo.currentData() or ''
                loc = f"{self.proj_combo.currentText()} / Z{z} / B{b} / C{c.container_index}"

        tags = [t.strip() for t in self.tags_input.text().split(',') if t.strip()]

        entry_id = save_worklog_entry(
            project_id       = pid,
            log_date         = self.date_edit.date().toString('yyyy-MM-dd'),
            description      = desc,
            category         = cat,
            container_id     = cid,
            equipment_serial = serial,
            site_location    = loc,
            tags             = tags,
            fault_name       = self.fault_combo.currentText().strip(),
            status           = self.status_combo.currentData(),
            sap_ticket       = self.sap_input.text().strip(),
            spare_parts      = self.parts_input.text().strip(),
        )

        photos = self.photo_picker.get_paths()
        failed = []
        for p in photos:
            try:
                save_image(entry_id, p)
            except Exception as ex:
                failed.append(f"{os.path.basename(p)}: {ex}")

        msg = f"Entry saved (ID: {entry_id[:8]}…)"
        if failed:
            msg += "\n\nFailed photos:\n" + "\n".join(failed)
            QMessageBox.warning(self, "Saved with warnings", msg)
        else:
            QMessageBox.information(self, "Saved", msg)

        self._clear_form()
        self.entry_saved.emit()

    def _clear_form(self):
        self.proj_combo.setCurrentIndex(0)
        self.date_edit.setDate(QDate.currentDate())
        self.desc_input.clear()
        self.tags_input.clear()
        self.serial_input.clear()
        self.loc_input.clear()
        self.cat_combo.setCurrentIndex(0)
        self.fault_combo.setCurrentText("")
        self.status_combo.setCurrentIndex(0)
        self.sap_input.clear()
        self.parts_input.clear()
        self.photo_picker.clear()


# ── PAGE 2: Records (timeline + export + report) ──────────────────────────────

class FieldLogRecordsPage(QWidget):
    """Filter bar + timeline of existing logs, export, and monthly report."""

    def __init__(self, parent=None):
        super().__init__(parent)
        self._build_ui()
        self._load_projects()

    def _build_ui(self):
        root = QVBoxLayout(self)
        root.setSpacing(6)

        title = QLabel("🗂  Field Log — Records")
        title.setFont(QFont("Segoe UI", 14, QFont.Bold))
        root.addWidget(title)

        # Filter bar
        flt_group = QGroupBox("Filter")
        flt_row = QHBoxLayout(flt_group)
        flt_row.setSpacing(6)

        flt_row.addWidget(QLabel("Project:"))
        self.flt_proj = QComboBox()
        self.flt_proj.setMinimumWidth(140)
        flt_row.addWidget(self.flt_proj)

        flt_row.addWidget(QLabel("From:"))
        self.flt_from = QDateEdit()
        self.flt_from.setDate(QDate.currentDate().addMonths(-1))
        self.flt_from.setCalendarPopup(True)
        self.flt_from.setDisplayFormat("yyyy-MM-dd")
        flt_row.addWidget(self.flt_from)

        flt_row.addWidget(QLabel("To:"))
        self.flt_to = QDateEdit()
        self.flt_to.setDate(QDate.currentDate())
        self.flt_to.setCalendarPopup(True)
        self.flt_to.setDisplayFormat("yyyy-MM-dd")
        flt_row.addWidget(self.flt_to)

        flt_row.addWidget(QLabel("Category:"))
        self.flt_cat = QComboBox()
        self.flt_cat.addItem("All", userData=None)
        for key in CATEGORIES:
            self.flt_cat.addItem(CATEGORY_LABELS[key], userData=key)
        flt_row.addWidget(self.flt_cat)

        flt_row.addWidget(QLabel("Search:"))
        self.flt_search = QLineEdit()
        self.flt_search.setPlaceholderText("text…")
        self.flt_search.setMaximumWidth(120)
        flt_row.addWidget(self.flt_search)

        load_btn = QPushButton("🔍 Load")
        load_btn.setStyleSheet(
            "QPushButton{background:#1976D2;color:white;"
            "border-radius:4px;padding:4px 12px;}"
            "QPushButton:hover{background:#1565C0;}"
        )
        load_btn.clicked.connect(self._load_timeline)
        flt_row.addWidget(load_btn)

        export_btn = QPushButton("📤 Export ▾")
        export_btn.setStyleSheet(
            "QPushButton{border:1px solid #90A4AE;border-radius:4px;padding:4px 12px;}"
            "QPushButton:hover{background:#E3F2FD;}"
        )
        export_btn.clicked.connect(self._show_export_menu)
        flt_row.addWidget(export_btn)

        report_btn = QPushButton("📄 Monthly Report")
        report_btn.setStyleSheet(
            "QPushButton{background:#455A64;color:white;"
            "border-radius:4px;padding:4px 12px;}"
            "QPushButton:hover{background:#37474F;}"
        )
        report_btn.clicked.connect(self._open_report_dialog)
        flt_row.addWidget(report_btn)

        flt_row.addStretch()
        root.addWidget(flt_group)

        self.result_label = QLabel("Click 'Load' to view logs.")
        self.result_label.setStyleSheet("color:#757575; font-style:italic;")
        root.addWidget(self.result_label)

        self._scroll = QScrollArea()
        self._scroll.setWidgetResizable(True)
        self._scroll.setStyleSheet("QScrollArea{border:none; background:#FAFAFA;}")

        self._timeline_widget = QWidget()
        self._timeline_widget.setObjectName("TimelineWidget")
        self._timeline_layout = QVBoxLayout(self._timeline_widget)
        self._timeline_layout.setContentsMargins(4, 4, 4, 4)
        self._timeline_layout.setSpacing(4)
        self._timeline_layout.addStretch()

        self._scroll.setWidget(self._timeline_widget)
        root.addWidget(self._scroll)

    # ── projects ───────────────────────────────────────────────────────────
    def refresh_projects(self):
        self._load_projects()

    def _load_projects(self):
        projects = get_all_projects()
        self.flt_proj.blockSignals(True)
        self.flt_proj.clear()
        self.flt_proj.addItem("— Select —", userData=None)
        for p in projects:
            self.flt_proj.addItem(p.name, userData=p.id)
        self.flt_proj.blockSignals(False)

    # ── timeline ───────────────────────────────────────────────────────────
    def _load_timeline(self):
        pid      = self.flt_proj.currentData()
        category = self.flt_cat.currentData()
        search   = self.flt_search.text().strip() or None
        date_from = self.flt_from.date().toString('yyyy-MM-dd')
        date_to   = self.flt_to.date().toString('yyyy-MM-dd')

        entries = get_worklog_entries(
            project_id  = pid,
            date_from   = date_from,
            date_to     = date_to,
            category    = category,
            search_text = search,
        )
        self._rebuild_timeline(entries)
        count = len(entries)
        self.result_label.setText(
            f"{count} entr{'ies' if count != 1 else 'y'} found."
        )

    def _rebuild_timeline(self, entries: list):
        while self._timeline_layout.count():
            item = self._timeline_layout.takeAt(0)
            if item.widget():
                item.widget().deleteLater()

        if not entries:
            empty = QLabel("No entries found.")
            empty.setAlignment(Qt.AlignCenter)
            empty.setStyleSheet("color:#9E9E9E; font-style:italic; padding:40px;")
            self._timeline_layout.addWidget(empty)
            self._timeline_layout.addStretch()
            return

        last_date = None
        for entry in entries:
            entry_date = entry.get('log_date', '')
            if entry_date != last_date:
                last_date = entry_date
                date_hdr = QLabel(f"  📅  {entry_date}")
                date_hdr.setStyleSheet(
                    "color:#37474F; font-weight:bold; font-size:13px;"
                    " padding:8px 4px 2px 4px; background:transparent;"
                )
                self._timeline_layout.addWidget(date_hdr)

            images = get_images_for_log(entry['id'])
            card   = LogCard(entry, images, parent=self._timeline_widget)
            card.edit_requested.connect(self._on_edit)
            card.delete_requested.connect(self._on_delete)
            self._timeline_layout.addWidget(card)

        self._timeline_layout.addStretch()

    # ── edit / delete ──────────────────────────────────────────────────────
    def _on_edit(self, entry_id: str):
        dlg = EditWorklogEntryDialog(entry_id, parent=self)
        if dlg.exec_() == QDialog.Accepted:
            self._load_timeline()

    def _on_delete(self, entry_id: str):
        answer = QMessageBox.question(
            self, "Delete?",
            "Delete this log entry and all attached photos?",
            QMessageBox.Yes | QMessageBox.No, QMessageBox.No
        )
        if answer == QMessageBox.Yes:
            delete_images_for_log(entry_id)
            delete_worklog_entry(entry_id)
            self._load_timeline()

    # ── export ─────────────────────────────────────────────────────────────
    def _show_export_menu(self):
        menu = QMenu(self)
        menu.addAction("📊  Export to Excel (.xlsx)", self._export_excel)
        menu.addAction("📄  Export to PDF (table)", self._export_pdf)
        menu.addSeparator()
        menu.addAction("🖼  Save photos (current filter)…", self._export_photos)
        btn = self.sender()
        menu.exec_(btn.mapToGlobal(btn.rect().bottomLeft()))

    def _export_photos(self):
        """Copy every photo of the currently-filtered entries to a chosen folder
        (one subfolder per entry). Remote-only photos are downloaded first."""
        entries = get_worklog_entries(
            project_id  = self.flt_proj.currentData(),
            date_from   = self.flt_from.date().toString('yyyy-MM-dd'),
            date_to     = self.flt_to.date().toString('yyyy-MM-dd'),
            category    = self.flt_cat.currentData(),
            search_text = self.flt_search.text().strip() or None,
        )
        with_photos = [e for e in entries if get_images_for_log(e['id'])]
        if not with_photos:
            QMessageBox.information(self, "Save photos",
                                    "No photos in the current filter.")
            return
        if QMessageBox.question(
            self, "Save photos",
            f"{len(with_photos)} entr{'ies' if len(with_photos) != 1 else 'y'} "
            f"with photos in the current filter.\n"
            f"Choose a folder — each entry gets its own subfolder.\n\nContinue?",
            QMessageBox.Yes | QMessageBox.No, QMessageBox.Yes
        ) != QMessageBox.Yes:
            return
        root = QFileDialog.getExistingDirectory(self, "Choose destination folder")
        if not root:
            return

        from services.image_service import export_entry_images, _safe_name
        total_saved, total_failed, seen = 0, 0, set()
        for e in with_photos:
            prefix = _entry_photo_prefix(e)
            sub = _safe_name(prefix)
            # de-duplicate subfolder names across entries with the same date/loc
            base_sub, k = sub, 1
            while sub in seen:
                sub = f"{base_sub}_{k}"; k += 1
            seen.add(sub)
            try:
                res = export_entry_images(e['id'], os.path.join(root, sub))
                total_saved += res['saved']
                total_failed += len(res['failed'])
            except Exception:
                total_failed += 1

        msg = f"Saved {total_saved} photo(s) from {len(with_photos)} entr" \
              f"{'ies' if len(with_photos) != 1 else 'y'} into:\n{root}"
        if total_failed:
            msg += f"\n\n{total_failed} photo(s) could not be saved (server-only, " \
                   f"needs sync)."
        if QMessageBox.question(
            self, "Photos Saved", msg + "\n\nOpen folder?",
            QMessageBox.Yes | QMessageBox.No
        ) == QMessageBox.Yes:
            try:
                os.startfile(root)
            except Exception:
                pass

    def _get_current_entries(self):
        from services.worklog_entry_service import get_worklog_entries_dataframe
        pid      = self.flt_proj.currentData()
        category = self.flt_cat.currentData()
        return get_worklog_entries_dataframe(
            project_id = pid,
            date_from  = self.flt_from.date().toString('yyyy-MM-dd'),
            date_to    = self.flt_to.date().toString('yyyy-MM-dd'),
            category   = category,
        )

    def _export_excel(self):
        df = self._get_current_entries()
        if df.empty:
            QMessageBox.information(self, "Export", "No entries to export.")
            return
        path, _ = QFileDialog.getSaveFileName(
            self, "Save Excel", "field_log.xlsx", "Excel files (*.xlsx)")
        if not path:
            return
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils.dataframe import dataframe_to_rows

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Field Log"
            header_fill = PatternFill("solid", fgColor="1A2B45")
            header_font = Font(bold=True, color="FFFFFF")
            for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
                ws.append(row)
                if r_idx == 1:
                    for cell in ws[1]:
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = Alignment(horizontal="center")
            for col in ws.columns:
                max_len = max((len(str(c.value or "")) for c in col), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 60)
            wb.save(path)
            if QMessageBox.question(
                self, "Exported", f"Saved to:\n{path}\n\nOpen file?",
                QMessageBox.Yes | QMessageBox.No
            ) == QMessageBox.Yes:
                os.startfile(path)
        except ImportError:
            QMessageBox.critical(self, "Missing library", "Install openpyxl:\n  pip install openpyxl")
        except Exception as ex:
            QMessageBox.critical(self, "Export Error", str(ex))

    def _export_pdf(self):
        df = self._get_current_entries()
        if df.empty:
            QMessageBox.information(self, "Export", "No entries to export.")
            return
        path, _ = QFileDialog.getSaveFileName(
            self, "Save PDF", "field_log.pdf", "PDF files (*.pdf)")
        if not path:
            return
        try:
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib import colors
            from reportlab.platypus import (
                SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer)
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.units import cm

            doc  = SimpleDocTemplate(path, pagesize=landscape(A4),
                                     leftMargin=1*cm, rightMargin=1*cm,
                                     topMargin=1.5*cm, bottomMargin=1.5*cm)
            styles = getSampleStyleSheet()
            story  = []
            story.append(Paragraph("BESS Field Log Report", styles["Title"]))
            date_range = (f"{self.flt_from.date().toString('yyyy-MM-dd')} – "
                          f"{self.flt_to.date().toString('yyyy-MM-dd')}")
            story.append(Paragraph(date_range, styles["Normal"]))
            story.append(Spacer(1, 0.4*cm))
            pdf_cols = ["Date", "Project", "Category", "Location", "Description"]
            pdf_cols = [c for c in pdf_cols if c in df.columns]
            sub = df[pdf_cols].fillna("")
            data = [pdf_cols] + sub.values.tolist()
            col_widths = [2.5*cm, 3.5*cm, 3*cm, 4*cm, 13*cm][:len(pdf_cols)]
            tbl = Table(data, colWidths=col_widths, repeatRows=1)
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1A2B45")),
                ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
                ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE",   (0, 0), (-1, 0), 9),
                ("FONTSIZE",   (0, 1), (-1, -1), 8),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
                ("GRID",       (0, 0), (-1, -1), 0.3, colors.HexColor("#CBD5E1")),
                ("VALIGN",     (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(tbl)
            doc.build(story)
            if QMessageBox.question(
                self, "Exported", f"Saved to:\n{path}\n\nOpen file?",
                QMessageBox.Yes | QMessageBox.No
            ) == QMessageBox.Yes:
                os.startfile(path)
        except ImportError:
            QMessageBox.critical(self, "Missing library", "Install reportlab:\n  pip install reportlab")
        except Exception as ex:
            QMessageBox.critical(self, "Export Error", str(ex))

    # ── monthly report (Phase 4) ────────────────────────────────────────────
    def _open_report_dialog(self):
        projects = get_all_projects()
        dlg = MonthlyReportDialog(projects, self.flt_proj.currentData(), parent=self)
        if dlg.exec_() != QDialog.Accepted:
            return
        v = dlg.values()
        ext = "pdf" if v["format"] == "pdf" else "docx"
        default_name = (f"Technical_Works_"
                        f"{v['year']}_{v['month']:02d}.{ext}")
        path, _ = QFileDialog.getSaveFileName(
            self, "Save report", default_name,
            "PDF (*.pdf)" if ext == "pdf" else "Word (*.docx)")
        if not path:
            return
        try:
            from services.worklog_report_service import generate_worklog_report
            generate_worklog_report(
                project_id     = v["project_id"],
                year           = v["year"],
                month          = v["month"],
                output_path    = path,
                output_format  = v["format"],
                include_photos = v["include_photos"],
                project_name   = v["project_name"],
            )
            if QMessageBox.question(
                self, "Done", f"Report saved:\n{path}\n\nOpen file?",
                QMessageBox.Yes | QMessageBox.No
            ) == QMessageBox.Yes:
                os.startfile(path)
        except Exception as ex:
            QMessageBox.critical(self, "Report error", str(ex))


# ── Backward-compatible alias ─────────────────────────────────────────────────
# The old single-page class name still works (points to the records page).
WorklogEntryForm = FieldLogRecordsPage

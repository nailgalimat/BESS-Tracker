"""services/stock_excel_service.py — spare-parts stock counts in from Excel

The owner receives stock snapshots as SAP MB52 exports and used to type every
line in by hand. This module reads such a file (and the app's own export, so
export → edit → import round-trips), works out what each row would do to a
warehouse, and applies only the rows that were confirmed.

Three steps, deliberately separate so the UI can show the middle one before
anything is written:

    parse(path)                  -> what the file says
    plan(warehouse_id, rows)     -> what it would change (read-only)
    apply(warehouse_id, rows, r) -> the writes

Two rules this module exists to keep:

* **A field the file does not carry is left alone.** An absent field means "no
  opinion", never "blank it" — the rule `_ENTRY_OPT_COLS` in `sync_client`
  spells out. For `min_quantity`, the low-stock alert threshold the owner sets
  by hand, that means three distinct states: no `Min Qty` column at all (a SAP
  export) leaves every threshold untouched; a `Min Qty` column with an empty
  cell leaves that row's threshold untouched; a filled cell is applied.
  `set_stock_level` is still never called from here — it writes quantity,
  threshold and unit in one go, so its `min_quantity=0` default would silently
  zero every alert in the warehouse. `ensure_stock_item` and `set_min_quantity`
  each write one field, so this module can only change what the file said.
  `Low Stock Alert` stays ignored: it is a computed flag, not an input.
* **The quantity moves by a transaction.** A counted quantity is applied as the
  difference against what is on the shelf now, through
  `stock_service.record_transaction` (`IN` when the count went up, `OUT` when it
  went down), so the audit trail still explains every quantity arithmetically
  instead of showing an unexplained jump.

There is no PyQt import here (three-layer rule); `ui/stock_page.py` is the thin
layer over this.
"""

import os
from datetime import date as _date
from typing import List, Optional

import openpyxl

from services import material_service as _materials
from services import stock_service as _stock

# How far down a sheet to look for the header row. SAP puts four junk rows and a
# date stamp above it; the app's own template puts an instruction row.
HEADER_SCAN_ROWS = 30

# Quantities are floats, so "the same" needs a tolerance rather than ==.
EPS = 1e-9

# Column synonyms, matched against the *stripped, lower-cased* header text.
# Order matters: the first name found wins, so the SAP spelling comes first.
_MATERIAL_HEADERS = ('material', 'material #', 'material number', 'material no',
                     'material_number', 'mat #', 'mat no', 'code')
_QTY_HEADERS      = ('unrestr.', 'unrestr', 'unrestricted', 'unrestricted stock',
                     'qty', 'quantity', 'stock', 'on hand')
_UNIT_HEADERS     = ('bun', 'unit', 'uom', 'base unit', 'unit of measure')
_MATDESC_HEADERS  = ('material description', 'material desc')
_DESC_HEADERS     = ('description', 'desc', 'name')
_PLANT_HEADERS    = ('plnt', 'plant')
_LOCATION_HEADERS = ('location', 'sloc', 'storage location', 'stor. loc.')
_WAREHOUSE_HEADERS = ('warehouse',)
_TRANSFER_HEADERS = ('trans./tfr', 'trans./ tfr', 'trans/tfr', 'in transfer')
# The alert threshold. An input when the column is there, per-row: see the
# module docstring. export_stock_excel writes this column, so export → set
# thresholds in Excel → import is a supported way to set 67 of them at once.
_MINQTY_HEADERS   = ('min qty', 'min quantity', 'min. qty', 'minimum',
                     'min stock', 'min_quantity')
# Read but deliberately never applied: 'Low Stock Alert' is computed from the
# threshold, not a field anyone can set.
_IGNORED_HEADERS  = ('low stock alert', 'notes')

# What write_template puts in the file. Kept in step with export_stock_excel's
# "Current Stock" sheet so all three files import through the same code path.
TEMPLATE_HEADERS = ('Material #', 'Description', 'Unit', 'Qty', 'Min Qty')
TEMPLATE_SHEET = 'Stock Count'
TEMPLATE_NOTE = (
    'Fill in one row per material. Qty = how many are on the shelf now (not the '
    'change). Min Qty is the low-stock alert threshold: fill it in to set it, '
    'leave it empty to keep the one already in the app. Delete these example '
    'rows before importing.'
)


# ── header / cell helpers ─────────────────────────────────────────────────────

def _norm(value) -> str:
    """Header text as we match it: no padding, single spaces, lower case.

    Every cell in an SAP MB52 export is left-padded with spaces so the columns
    line up visually, so nothing can be compared before it is stripped.
    """
    if value is None:
        return ''
    return ' '.join(str(value).split()).strip().lower()


def _text(row, index: Optional[int]) -> str:
    """Stripped text of one cell, tolerant of short rows and None."""
    if index is None or index < 0 or index >= len(row):
        return ''
    value = row[index]
    if value is None:
        return ''
    return str(value).replace(' ', ' ').strip()


def _raw(row, index: Optional[int]):
    if index is None or index < 0 or index >= len(row):
        return None
    return row[index]


def _looks_like_header(row) -> bool:
    """Is this row the header row of a stock list?

    A material column plus either a plant column (SAP) or a quantity column
    (the app's own layout). The app export's second sheet — the transaction
    log — also has Material #/Qty columns, so it is rejected explicitly.
    """
    cells = [_norm(c) for c in row]
    if not any(c in _MATERIAL_HEADERS for c in cells):
        return False
    if 'type' in cells and 'date' in cells:
        return False
    return any(c in _QTY_HEADERS or c in _PLANT_HEADERS for c in cells)


def _pick(headers: List[str], names) -> Optional[int]:
    """Index of the first column whose header is exactly one of `names`."""
    for want in names:
        for index, header in enumerate(headers):
            if header == want:
                return index
    return None


def _parse_quantity(raw):
    """(value, None) on success, (None, reason) when the cell cannot be read.

    Tolerant on purpose — SAP quantities arrive as padded text, sometimes with a
    thousands separator or a trailing minus. What it will not do is guess: an
    unreadable cell is reported so the row can be shown as a problem instead of
    being imported as zero.
    """
    if raw is None:
        return None, 'quantity cell is empty'
    if isinstance(raw, bool):
        return None, 'quantity is not a number: {!r}'.format(raw)

    if isinstance(raw, (int, float)):
        qty = float(raw)
    else:
        text = str(raw).replace(' ', ' ').strip()
        if not text:
            return None, 'quantity cell is empty'
        negative = text.endswith('-')            # SAP writes "5-" for -5
        if negative:
            text = text[:-1].strip()
        text = text.replace(' ', '')
        if ',' in text and '.' in text:
            # whichever separator comes last is the decimal one
            text = (text.replace('.', '') if text.rindex(',') > text.rindex('.')
                    else text.replace(',', ''))
            text = text.replace(',', '.')
        elif ',' in text:
            head, _, tail = text.rpartition(',')
            # "1,500" is fifteen hundred; "1,5" is one and a half
            text = (head + tail) if (head and len(tail) == 3) else text.replace(',', '.')
        try:
            qty = float(text)
        except ValueError:
            return None, 'quantity is not a number: {!r}'.format(str(raw).strip())
        if negative:
            qty = -qty

    if qty != qty or qty in (float('inf'), float('-inf')):
        return None, 'quantity is not a finite number'
    if qty < 0:
        return None, 'quantity is negative ({:g})'.format(qty)
    return qty, None


# ── parse ─────────────────────────────────────────────────────────────────────

def parse(path: str) -> dict:
    """Read a stock workbook and return what we understood of it.

    Never raises on a merely odd file — an unreadable row becomes an entry in
    `problems` and the rest is still offered. A file with no recognisable header
    row, or no quantity column, is a clear error (`ValueError`).

    Returns:
        {
          'file', 'sheet', 'layout' ('sap'|'template'), 'header_row' (1-based),
          'plants', 'locations', 'warehouses',   # what the file says it is about
          'has_min_column': bool,                # is the threshold an input here?
          'ignored_columns',                     # e.g. Low Stock Alert
          'rows':     [{'row','material_number','description','unit',
                        'quantity','in_transfer','min_quantity'}],
          'problems': [{'row','kind','material_number','message'}],
          'skipped_blank': int,
        }

    A row's `min_quantity` is None when the file has no opinion about the
    threshold — either the column is absent or that cell is empty — and a
    number when it should be applied.
    """
    try:
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:                          # corrupt / not a workbook
        raise ValueError('Could not read Excel file: {}'.format(exc))

    try:
        found = None
        for sheet in book.worksheets:
            grid = [r for r in sheet.iter_rows(values_only=True)]
            for index in range(min(HEADER_SCAN_ROWS, len(grid))):
                if _looks_like_header(grid[index]):
                    found = (sheet.title, grid, index)
                    break
            if found:
                break

        if not found:
            raise ValueError(
                'No stock table found in this file.\n\n'
                'Expected a header row with a "Material" column (SAP MB52 '
                'export) or a "Material #" column (this app\'s export or '
                'template), somewhere in the first {} rows.'
                .format(HEADER_SCAN_ROWS))

        sheet_name, grid, header_index = found
        headers = [_norm(c) for c in grid[header_index]]

        i_material = _pick(headers, _MATERIAL_HEADERS)
        i_quantity = _pick(headers, _QTY_HEADERS)
        i_unit = _pick(headers, _UNIT_HEADERS)
        i_plant = _pick(headers, _PLANT_HEADERS)
        i_location = _pick(headers, _LOCATION_HEADERS)
        i_warehouse = _pick(headers, _WAREHOUSE_HEADERS)
        i_transfer = _pick(headers, _TRANSFER_HEADERS)
        i_min = _pick(headers, _MINQTY_HEADERS)
        # In an SAP export column F is the material description and column G is
        # the storage-location description — so the specific name wins.
        i_description = _pick(headers, _MATDESC_HEADERS)
        if i_description is None:
            i_description = _pick(headers, _DESC_HEADERS)

        if i_quantity is None:
            raise ValueError(
                'Found a material column but no quantity column on sheet '
                '"{}".\n\nExpected "Unrestr." (SAP) or "Qty".'.format(sheet_name))

        layout = ('sap' if (i_plant is not None or 'unrestr.' in headers)
                  else 'template')
        ignored = [' '.join(str(c).split()) for c in grid[header_index]
                   if _norm(c) in _IGNORED_HEADERS]

        rows, problems = [], []
        plants, locations, warehouses = [], [], []
        seen = {}
        skipped_blank = 0

        for offset, raw_row in enumerate(grid[header_index + 1:]):
            number = header_index + 2 + offset          # 1-based Excel row

            if not any(_text(raw_row, i) for i in range(len(raw_row))):
                skipped_blank += 1
                continue
            if _looks_like_header(raw_row):             # a repeated header block
                skipped_blank += 1
                continue

            material = _text(raw_row, i_material).upper()
            if not material:
                problems.append({
                    'row': number, 'kind': 'no_material', 'material_number': '',
                    'message': 'no material number — row ignored'})
                continue

            for index, bucket in ((i_plant, plants), (i_location, locations),
                                  (i_warehouse, warehouses)):
                value = _text(raw_row, index)
                if value and value not in bucket:
                    bucket.append(value)

            if material in seen:
                problems.append({
                    'row': number, 'kind': 'duplicate',
                    'material_number': material,
                    'message': 'appears again (first seen on row {}) — this '
                               'row is ignored'.format(seen[material])})
                continue
            seen[material] = number

            quantity, reason = _parse_quantity(_raw(raw_row, i_quantity))
            if quantity is None:
                problems.append({
                    'row': number, 'kind': 'quantity',
                    'material_number': material, 'message': reason})
                continue

            in_transfer, _ = _parse_quantity(_raw(raw_row, i_transfer)) \
                if i_transfer is not None else (None, None)

            # The threshold: None means "the file has no opinion", so an absent
            # column and an empty cell both leave the stored one alone. A cell
            # that is filled but unreadable is a problem of its own — it must
            # not cost the row its quantity, which parsed fine.
            min_quantity = None
            if i_min is not None and _text(raw_row, i_min):
                min_quantity, min_reason = _parse_quantity(_raw(raw_row, i_min))
                if min_quantity is None:
                    problems.append({
                        'row': number, 'kind': 'min_quantity',
                        'material_number': material,
                        'message': 'minimum stock ignored, threshold left as it '
                                   'is — {}'.format(min_reason)})

            rows.append({
                'row': number,
                'material_number': material,
                'description': _text(raw_row, i_description),
                'unit': _text(raw_row, i_unit),
                'quantity': quantity,
                'in_transfer': in_transfer or 0.0,
                'min_quantity': min_quantity,
            })

        return {
            'file': os.path.basename(path),
            'sheet': sheet_name,
            'layout': layout,
            'header_row': header_index + 1,
            'plants': plants,
            'locations': locations,
            'warehouses': warehouses,
            'has_min_column': i_min is not None,
            'ignored_columns': ignored,
            'rows': rows,
            'problems': problems,
            'skipped_blank': skipped_blank,
        }
    finally:
        book.close()


# ── plan ──────────────────────────────────────────────────────────────────────

def plan(warehouse_id: int, rows: List[dict]) -> List[dict]:
    """Classify each parsed row against the warehouse — writes nothing.

    Each returned row carries its parsed fields plus:
        status      'new' | 'unchanged' | 'changed' — 'changed' if the quantity
                    or the threshold would move, so a file that only sets
                    thresholds is a real import and not an empty no-op
        before      quantity on the shelf now (0.0 for a material not stocked yet)
        after       quantity the file counted
        delta       after - before
        min_before  threshold now
        min_after   threshold afterwards (= min_before when the file is silent)
        min_changed would this row move the threshold?

    The row's parsed `min_quantity` is left exactly as parse() set it — None
    where the file has no opinion. apply() reads that, not `min_after`: a plan
    built before someone edited a threshold by hand must not write the value it
    saw back over theirs.
    """
    current = {r['material_number']: r for r in _stock.get_stock(warehouse_id)}
    out = []
    for row in rows:
        existing = current.get(row['material_number'])
        before = float(existing['quantity'] or 0) if existing else 0.0
        after = float(row['quantity'])
        delta = after - before

        min_before = float(existing['min_quantity'] or 0) if existing else 0.0
        wanted = row.get('min_quantity')
        min_after = min_before if wanted is None else float(wanted)
        min_changed = abs(min_after - min_before) > EPS

        if existing is None:
            status = 'new'
        elif abs(delta) < EPS and not min_changed:
            status = 'unchanged'
        else:
            status = 'changed'

        planned = dict(row)
        planned.update({
            'status': status,
            'before': before,
            'after': after,
            'delta': 0.0 if abs(delta) < EPS else delta,
            'min_before': min_before,
            'min_after': min_after,
            'min_changed': min_changed,
            'current_description': (existing or {}).get('description', '') or '',
        })
        out.append(planned)
    return out


def summarise(plan_rows: List[dict]) -> dict:
    """Counts per status — for the preview's one-line summary."""
    return {
        'new': sum(1 for r in plan_rows if r['status'] == 'new'),
        'changed': sum(1 for r in plan_rows if r['status'] == 'changed'),
        'unchanged': sum(1 for r in plan_rows if r['status'] == 'unchanged'),
        'thresholds': sum(1 for r in plan_rows if r.get('min_changed')),
        'total': len(plan_rows),
    }


# ── apply ─────────────────────────────────────────────────────────────────────

def apply(warehouse_id: int, plan_rows: List[dict], reference: str,
          transaction_date: Optional[str] = None,
          project_id: Optional[int] = None) -> dict:
    """Write the given plan rows into the warehouse. Returns counts.

    `plan_rows` is already what the user confirmed — this function writes every
    row it is given.

    The counted quantity is reached by a difference through `record_transaction`
    (`IN` / `OUT`), and the difference is measured against the quantity on the
    shelf *now*, re-read per row, so a plan built a few minutes ago cannot
    overshoot if stock moved meanwhile.

    The alert threshold is written only for a row whose parsed `min_quantity` is
    a number — i.e. the file carried a `Min Qty` column and that cell was filled.
    It reads the row's own parsed value rather than the plan's `min_after`, so a
    plan built before someone edited a threshold by hand cannot write the stale
    value back. `set_stock_level` is deliberately not called from here at all:
    it writes quantity, threshold and unit together, and its `min_quantity=0`
    default would zero every alert in the warehouse.

    Descriptions and units are merged into the `materials` catalogue; an empty
    cell never blanks a description that is already there.
    """
    ref = (reference or '').strip() or 'STOCK IMPORT'
    day = transaction_date or _date.today().isoformat()
    if project_id is None:
        project_id = _warehouse_project(warehouse_id)

    counts = {'created': 0, 'updated': 0, 'unchanged': 0, 'transactions': 0,
              'thresholds_set': 0, 'materials_created': 0,
              'materials_updated': 0, 'quantity_in': 0.0, 'quantity_out': 0.0}

    for row in plan_rows:
        material = (row.get('material_number') or '').strip().upper()
        if not material:
            continue
        target = float(row.get('after', row.get('quantity', 0)) or 0)
        unit = (row.get('unit') or '').strip()

        was_new = _stock.ensure_stock_item(warehouse_id, material, unit)

        held = _stock.get_stock_quantity(warehouse_id, material) or {}
        before = float(held.get('quantity') or 0)
        delta = target - before

        if abs(delta) > EPS:
            _stock.record_transaction(
                warehouse_id=warehouse_id,
                material_number=material,
                transaction_type='IN' if delta > 0 else 'OUT',
                quantity=abs(delta),
                transaction_date=day,
                project_id=project_id,
                reference=ref,
                notes='Stock count import: {:g} → {:g}'.format(before, target),
            )
            counts['transactions'] += 1
            if delta > 0:
                counts['quantity_in'] += delta
            else:
                counts['quantity_out'] += -delta

        # The threshold: only when the file actually carried one for this row.
        threshold_moved = False
        wanted_min = row.get('min_quantity')
        if wanted_min is not None:
            threshold_moved = _stock.set_min_quantity(
                warehouse_id, material, float(wanted_min))
            if threshold_moved:
                counts['thresholds_set'] += 1

        if was_new:
            counts['created'] += 1
        elif abs(delta) > EPS or threshold_moved:
            counts['updated'] += 1
        else:
            counts['unchanged'] += 1

        _upsert_material(material, row.get('description') or '', unit, counts)

    return counts


def _warehouse_project(warehouse_id: int) -> Optional[int]:
    """The project a warehouse belongs to, so the audit row carries it."""
    for warehouse in _stock.get_all_warehouses():
        if warehouse['id'] == warehouse_id:
            return warehouse.get('project_id')
    return None


def _upsert_material(material_number: str, description: str, unit: str,
                     counts: dict):
    """Merge the file's description/unit into the materials catalogue.

    An empty cell means "the file says nothing", so it never blanks a
    description or unit that is already in the catalogue — the same rule the
    rest of the app follows for absent fields. Notes are left untouched.
    """
    description = (description or '').strip()
    unit = (unit or '').strip()
    if not description and not unit:
        return

    existing = _materials.get_material(material_number)
    if existing is None:
        _materials.save_material(material_number, description, unit, '')
        counts['materials_created'] += 1
        return

    had_description = (existing.get('description') or '').strip()
    had_unit = (existing.get('unit') or '').strip()
    want_description = description or had_description
    want_unit = unit or had_unit
    if want_description != had_description or want_unit != had_unit:
        _materials.save_material(material_number, want_description, want_unit,
                                 existing.get('notes') or '')
        counts['materials_updated'] += 1


# ── template ──────────────────────────────────────────────────────────────────

def write_template(path: str):
    """Write an importable blank stock-count template.

    The instruction row above the header is exactly why `parse` scans for the
    header rather than trusting a row number — the same reason it copes with
    SAP's four junk rows.
    """
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = TEMPLATE_SHEET

    sheet.cell(row=1, column=1, value=TEMPLATE_NOTE)
    for column, header in enumerate(TEMPLATE_HEADERS, start=1):
        sheet.cell(row=3, column=column, value=header)
    # One example sets a threshold, the other leaves it empty — which is how the
    # file says "keep the threshold already in the app".
    for offset, example in enumerate((
        ('EXAMPLE-MAT-1', 'Example row — delete before importing', 'PCS', 0, 2),
        ('EXAMPLE-MAT-2', 'Example row — delete before importing', 'KG', 0, None),
    )):
        for column, value in enumerate(example, start=1):
            if value is None:
                continue
            sheet.cell(row=4 + offset, column=column, value=value)

    for column, width in enumerate((22, 54, 10, 10, 10), start=1):
        sheet.column_dimensions[sheet.cell(row=3, column=column).column_letter]\
            .width = width

    book.save(path)

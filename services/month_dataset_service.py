"""services/month_dataset_service.py — the month's SCADA files, kept with the month

The Monthly Report used to ask for 13 files (Tashkent) or 6 (Bukhara) every
time it was generated, from wherever they lay on the Desktop. Now a month owns
its data set:

  add_files      recognises each file's type from its sheets and column headers
                 (header rows only — a few hundredths of a second per file),
                 copies it into report_data/p<project>/<YYYY-MM>/inputs/ next to
                 the database, and remembers it with its SHA-256. A file whose
                 type cannot be told comes back 'unrecognised' / 'ambiguous'
                 for the user to choose; nothing is guessed. A newer file of a
                 type replaces the older one (the older copy stays: an earlier
                 version's manifest names it).
  analyse_file   reads the copy once (through parse_cache, so the report run
                 after it reads nothing from xlsx) and stores its coverage: days
                 with data, part days, gaps of more than an hour.
  generation_paths  the stored copies as the generator's *_path arguments —
                 a re-generation never picks files again.

The folder is derived from database.db_manager.DB_PATH at call time: next to
dist\\pv_bess_tracker.db for the exe (survives a rebuild; never inside
PyInstaller's _MEI temp folder), and inside the test's work folder under the
harness. A locked (sent) month's data set cannot be changed.
"""

import calendar
import datetime
import json
import os
import re
import shutil
from typing import Dict, List, Optional

import pandas as pd

from database.db_manager import get_connection
import services.parse_cache as pc
import services.report_workflow_service as rw

data_root = pc.data_root


def _t(key, label, site, group, required, param, manual=False):
    return {'key': key, 'label': label, 'site': site, 'group': group,
            'required': required, 'param': param, 'manual': manual}


# group: series (timestamped samples, shown in the coverage grid), events (the
# alarm list), daily (one row a day), month_end (snapshots and the customer's
# file — the "month end" checklist). `param` is the generator argument.
TYPES = [
    _t('working_status', 'LC working status', 'tashkent', 'series', True, 'working_status_path'),
    _t('pcs_cd', 'PCS charge/discharge status', 'tashkent', 'series', True, 'pcs_cd_path'),
    _t('pcs_fault', 'PCS fault status', 'tashkent', 'series', False, 'pcs_fault_path'),
    _t('soc', 'SOC', 'tashkent', 'series', True, 'soc_path'),
    _t('alarms', 'Alarms (Production + Warning)', 'tashkent', 'events', True, 'alarm_path'),
    _t('lc_charge', 'LC daily charge', 'tashkent', 'series', True, 'lc_charge_path'),
    _t('lc_discharge', 'LC daily discharge', 'tashkent', 'series', True, 'lc_discharge_path'),
    _t('hv_meter', 'HV meter daily', 'tashkent', 'daily', True, 'hv_meter_daily_path'),
    _t('cycles_first_day', 'Cycles first day', 'tashkent', 'month_end', False, 'cycles_first_day_path'),
    _t('cycles_last_day', 'Cycles last day', 'tashkent', 'month_end', False, 'cycles_last_day_path'),
    _t('soh_snapshot', 'SOH last day', 'tashkent', 'month_end', True, 'soh_snapshot_path'),
    _t('lc_total_charge', 'LC total charge', 'tashkent', 'month_end', False, 'lc_total_charge_path'),
    _t('lc_total_discharge', 'LC total discharge', 'tashkent', 'month_end', False, 'lc_total_discharge_path'),
    _t('bk_site_kpi', 'Main KPI 5-min', 'bukhara', 'series', True, 'site_kpi_path'),
    _t('bk_overall_lc', 'LC data total', 'bukhara', 'daily', True, 'overall_lc_path'),
    _t('bk_meter_daily', 'Main meter daily', 'bukhara', 'daily', True, 'meter_daily_path'),
    _t('bk_alarms', 'Monthly alarm report', 'bukhara', 'events', True, 'alarm_path'),
    _t('bk_battery_unit', 'Battery unit data (cycles)', 'bukhara', 'month_end', True, 'battery_unit_path'),
    # not read by the report; chosen by hand (its format is the customer's)
    _t('customer_poi', 'Customer POI file', '*', 'month_end', False, None, manual=True),
]
TYPE = {t['key']: t for t in TYPES}

# Exports recognised but not read by the report — reported as such, not copied.
IGNORED = {
    'bk_availability': 'System availability (not read by the report)',
    'bk_total_5min': 'Total 5-min data (not read by the report)',
    'bk_daily_kpi': 'Daily KPI / total data (not read by the report)',
    'bk_meter_5min': 'Main meter 5-min data (not read by the report)',
}
_IGNORED_SITE = 'bukhara'

EXCEL_EXT = ('.xlsx', '.xlsm', '.xls')


def site_type(project_id: int) -> str:
    return (rw.get_project_config(project_id).get('site_type') or 'tashkent')


def types_for(site: str) -> List[dict]:
    return [t for t in TYPES if t['site'] in (site, '*')]


def month_dir(project_id: int, year: int, month: int) -> str:
    return os.path.join(data_root(), f'p{int(project_id)}', f'{int(year):04d}-{int(month):02d}')


def rel(path: str) -> str:
    return os.path.relpath(path, data_root()).replace('\\', '/')


def absolute(stored: str) -> str:
    return os.path.join(data_root(), *str(stored).split('/'))


# ── recognition ──────────────────────────────────────────────────────────────

def probe_file(path: str) -> dict:
    """Sheet names, the first rows of the first sheets, and the first two
    timestamps — read in openpyxl's read-only mode, never the whole file."""
    import openpyxl
    out = {'sheets': [], 'rows': {}, 'first_ts': None, 'step_min': None, 'error': ''}
    if not str(path).lower().endswith(EXCEL_EXT) or str(path).lower().endswith('.xls'):
        out['error'] = 'not an .xlsx workbook'
        return out
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:                                  # noqa: BLE001
        out['error'] = f'cannot be opened: {e}'
        return out
    try:
        out['sheets'] = list(wb.sheetnames)
        wanted = [wb.sheetnames[0]] + [s for s in wb.sheetnames
                                       if s.lower() in ('production', 'warning', 'total')]
        for name in dict.fromkeys(wanted):
            rows = []
            for i, r in enumerate(wb[name].iter_rows(values_only=True)):
                rows.append(list(r or []))
                if i >= 3:
                    break
            out['rows'][name] = rows
    finally:
        wb.close()
    first = out['rows'].get(out['sheets'][0]) if out['sheets'] else None
    if first and len(first) >= 2:
        hdr = [str(v or '').strip() for v in first[0]]
        ts = [_row_ts(r, hdr) for r in first[1:]]
        ts = [t for t in ts if t is not None]
        if ts:
            out['first_ts'] = ts[0]
        if len(ts) >= 2 and ts[1] > ts[0]:
            out['step_min'] = (ts[1] - ts[0]).total_seconds() / 60.0
    return out


def _row_ts(row, hdr):
    if not row or not hdr or hdr[0].lower() != 'date':
        return None
    d = row[0]
    t = row[1] if len(hdr) > 1 and hdr[1].lower() == 'time' and len(row) > 1 else None
    try:
        if isinstance(d, datetime.datetime) and t is None:
            return pd.Timestamp(d)
        s = str(d.date() if isinstance(d, datetime.datetime) else d)
        if t is not None:
            s += ' ' + str(t)
        v = pd.to_datetime(s, errors='coerce')
        return None if pd.isna(v) else v
    except (TypeError, ValueError):
        return None


def _prev_month(year, month):
    return (year - 1, 12) if month == 1 else (year, month - 1)


def recognise(probe: dict, year: int = None, month: int = None) -> List[str]:
    """Candidate type keys for a probed file, over both sites (an empty list:
    not recognised; more than one: ask)."""
    sheets = [s.lower() for s in probe.get('sheets') or []]
    if not sheets:
        return []

    def header(name):
        rows = probe['rows'].get(name) or []
        return ' | '.join(str(v) for v in (rows[0] if rows else []) if v is not None).upper()

    first = probe['sheets'][0]
    H = header(first)
    step = probe.get('step_min')
    got = []
    if {'production', 'warning'} <= set(sheets):
        prod = next(s for s in probe['sheets'] if s.lower() == 'production')
        if 'TRIGGER NAME' in header(prod):
            got += ['alarms', 'bk_alarms']
    if 'lc_energy' in sheets:
        got.append('bk_overall_lc')
    if 'total' in sheets:
        tot = header(next(s for s in probe['sheets'] if s.lower() == 'total'))
        if 'BLOCK NUMBER' in tot and 'CYCLES' in tot:
            got.append('bk_battery_unit')
    lc = 'LC200' in H
    for needle, key in (('WORKING STATUS', 'working_status'), ('SYSTEM SOC', 'soc'),
                        ('SYSTEM SOH', 'soh_snapshot'),
                        ('DAILY CHARGE ENERGY', 'lc_charge'),
                        ('DAILY DISCHARGE ENERGY', 'lc_discharge'),
                        ('TOTAL CHARGE ENERGY', 'lc_total_charge'),
                        ('TOTAL DISCHARGE ENERGY', 'lc_total_discharge')):
        if lc and needle in H:
            got.append(key)
    if re.search(r'PCS\s+\d+\.\d+\.\d+', H):
        if 'CHARGE AND DISCHARGE STATUS' in H:
            got.append('pcs_cd')
        if 'FAULT STATUS' in H:
            got.append('pcs_fault')
    if 'CMU' in H and 'CHARGE AND DISCHARGE CYCLES' in H:
        got += _first_or_last(probe.get('first_ts'), year, month)
    if 'METER' in H and 'ENERGY EXPORTED' in H and 'ENERGY IMPORTED' in H:
        got.append('hv_meter')
    if 'UNIT FAULT' in H and first.lower().startswith('block'):
        got.append('bk_availability')
    if 'AVERAGE SOC' in H and 'NUMBER OF RACK' in H:
        if 'DAILY - ' in H or (step is not None and step >= 12 * 60):
            got.append('bk_daily_kpi')
        elif 'TOTAL ACTIVE POWER' in H:
            got.append('bk_total_5min')
        else:
            got.append('bk_site_kpi')
    if 'METER BESS' in H and 'AC CURRENT PHASE' in H:
        got.append('bk_meter_daily' if step is not None and step >= 12 * 60 else 'bk_meter_5min')
    return list(dict.fromkeys(got))


def _first_or_last(ts, year, month):
    """The two CMU cycle snapshots look alike; the date tells them apart."""
    if ts is None or not year or not month:
        return ['cycles_first_day', 'cycles_last_day']
    py, pm = _prev_month(year, month)
    if ((ts.year, ts.month) == (year, month) and ts.day <= 3) or \
            ((ts.year, ts.month) == (py, pm) and ts.day >= 28):
        return ['cycles_first_day']
    if (ts.year, ts.month) == (year, month) and ts.day >= 26:
        return ['cycles_last_day']
    return ['cycles_first_day', 'cycles_last_day']


def classify(path: str, site: str, year: int, month: int) -> dict:
    """{'status': 'recognised'|'ignored'|'unrecognised'|'ambiguous', 'type',
    'candidates', 'probe', 'note'} for one file and one project's site."""
    probe = probe_file(path)
    cands = recognise(probe, year, month)
    mine = [k for k in cands if k in TYPE and TYPE[k]['site'] in (site, '*')]
    ignored = [k for k in cands if k in IGNORED and _IGNORED_SITE == site]
    other = [k for k in cands if k not in mine and k not in ignored]
    if len(mine) == 1:
        return {'status': 'recognised', 'type': mine[0], 'candidates': mine, 'probe': probe, 'note': ''}
    if len(mine) > 1:
        return {'status': 'ambiguous', 'type': None, 'candidates': mine, 'probe': probe,
                'note': 'could be ' + ' or '.join(TYPE[k]['label'] for k in mine)}
    if ignored:
        return {'status': 'ignored', 'type': ignored[0], 'candidates': [], 'probe': probe,
                'note': IGNORED[ignored[0]]}
    note = probe.get('error') or ''
    if other:
        note = 'looks like an export of another site type'
    return {'status': 'unrecognised', 'type': None, 'candidates': [], 'probe': probe, 'note': note}


# ── the data set ─────────────────────────────────────────────────────────────

def list_files(project_id: int, year: int, month: int, status: Optional[str] = 'active') -> List[dict]:
    q = "SELECT * FROM month_files WHERE project_id=? AND year=? AND month=?"
    p = [project_id, year, month]
    if status:
        q += " AND status=?"
        p.append(status)
    q += " ORDER BY id"
    conn = get_connection()
    try:
        rows = [dict(r) for r in conn.execute(q, p).fetchall()]
    finally:
        conn.close()
    for r in rows:
        t = TYPE.get(r['data_type']) or {}
        r['label'] = t.get('label', r['data_type'])
        r['group'] = t.get('group', '')
        r['abs_path'] = absolute(r['stored_path'])
        try:
            r['coverage'] = json.loads(r.get('coverage') or '{}')
        except ValueError:
            r['coverage'] = {}
    return rows


def get_file(file_id: int) -> Optional[dict]:
    conn = get_connection()
    try:
        r = conn.execute("SELECT project_id, year, month FROM month_files WHERE id=?",
                         (int(file_id),)).fetchone()
    finally:
        conn.close()
    if not r:
        return None
    return next((f for f in list_files(r[0], r[1], r[2], status=None) if f['id'] == int(file_id)), None)


def active_by_type(project_id: int, year: int, month: int) -> Dict[str, dict]:
    return {f['data_type']: f for f in list_files(project_id, year, month)}


def _safe_name(name: str) -> str:
    return re.sub(r'[^\w.\- ()]+', '_', name).strip() or 'file'


def store_file(project_id: int, year: int, month: int, path: str, data_type: str,
               recognised: str = 'auto') -> dict:
    """Copy one file into the month as `data_type`. Returns {'status':
    'added'|'replaced'|'duplicate'|'restored', 'file_id', 'replaced': name|None}."""
    rw.assert_month_open(project_id, year, month)
    if data_type not in TYPE:
        raise ValueError(f'unknown data type {data_type!r}')
    sha = pc.file_sha256(path)
    current = active_by_type(project_id, year, month).get(data_type)
    if current and current['sha256'] == sha:
        stored = current['abs_path']
        if os.path.isfile(stored) and pc.file_sha256(stored) == sha:
            return {'status': 'duplicate', 'file_id': current['id'], 'replaced': None}
        # the stored copy was lost or changed on disk: the same file puts it back
        os.makedirs(os.path.dirname(stored), exist_ok=True)
        if os.path.exists(stored):
            os.chmod(stored, 0o666)
        tmp = stored + '.part'
        shutil.copyfile(path, tmp)
        os.replace(tmp, stored)
        return {'status': 'restored', 'file_id': current['id'], 'replaced': None}
    name = os.path.basename(path)
    dest_dir = os.path.join(month_dir(project_id, year, month), 'inputs')
    os.makedirs(dest_dir, exist_ok=True)
    dest = os.path.join(dest_dir, f'{data_type}__{sha[:10]}__{_safe_name(name)}')
    if not (os.path.exists(dest) and pc.file_sha256(dest) == sha):
        tmp = dest + '.part'
        shutil.copyfile(path, tmp)
        os.replace(tmp, dest)
    conn = get_connection()
    try:
        if current:
            conn.execute("UPDATE month_files SET status='replaced' WHERE id=?", (current['id'],))
        cur = conn.execute("""
            INSERT INTO month_files (project_id, year, month, data_type, original_name,
                                     source_path, stored_path, sha256, size_bytes, recognised)
            VALUES (?,?,?,?,?,?,?,?,?,?)""",
            (project_id, year, month, data_type, name, os.path.abspath(path), rel(dest), sha,
             os.path.getsize(dest), recognised))
        conn.commit()
        fid = cur.lastrowid
    finally:
        conn.close()
    return {'status': 'replaced' if current else 'added', 'file_id': fid,
            'replaced': current['original_name'] if current else None}


def add_files(project_id: int, year: int, month: int, paths, chosen: Optional[dict] = None) -> List[dict]:
    """Add files to the month's data set. `chosen` = {path: type key | 'skip'}
    answers files that were not recognised. One result per path:
    {'path', 'name', 'status', 'type', 'label', 'candidates', 'note', 'period',
    'file_id'} — status added | replaced | duplicate | restored | ignored | skipped |
    unrecognised | ambiguous | error."""
    rw.assert_month_open(project_id, year, month)
    site = site_type(project_id)
    chosen = chosen or {}
    out = []
    for path in paths:
        r = {'path': path, 'name': os.path.basename(path), 'status': '', 'type': None,
             'label': '', 'candidates': [], 'note': '', 'period': '', 'file_id': None}
        out.append(r)
        pick = chosen.get(path)
        if pick == 'skip':
            r['status'] = 'skipped'
            continue
        try:
            if pick:
                if pick not in TYPE:
                    raise ValueError(f'unknown data type {pick!r}')
                key, how = pick, 'chosen'
                probe = probe_file(path)
            else:
                c = classify(path, site, year, month)
                probe = c['probe']
                r.update(candidates=c['candidates'], note=c['note'])
                if c['status'] in ('unrecognised', 'ambiguous'):
                    r['status'] = c['status']
                    continue
                if c['status'] == 'ignored':
                    r.update(status='ignored', label=c['note'])
                    continue
                key, how = c['type'], 'auto'
            ts = probe.get('first_ts')
            if ts is not None:
                r['period'] = f'from {ts:%d.%m.%Y %H:%M}'
                if (ts.year, ts.month) not in ((year, month), _prev_month(year, month)):
                    r['note'] = (r['note'] + '; ' if r['note'] else '') + \
                        f'data starts {ts:%d.%m.%Y} — not {rw.MONTHS_EN[month]} {year}'
            res = store_file(project_id, year, month, path, key, how)
            r.update(status=res['status'], type=key, label=TYPE[key]['label'],
                     file_id=res['file_id'])
            if res['replaced']:
                r['note'] = (r['note'] + '; ' if r['note'] else '') + f"replaces {res['replaced']}"
        except rw.MonthLockedError:
            raise
        except Exception as e:                              # noqa: BLE001
            r.update(status='error', note=str(e))
    return out


def folder_files(folder: str) -> List[str]:
    """The Excel files directly in a folder (Office lock files left out)."""
    try:
        names = sorted(os.listdir(folder))
    except OSError:
        return []
    return [os.path.join(folder, n) for n in names
            if n.lower().endswith(EXCEL_EXT) and not n.startswith('~$')
            and os.path.isfile(os.path.join(folder, n))]


def remove_file(file_id: int):
    """Take a file out of the month (its copy stays for earlier manifests)."""
    f = get_file(file_id)
    if not f:
        return
    rw.assert_month_open(f['project_id'], f['year'], f['month'])
    conn = get_connection()
    try:
        conn.execute("UPDATE month_files SET status='removed' WHERE id=?", (int(file_id),))
        conn.commit()
    finally:
        conn.close()


# ── coverage ─────────────────────────────────────────────────────────────────

def _timestamps(df) -> pd.Series:
    if df is None or getattr(df, 'empty', True) or 'Date' not in df.columns:
        return pd.Series(dtype='datetime64[ns]')
    if 'Time' in df.columns:
        ts = pd.to_datetime(df['Date'].astype(str) + ' ' + df['Time'].astype(str), errors='coerce')
    else:
        ts = pd.to_datetime(df['Date'], errors='coerce')
    values = df.drop(columns=[c for c in ('Date', 'Time') if c in df.columns])
    if values.shape[1]:
        ts = ts[values.notna().any(axis=1).values]
    return ts.dropna()


def series_coverage(ts: pd.Series) -> dict:
    """Days with data ('full' / 'partial') and gaps > max(1 h, 3 samples)."""
    ts = pd.Series(pd.to_datetime(ts)).dropna().drop_duplicates().sort_values()
    if ts.empty:
        return {'rows': 0, 'days': {}, 'gaps': []}
    diffs = ts.diff().dropna().dt.total_seconds() / 60.0
    step = float(diffs[diffs > 0].median()) if (diffs > 0).any() else 1440.0
    expected = max(1, int(round(1440.0 / step))) if step < 1440 else 1
    per_day = ts.dt.date.value_counts()
    days = {d.isoformat(): ('full' if n >= 0.95 * expected else 'partial')
            for d, n in sorted(per_day.items())}
    limit = max(60.0, 3 * step)
    gaps = []
    prev = ts.iloc[0]
    for t in ts.iloc[1:]:
        dm = (t - prev).total_seconds() / 60.0
        if dm > limit and step < 1440:
            a = prev + pd.Timedelta(minutes=step)
            gaps.append([f'{a:%Y-%m-%d %H:%M}', f'{t:%Y-%m-%d %H:%M}',
                         round((t - a).total_seconds() / 3600.0, 2)])
        prev = t
    return {'rows': int(len(ts)), 'first': f'{ts.iloc[0]:%Y-%m-%d %H:%M}',
            'last': f'{ts.iloc[-1]:%Y-%m-%d %H:%M}', 'step_min': round(step, 2),
            'days': days, 'gaps': gaps[:50], 'n_gaps': len(gaps)}


def analyse_file(file_id: int) -> dict:
    """Read a stored file once (warming the parse cache with exactly the
    reads the report loaders make) and store its coverage."""
    f = get_file(file_id)
    if not f:
        return {}
    path, key, group = f['abs_path'], f['data_type'], f['group']
    cov = {}
    try:
        if key == 'customer_poi':
            cov = {'present': True}
        elif group == 'events':
            ts = []
            for sheet in ('Production', 'Warning'):
                try:
                    df = pc.read_excel(path, sheet_name=sheet)
                except Exception:                          # noqa: BLE001
                    continue
                if 'Activated' in df.columns:
                    ts.append(pd.to_datetime(df['Activated'], errors='coerce').dropna())
            allts = pd.concat(ts) if ts else pd.Series(dtype='datetime64[ns]')
            if allts.empty:
                cov = {'rows': 0, 'days': {}, 'gaps': []}
            else:
                d0, d1 = allts.min().date(), allts.max().date()
                days = {(d0 + datetime.timedelta(days=i)).isoformat(): 'full'
                        for i in range((d1 - d0).days + 1)}
                cov = {'rows': int(len(allts)), 'first': f'{allts.min():%Y-%m-%d %H:%M}',
                       'last': f'{allts.max():%Y-%m-%d %H:%M}', 'days': days, 'gaps': []}
        elif key == 'bk_battery_unit':
            df = pc.read_excel(path, sheet_name='Total')
            cov = {'rows': int(len(df))}
        elif key == 'bk_overall_lc':
            cov = series_coverage(_timestamps(pc.read_excel(path, sheet_name='LC_Energy')))
        else:
            cov = series_coverage(_timestamps(pc.read_excel(path, sheet_name=0)))
    except Exception as e:                                  # noqa: BLE001
        cov = {'error': str(e)}
    cov['analysed_at'] = datetime.datetime.now().isoformat(timespec='seconds')
    conn = get_connection()
    try:
        conn.execute("UPDATE month_files SET coverage=? WHERE id=?",
                     (json.dumps(cov, default=str), int(file_id)))
        conn.commit()
    finally:
        conn.close()
    return cov


def analyse_pending(project_id: int, year: int, month: int, progress=None) -> int:
    """Analyse every active file of the month that has no coverage yet."""
    n = 0
    for f in list_files(project_id, year, month):
        if f['coverage'].get('analysed_at'):
            continue
        if progress:
            progress(f"Reading {f['original_name']}…")
        analyse_file(f['id'])
        n += 1
    return n


def expected_days(year: int, month: int, today: datetime.date = None) -> List[str]:
    """Days of the month data can exist for: the whole month once it is over,
    up to yesterday while it runs."""
    today = today or datetime.date.today()
    last = calendar.monthrange(year, month)[1]
    end = datetime.date(year, month, last)
    if (year, month) == (today.year, today.month):
        end = today - datetime.timedelta(days=1)
    start = datetime.date(year, month, 1)
    return [(start + datetime.timedelta(days=i)).isoformat()
            for i in range(max(0, (end - start).days + 1))]


def coverage_grid(project_id: int, year: int, month: int) -> List[dict]:
    """One row per grid type of the site: {'type', 'label', 'required', 'file',
    'days': {iso: 'full'|'partial'}, 'analysed'}."""
    files = active_by_type(project_id, year, month)
    rows = []
    for t in types_for(site_type(project_id)):
        if t['group'] not in ('series', 'events', 'daily'):
            continue
        f = files.get(t['key'])
        cov = (f or {}).get('coverage') or {}
        prefix = f'{year:04d}-{month:02d}-'
        rows.append({'type': t['key'], 'label': t['label'], 'required': t['required'],
                     'file': f, 'analysed': bool(cov.get('analysed_at')),
                     'days': {d: v for d, v in (cov.get('days') or {}).items() if d.startswith(prefix)},
                     'gaps': cov.get('gaps') or []})
    return rows


def month_end(project_id: int, year: int, month: int) -> List[dict]:
    """The month-end checklist: snapshots and the customer's file, present or not."""
    files = active_by_type(project_id, year, month)
    out = []
    for t in types_for(site_type(project_id)):
        if t['group'] != 'month_end':
            continue
        f = files.get(t['key'])
        cov = (f or {}).get('coverage') or {}
        out.append({'type': t['key'], 'label': t['label'], 'required': t['required'],
                    'present': bool(f), 'file': f, 'first': cov.get('first')})
    return out


def missing_required(project_id: int, year: int, month: int) -> List[str]:
    files = active_by_type(project_id, year, month)
    return [t['label'] for t in types_for(site_type(project_id))
            if t['required'] and t['key'] not in files]


def generation_paths(project_id: int, year: int, month: int) -> dict:
    """{generator argument: stored copy} for the month's active files."""
    out = {}
    for key, f in active_by_type(project_id, year, month).items():
        t = TYPE.get(key)
        if not t or not t['param']:
            continue
        if not os.path.isfile(f['abs_path']):
            raise FileNotFoundError(f"{t['label']}: the stored copy is missing ({f['stored_path']}) "
                                    "— add the file again")
        out[t['param']] = f['abs_path']
    return out

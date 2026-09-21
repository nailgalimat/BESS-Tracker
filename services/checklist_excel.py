"""services/checklist_excel.py — the customer's own checklist workbook, read
and written back.

The PM checklists come as Excel files the customer issued ("01) PCS
Checklist.xlsx", "02) BESS Checklist.xlsx"). They must go back looking exactly
as they came: same logo, same layout, same print setup — and, above all, with
the **Excel 365 cell checkboxes** in column E still being checkboxes. Those
live in `xl/featurePropertyBag/featurePropertyBag.xml`, which openpyxl drops on
save (it also drops the printer settings), so a load-and-save round trip hands
the customer a different-looking document.

So: read with openpyxl (safe, read-only), write by **patching the sheet XML
inside the zip** and copying every other part byte for byte.

Layout of both files (checked against the real ones):

      rows 1-3   title
      row  4-6   header fields, merged: C4 Plant Name, E4 Date, C5 Location,
                 E5 Equipment Serial Number, C6 Name, E6 Signature
      row  7     column headings
      rows 8..   items: A S.No · B Equipment · C Activity (merged per group)
                 D Description · E checkbox · F "Verified / Done" · G Comments
      Progress   the row whose B says "Progress": COUNTIF(E…,TRUE)/COUNTA(E…)
      then       footnotes

Mapping, as the user set it:
    OK   -> TRUE       NOK -> FALSE + comment       N/A -> empty
    excluded -> empty, the row kept, with whatever note the user wrote
Progress is OK / (OK + NOK), so every in-scope row must be written — a value
left over from the template would otherwise count.
"""
import io
import os
import re
import shutil
import zipfile
from xml.sax.saxutils import escape

import openpyxl

HEADER_ANCHORS = {'plant': 'C4', 'date': 'E4', 'location': 'C5',
                  'serial': 'E5', 'name': 'C6', 'signature': 'E6'}
FIRST_ITEM_ROW = 8
COL = {'no': 'A', 'equipment': 'B', 'activity': 'C', 'text': 'D',
       'status': 'E', 'done': 'F', 'comment': 'G'}

# result codes stored per item
OK, NOK, NA, EXCLUDED, PENDING = 'OK', 'NOK', 'N/A', 'Excluded', ''


def parse_template(path: str) -> dict:
    """Read a checklist workbook: its items, in order, with the Excel row each
    one sits on. Returns {'sheet', 'progress_row', 'items': [...]}, where an
    item is {no, equipment, activity, text, excel_row}."""
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb.active
    merged = {}
    for rng in ws.merged_cells.ranges:
        top = ws.cell(row=rng.min_row, column=rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged[(r, c)] = top

    def val(row, col_letter):
        c = ws[f'{col_letter}{row}']
        v = merged.get((c.row, c.column), c.value)
        return '' if v is None else str(v).strip()

    items, progress_row = [], None
    for row in range(FIRST_ITEM_ROW, ws.max_row + 1):
        equipment = val(row, COL['equipment'])
        if equipment.lower() == 'progress':
            progress_row = row
            break
        text = val(row, COL['text'])
        if not text:
            continue
        items.append({
            'no': val(row, COL['no']),
            'equipment': equipment,
            'activity': re.sub(r'\s*\n\s*', ' ', val(row, COL['activity'])),
            'text': text,
            'excel_row': row,
        })
    wb.close()
    return {'sheet': ws.title, 'progress_row': progress_row, 'items': items}


# ── writing the workbook back ────────────────────────────────────────────────

_CELL = re.compile(r'(\$?)([A-Z]{1,3})(\$?)(\d+)')


def _shift_row(r, at, n):
    return r + n if r >= at else r


def _shift_refs(text, at, n):
    return _CELL.sub(lambda m: f'{m.group(1)}{m.group(2)}{m.group(3)}'
                               f'{_shift_row(int(m.group(4)), at, n)}', text)


def _shift_range(rng, at, n, extend_if_ends_before=False):
    """A range spanning the insertion point grows; one wholly below moves.
    With extend_if_ends_before, the group being extended grows too."""
    if ':' not in rng:
        return _shift_refs(rng, at, n)
    a, b = rng.split(':')
    ra, rb = int(_CELL.match(a).group(4)), int(_CELL.match(b).group(4))
    if extend_if_ends_before and rb == at - 1 and ra < at:
        return f'{a}:{_shift_refs(b, rb, n)}'
    return f'{_shift_refs(a, at, n)}:{_shift_refs(b, at, n)}'


_REF = re.compile(r'(\$?[A-Z]{1,3}\$?\d+)(?::(\$?[A-Z]{1,3}\$?\d+))?')


def _shift_formula(text, at, n):
    """Shift a formula's references. A range that ends on the row just above
    the new one grows over it — an item added at the end of the list has to
    count in Progress, and has to stay inside the shared formula's range, or
    Excel calls the file damaged."""
    def one(m):
        if m.group(2):
            return _shift_range(f'{m.group(1)}:{m.group(2)}', at, n,
                                extend_if_ends_before=True)
        return _shift_refs(m.group(1), at, n)
    return _REF.sub(one, text)


def _shift_attr_ref(attrs, at, n):
    return re.sub(r'ref="([^"]+)"',
                  lambda m: f'ref="{_shift_range(m.group(1), at, n, True)}"', attrs)


# Excel's hard limit on the characters in one cell. A longer string makes
# Excel call the whole file damaged and offer to repair it, which for the
# customer means the checklist did not arrive.
_MAX_CELL_CHARS = 32767


def _cell_xml(ref: str, style: str, value):
    """One <c> element, keeping the template cell's style (the checkbox format
    in column E is carried by the style, not by the value)."""
    if value is True or value is False:
        return f'<c r="{ref}"{style} t="b"><v>{int(value)}</v></c>'
    if value in (None, ''):
        return f'<c r="{ref}"{style}/>'
    text = str(value)
    if len(text) > _MAX_CELL_CHARS:
        text = text[:_MAX_CELL_CHARS - 1] + '…'
    return (f'<c r="{ref}"{style} t="inlineStr"><is>'
            f'<t xml:space="preserve">{escape(text)}</t></is></c>')


def _set_cell(sheet_xml: str, ref: str, value) -> str:
    """Replace one cell's value, keeping its style; append it to its row when
    the template has no cell there at all."""
    m = re.search(r'<c r="%s"([^>]*?)(/>|>.*?</c>)' % ref, sheet_xml, flags=re.S)
    style = ''
    if m:
        s = re.search(r'\ss="\d+"', m.group(1))
        style = s.group(0) if s else ''
        return sheet_xml.replace(m.group(0), _cell_xml(ref, style, value), 1)
    row = re.match(r'([A-Z]+)(\d+)', ref).group(2)
    rm = re.search(r'(<row r="%s"[^>]*>)(.*?)(</row>)' % row, sheet_xml, flags=re.S)
    if not rm:
        return sheet_xml
    return sheet_xml.replace(rm.group(0),
                             rm.group(1) + rm.group(2) + _cell_xml(ref, '', value)
                             + rm.group(3), 1)


def _insert_row(parts: dict, sheet_name: str, after_row: int, cells: dict,
                group_col: str = 'C', height=None) -> int:
    """Insert one row straight after *after_row*, in the sheet XML, keeping
    checkboxes, merges, the group label, formulas, comments and print setup.
    Returns the new row number."""
    at, n = after_row + 1, 1
    s = parts[sheet_name].decode('utf-8')

    def move_row(m):                                   # rows below move down
        r = int(m.group(1))
        row_xml = m.group(0)
        if r < at:
            return row_xml
        row_xml = re.sub(r'<row r="\d+"', f'<row r="{r + n}"', row_xml, count=1)
        return re.sub(r'<c r="([A-Z]+)(\d+)"',
                      lambda c: f'<c r="{c.group(1)}{int(c.group(2)) + n}"', row_xml)
    s = re.sub(r'<row r="(\d+)"[^>]*>.*?</row>', move_row, s, flags=re.S)

    s = re.sub(r'<f([^>]*)>([^<]*)</f>',
               lambda m: f'<f{_shift_attr_ref(m.group(1), at, n)}>'
                         f'{_shift_formula(m.group(2), at, n)}</f>', s)
    s = re.sub(r'<f([^>]*ref="[^"]*"[^>]*)/>',
               lambda m: f'<f{_shift_attr_ref(m.group(1), at, n)}/>', s)
    s = re.sub(r'<mergeCell ref="([^"]+)"/>',
               lambda m: '<mergeCell ref="%s"/>' % _shift_range(
                   m.group(1), at, n,
                   extend_if_ends_before=m.group(1).startswith(group_col)), s)
    s = re.sub(r'sqref="([^"]+)"',
               lambda m: 'sqref="' + ' '.join(_shift_range(x, at, n, True)
                                              for x in m.group(1).split()) + '"', s)
    s = re.sub(r'<dimension ref="([^"]+)"/>',
               lambda m: f'<dimension ref="{_shift_range(m.group(1), at, n)}"/>', s)
    s = re.sub(r'<brk id="(\d+)"',
               lambda m: f'<brk id="{_shift_row(int(m.group(1)), at, n)}"', s)

    tmpl = re.search(r'<row r="%d"([^>]*)>(.*?)</row>' % after_row, s, flags=re.S)
    row_attrs = re.sub(r'\s+ht="[^"]*"', '', tmpl.group(1))
    if height:
        row_attrs += f' ht="{height}" customHeight="1"'
    new_cells = []
    for c in re.finditer(r'<c r="([A-Z]+)%d"([^>]*?)(/>|>(.*?)</c>)' % after_row,
                         tmpl.group(2), flags=re.S):
        col = c.group(1)
        st = re.search(r'\ss="\d+"', c.group(2))
        st = st.group(0) if st else ''
        body = c.group(4) or ''
        if col == COL['done'] and '<f' in body:        # keep "Verified / Done"
            si = re.search(r'si="(\d+)"', body)
            new_cells.append(
                f'<c r="{col}{at}"{st} t="str"><f t="shared" si="{si.group(1)}"/><v/></c>'
                if si else f'<c r="{col}{at}"{st}/>')
        else:
            new_cells.append(_cell_xml(f'{col}{at}', st, cells.get(col)))
    s = s.replace(tmpl.group(0),
                  tmpl.group(0) + f'<row r="{at}"{row_attrs}>{"".join(new_cells)}</row>', 1)
    parts[sheet_name] = s.encode('utf-8')

    w = parts['xl/workbook.xml'].decode('utf-8')
    w = re.sub(r'(<definedName[^>]*>)([^<]*)(</definedName>)',
               lambda m: m.group(1) + _shift_refs(m.group(2), at, n) + m.group(3), w)
    parts['xl/workbook.xml'] = w.encode('utf-8')

    for name in list(parts):                           # comments and their anchors
        if name.startswith('xl/comments'):
            parts[name] = re.sub(
                r'ref="([^"]+)"', lambda m: f'ref="{_shift_refs(m.group(1), at, n)}"',
                parts[name].decode('utf-8')).encode('utf-8')
        if name.startswith('xl/drawings/vmlDrawing'):
            v = parts[name].decode('utf-8')
            v = re.sub(r'<x:Row>(\d+)</x:Row>',
                       lambda m: f'<x:Row>{_shift_row(int(m.group(1)), at - 1, n)}</x:Row>', v)

            def anchor(m):
                p = [x.strip() for x in m.group(1).split(',')]
                p[2] = str(_shift_row(int(p[2]), at - 1, n))
                p[6] = str(_shift_row(int(p[6]), at - 1, n))
                return '<x:Anchor>' + ', '.join(p) + '</x:Anchor>'
            parts[name] = re.sub(r'<x:Anchor>([^<]*)</x:Anchor>', anchor, v,
                                 flags=re.S).encode('utf-8')
    return at


def _drop_calc_chain(parts: dict):
    """calcChain lists formula cells by address; after a row insert it is
    stale. Excel rebuilds it."""
    if 'xl/calcChain.xml' in parts:
        del parts['xl/calcChain.xml']
        parts['[Content_Types].xml'] = re.sub(
            r'<Override[^>]*calcChain[^>]*/>', '',
            parts['[Content_Types].xml'].decode('utf-8')).encode('utf-8')
        rels = 'xl/_rels/workbook.xml.rels'
        parts[rels] = re.sub(r'<Relationship[^>]*calcChain[^>]*/>', '',
                             parts[rels].decode('utf-8')).encode('utf-8')


def _force_recalc(parts: dict):
    w = parts['xl/workbook.xml'].decode('utf-8')
    if '<calcPr' in w:
        w = re.sub(r'<calcPr([^>]*?)\s*/>',
                   lambda m: '<calcPr' + re.sub(r'\s+fullCalcOnLoad="[^"]*"', '',
                                                m.group(1)) + ' fullCalcOnLoad="1"/>', w)
    else:
        w = w.replace('</workbook>', '<calcPr fullCalcOnLoad="1"/></workbook>')
    parts['xl/workbook.xml'] = w.encode('utf-8')


def write_filled(src_path: str, out_path: str, header: dict, results: dict,
                 added: list = None) -> str:
    """Write the filled checklist as a copy of the customer's own file.

    header  {plant, date, location, serial, name, signature} — blanks skipped
    results {excel_row: {'result': OK|NOK|N/A|Excluded|'', 'comment': str}}
    added   [{'after_row': int, 'no': str, 'equipment': str, 'text': str,
              'result': …, 'comment': str}] — items the desktop added, each
            placed at the end of its group
    """
    with zipfile.ZipFile(src_path) as zin:
        order = list(zin.infolist())
        parts = {i.filename: zin.read(i.filename) for i in order}
    sheet = 'xl/worksheets/sheet1.xml'
    if sheet not in parts:
        sheet = next(n for n in parts if n.startswith('xl/worksheets/sheet'))

    # added rows first: they move the rows below them, so the results written
    # afterwards land on the right ones. Two items added to the same group
    # share an after_row, and each insert goes immediately after it — so they
    # have to be written last-first, or the customer's file shows them in
    # reverse order. Sorting on after_row alone left that to chance.
    shifted = {}
    for _n, extra in sorted(enumerate(added or []),
                            key=lambda t: (t[1]['after_row'], t[0]), reverse=True):
        at = _insert_row(parts, sheet, extra['after_row'], {
            COL['no']: extra.get('no', ''),
            COL['equipment']: extra.get('equipment', ''),
            COL['text']: extra.get('text', ''),
            COL['status']: _status_value(extra.get('result')),
            COL['comment']: extra.get('comment', ''),
        }, height=30)
        for row in list(results):
            if row >= at:
                shifted[row] = shifted.get(row, row) + 1
    if added:
        _drop_calc_chain(parts)

    s = parts[sheet].decode('utf-8')
    for field, ref in HEADER_ANCHORS.items():
        if header.get(field):
            s = _set_cell(s, ref, header[field])
    for row, res in sorted(results.items()):
        r = shifted.get(row, row)
        s = _set_cell(s, f"{COL['status']}{r}", _status_value(res.get('result')))
        if res.get('comment'):
            s = _set_cell(s, f"{COL['comment']}{r}", res['comment'])
    parts[sheet] = s.encode('utf-8')
    _force_recalc(parts)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    with zipfile.ZipFile(out_path, 'w', zipfile.ZIP_DEFLATED) as zout:
        for info in order:
            if info.filename in parts:
                zout.writestr(info, parts[info.filename])
    return out_path


def _status_value(result):
    """OK ticks the box, NOK unticks it, N/A and excluded leave it empty —
    an empty cell is out of the Progress count, a FALSE one is not."""
    if result == OK:
        return True
    if result == NOK:
        return False
    return ''

"""
services/worklog_report_service.py
-----------------------------------
Monthly "Technical Works" report from the Field Log (work_log_entries).

Produces a PDF (ReportLab) or Word (python-docx via docx_renderer) document:
  1. Cover / title
  2. Summary   — total works, breakdown by category and by status,
                 counts of faults / spare-part usage
  3. Works by block/equipment — one table per block (site-wide items grouped
     separately), sorted by date
  4. Photo appendix (optional) — thumbnails grouped per entry

Both renderers share build_report_data() so the two formats never drift.
"""
from __future__ import annotations

import calendar
import os
import re
from datetime import datetime
from typing import Optional

from services.worklog_entry_service import (
    get_worklog_entries, CATEGORY_LABELS,
)
from services.image_service import get_images_for_log, download_remote_image


# ── Text / font helpers ───────────────────────────────────────────────────────

# No common PDF/DOCX font ships emoji glyphs, so they render as "tofu" boxes.
# Strip them from any label that may carry one (category labels start with 🔧/⚡…).
_EMOJI_RE = re.compile(
    "["
    "\U0001F000-\U0001FAFF"   # symbols, pictographs, emoji
    "\U00002600-\U000027BF"   # misc symbols + dingbats
    "\U00002B00-\U00002BFF"   # arrows/symbols
    "\U00002190-\U000021FF"   # arrows
    "️‍"            # variation selector, ZWJ
    "]+",
    flags=re.UNICODE,
)


def _clean(s) -> str:
    """Drop emoji and collapse whitespace so text renders in a plain TTF font."""
    if not s:
        return ""
    return re.sub(r"\s{2,}", " ", _EMOJI_RE.sub("", str(s))).strip()


def _cat_label(cat: str) -> str:
    return _clean(CATEGORY_LABELS.get(cat, cat))


# ReportLab's built-in fonts (Helvetica/Times) cover only Latin-1 — Cyrillic
# comes out as black boxes. Register a Unicode TTF once and thread its name
# through every style. DejaVuSans is shipped with matplotlib (a hard dependency,
# bundled into the frozen exe), so it is available on every machine.
_FONT_CACHE: dict = {}


def _register_cyrillic_fonts() -> tuple:
    """Return (regular, bold) font names registered with ReportLab (Cyrillic-safe)."""
    if _FONT_CACHE:
        return _FONT_CACHE["reg"], _FONT_CACHE["bold"]

    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    reg_name, bold_name = "Helvetica", "Helvetica-Bold"
    reg_path = bold_path = None

    # 1) matplotlib's bundled DejaVuSans (works frozen + on any OS)
    try:
        from matplotlib import font_manager
        reg_path = font_manager.findfont("DejaVu Sans", fallback_to_default=True)
        bold_path = font_manager.findfont(
            font_manager.FontProperties(family="DejaVu Sans", weight="bold"),
            fallback_to_default=True,
        )
    except Exception:
        reg_path = bold_path = None

    # 2) Windows Arial fallback
    if not reg_path or not os.path.isfile(reg_path):
        win = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")
        cand = os.path.join(win, "arial.ttf")
        candb = os.path.join(win, "arialbd.ttf")
        if os.path.isfile(cand):
            reg_path = cand
            bold_path = candb if os.path.isfile(candb) else cand

    try:
        if reg_path and os.path.isfile(reg_path):
            pdfmetrics.registerFont(TTFont("ReportBody", reg_path))
            reg_name = "ReportBody"
            if bold_path and os.path.isfile(bold_path):
                pdfmetrics.registerFont(TTFont("ReportBody-Bold", bold_path))
                bold_name = "ReportBody-Bold"
            else:
                bold_name = "ReportBody"
    except Exception:
        reg_name, bold_name = "Helvetica", "Helvetica-Bold"

    _FONT_CACHE["reg"], _FONT_CACHE["bold"] = reg_name, bold_name
    return reg_name, bold_name


MONTHS_EN = {
    1: "January", 2: "February", 3: "March", 4: "April",
    5: "May", 6: "June", 7: "July", 8: "August",
    9: "September", 10: "October", 11: "November", 12: "December",
}

STATUS_LABELS = {"open": "Open", "done": "Done", "": "—"}


# ── Data preparation ──────────────────────────────────────────────────────────

def _month_bounds(year: int, month: int):
    last = calendar.monthrange(year, month)[1]
    return f"{year}-{month:02d}-01", f"{year}-{month:02d}-{last:02d}"


def _block_key(e: dict):
    """Sort/group key: (has_block, zone, block) so site-wide items sort last."""
    z = e.get("zone_number")
    b = e.get("block_number")
    if b is None:
        return (1, 9999, 9999)
    return (0, z if z is not None else 0, b)


def _block_title(e: dict) -> str:
    z = e.get("zone_number")
    b = e.get("block_number")
    if b is None:
        return "Site-wide"
    t = f"Block {b}"
    if z is not None:
        t += f" (Zone {z})"
    return t


def _entry_location(e: dict) -> str:
    if e.get("block_number") is not None:
        z = e.get("zone_number", "")
        b = e.get("block_number", "")
        ci = e.get("container_index", "")
        loc = f"Z{z}·B{b}·C{ci}"
        if e.get("container_type"):
            loc += f" ({e['container_type']})"
        return loc
    return e.get("site_location") or "—"


def build_report_data(project_id: Optional[int], year: int, month: int,
                      project_name: str = "") -> dict:
    """Fetch entries for the month and pre-aggregate everything the renderers need."""
    date_from, date_to = _month_bounds(year, month)
    entries = get_worklog_entries(
        project_id=project_id, date_from=date_from, date_to=date_to
    )
    # get_worklog_entries returns newest-first; report reads best oldest-first
    entries = list(reversed(entries))

    # Summary aggregates
    by_category: dict = {}
    by_status = {"open": 0, "done": 0, "": 0}
    n_faults = 0
    n_parts = 0
    for e in entries:
        cat = e.get("category", "other")
        by_category[cat] = by_category.get(cat, 0) + 1
        st = (e.get("status") or "").lower()
        by_status[st if st in by_status else ""] = by_status.get(st, 0) + 1
        if (e.get("fault_name") or "").strip():
            n_faults += 1
        if (e.get("spare_parts") or "").strip():
            n_parts += 1

    # Group by block
    groups: dict = {}
    for e in sorted(entries, key=_block_key):
        key = _block_key(e)
        groups.setdefault(key, {"title": _block_title(e), "entries": []})
        groups[key]["entries"].append(e)

    return {
        "project_name": project_name or "(all projects)",
        "year": year,
        "month": month,
        "month_name": MONTHS_EN.get(month, str(month)),
        "period": f"{date_from} — {date_to}",
        "date_from": date_from,
        "date_to": date_to,
        "entries": entries,
        "total": len(entries),
        "by_category": by_category,
        "by_status": by_status,
        "n_faults": n_faults,
        "n_parts": n_parts,
        "groups": [groups[k] for k in sorted(groups.keys())],
    }


def _resolve_photo(img: dict) -> Optional[str]:
    """Return a local path for an image, downloading from the server if needed."""
    fp = img.get("file_path") or ""
    if fp and os.path.isfile(fp):
        return fp
    try:
        path = download_remote_image(img.get("id", ""), img.get("work_log_id", ""))
        if path and os.path.isfile(path):
            return path
    except Exception:
        pass
    return None


# ── Public entry point ────────────────────────────────────────────────────────

def generate_worklog_report(
    project_id: Optional[int],
    year: int,
    month: int,
    output_path: str,
    output_format: str = "pdf",
    include_photos: bool = True,
    project_name: str = "",
) -> str:
    data = build_report_data(project_id, year, month, project_name)
    if output_format.lower() == "docx":
        return _render_docx(data, output_path, include_photos)
    return _render_pdf(data, output_path, include_photos)


# ── PDF renderer (ReportLab) ──────────────────────────────────────────────────

def _render_pdf(data: dict, output_path: str, include_photos: bool) -> str:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
        Image as RLImage, KeepTogether, PageBreak,
    )

    NAVY = colors.HexColor("#1A2B45")
    LIGHT = colors.HexColor("#F1F5F9")
    GREY = colors.HexColor("#CBD5E1")

    FONT, FONTB = _register_cyrillic_fonts()

    styles = getSampleStyleSheet()
    h_title = ParagraphStyle("t", parent=styles["Title"], textColor=NAVY, fontSize=22,
                             fontName=FONTB)
    h_sec = ParagraphStyle("s", parent=styles["Heading2"], textColor=NAVY, fontSize=13,
                           spaceBefore=10, spaceAfter=4, fontName=FONTB)
    h_block = ParagraphStyle("b", parent=styles["Heading3"], textColor=NAVY, fontSize=11,
                             spaceBefore=8, spaceAfter=2, fontName=FONTB)
    body = ParagraphStyle("body", parent=styles["Normal"], fontSize=8, leading=10,
                          fontName=FONT)
    muted = ParagraphStyle("m", parent=styles["Normal"], fontSize=9,
                           textColor=colors.HexColor("#6B7A8D"), fontName=FONT)

    doc = SimpleDocTemplate(
        output_path, pagesize=A4,
        leftMargin=1.4*cm, rightMargin=1.4*cm,
        topMargin=1.4*cm, bottomMargin=1.4*cm,
        title=f"Technical Works Report — {data['month_name']} {data['year']}",
    )
    story = []

    # Cover
    story.append(Paragraph("Technical Works Report", h_title))
    story.append(Paragraph(f"{data['month_name']} {data['year']}", muted))
    story.append(Spacer(1, 0.3*cm))
    cover_rows = [
        ["Project:", data["project_name"]],
        ["Period:", data["period"]],
        ["Total works:", str(data["total"])],
        ["Generated:", datetime.now().strftime("%Y-%m-%d %H:%M")],
    ]
    ct = Table(cover_rows, colWidths=[4*cm, 13*cm])
    ct.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), LIGHT),
        ("TEXTCOLOR", (0, 0), (-1, -1), NAVY),
        ("FONTNAME", (0, 0), (-1, -1), FONT),
        ("FONTNAME", (0, 0), (0, -1), FONTB),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.3, GREY),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(ct)
    story.append(Spacer(1, 0.4*cm))

    if data["total"] == 0:
        story.append(Paragraph("No entries for the selected period.", h_sec))
        doc.build(story)
        return output_path

    # Summary
    story.append(Paragraph("1. Summary", h_sec))
    cat_lines = ", ".join(
        f"{_cat_label(k)}: {v}"
        for k, v in sorted(data["by_category"].items(), key=lambda x: -x[1])
    )
    st = data["by_status"]
    summ_rows = [
        ["Total works", str(data["total"])],
        ["By category", cat_lines or "—"],
        ["Statuses", f"Done: {st.get('done',0)}   Open: {st.get('open',0)}   "
                     f"No status: {st.get('',0)}"],
        ["With a recorded fault", str(data["n_faults"])],
        ["With spare parts used", str(data["n_parts"])],
    ]
    stbl = Table(summ_rows, colWidths=[6*cm, 11*cm])
    stbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), LIGHT),
        ("TEXTCOLOR", (0, 0), (-1, -1), NAVY),
        ("FONTNAME", (0, 0), (-1, -1), FONT),
        ("FONTNAME", (0, 0), (0, -1), FONTB),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.3, GREY),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(stbl)
    story.append(Spacer(1, 0.3*cm))

    # Works by block
    story.append(Paragraph("2. Works by Block / Equipment", h_sec))
    headers = ["Date", "Location", "Category", "Fault",
               "Description", "Status", "SAP", "Spare Parts"]
    col_w = [1.7*cm, 2.6*cm, 2.0*cm, 2.4*cm, 5.0*cm, 1.4*cm, 1.6*cm, 2.3*cm]

    for g in data["groups"]:
        block_flow = [Paragraph(g["title"], h_block)]
        rows = [headers]
        for e in g["entries"]:
            rows.append([
                Paragraph(e.get("log_date", ""), body),
                Paragraph(_entry_location(e), body),
                Paragraph(_cat_label(e.get("category", "")), body),
                Paragraph(e.get("fault_name") or "—", body),
                Paragraph((e.get("description") or "").replace("\n", "<br/>"), body),
                Paragraph(STATUS_LABELS.get((e.get("status") or "").lower(), "—"), body),
                Paragraph(e.get("sap_ticket") or "—", body),
                Paragraph(e.get("spare_parts") or "—", body),
            ])
        tbl = Table(rows, colWidths=col_w, repeatRows=1)
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), NAVY),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), FONTB),
            ("FONTSIZE", (0, 0), (-1, 0), 8),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]),
            ("GRID", (0, 0), (-1, -1), 0.3, GREY),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        block_flow.append(tbl)
        block_flow.append(Spacer(1, 0.25*cm))
        # Keep small blocks (heading + short table) on one page; let big ones flow.
        if len(g["entries"]) <= 4:
            story.append(KeepTogether(block_flow))
        else:
            story.extend(block_flow)

    # Photo appendix
    if include_photos:
        photo_entries = []
        for e in data["entries"]:
            imgs = get_images_for_log(e["id"])
            if imgs:
                photo_entries.append((e, imgs))
        if photo_entries:
            story.append(PageBreak())
            story.append(Paragraph("3. Photo Appendix", h_sec))
            for e, imgs in photo_entries:
                cap = f"{e.get('log_date','')} · {_entry_location(e)} · " \
                      f"{_cat_label(e.get('category',''))}"
                cells = []
                row = []
                for img in imgs:
                    path = _resolve_photo(img)
                    if not path:
                        continue
                    try:
                        iw, ih = ImageReader(path).getSize()
                        w = 5.2*cm
                        h = w * (ih / iw) if iw else 4*cm
                        h = min(h, 5.2*cm)
                        row.append(RLImage(path, width=w, height=h))
                    except Exception:
                        continue
                    if len(row) == 3:
                        cells.append(row); row = []
                if row:
                    cells.append(row)
                if not cells:
                    continue
                pt = Table(cells, colWidths=[5.6*cm]*3)
                pt.setStyle(TableStyle([
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]))
                story.append(KeepTogether([Paragraph(cap, h_block), pt,
                                           Spacer(1, 0.2*cm)]))

    doc.build(story)
    return output_path


# ── DOCX renderer ─────────────────────────────────────────────────────────────

def _render_docx(data: dict, output_path: str, include_photos: bool) -> str:
    from docx.shared import Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from services import docx_renderer as dr

    doc = dr.make_doc(margin_cm=1.8)

    dr.add_heading(doc, "Technical Works Report", level=1)
    dr.add_paragraph(doc, f"{data['month_name']} {data['year']}",
                     size=12, color=dr.TEXT_MUTE)
    dr.add_styled_table(
        doc,
        ["Parameter", "Value"],
        [
            ["Project", data["project_name"]],
            ["Period", data["period"]],
            ["Total works", str(data["total"])],
            ["Generated", datetime.now().strftime("%Y-%m-%d %H:%M")],
        ],
    )
    doc.add_paragraph()

    if data["total"] == 0:
        dr.add_paragraph(doc, "No entries for the selected period.", size=10)
        dr.save_doc(doc, output_path)
        return output_path

    # Summary
    dr.add_heading(doc, "1. Summary", level=2)
    st = data["by_status"]
    cat_lines = ", ".join(
        f"{_cat_label(k)}: {v}"
        for k, v in sorted(data["by_category"].items(), key=lambda x: -x[1])
    )
    dr.add_styled_table(
        doc,
        ["Metric", "Value"],
        [
            ["Total works", str(data["total"])],
            ["By category", cat_lines or "—"],
            ["Statuses",
             f"Done: {st.get('done',0)}; Open: {st.get('open',0)}; "
             f"No status: {st.get('',0)}"],
            ["With a fault", str(data["n_faults"])],
            ["With spare parts", str(data["n_parts"])],
        ],
    )
    doc.add_paragraph()

    # Works by block
    dr.add_heading(doc, "2. Works by Block / Equipment", level=2)
    headers = ["Date", "Location", "Category", "Fault",
               "Description", "Status", "SAP", "Spare Parts"]
    for g in data["groups"]:
        dr.add_heading(doc, g["title"], level=3)
        rows = []
        for e in g["entries"]:
            rows.append([
                e.get("log_date", ""),
                _entry_location(e),
                _cat_label(e.get("category", "")),
                e.get("fault_name") or "—",
                e.get("description") or "",
                STATUS_LABELS.get((e.get("status") or "").lower(), "—"),
                e.get("sap_ticket") or "—",
                e.get("spare_parts") or "—",
            ])
        dr.add_styled_table(doc, headers, rows)
        doc.add_paragraph()

    # Photo appendix
    if include_photos:
        photo_entries = []
        for e in data["entries"]:
            imgs = get_images_for_log(e["id"])
            if imgs:
                photo_entries.append((e, imgs))
        if photo_entries:
            doc.add_page_break()
            dr.add_heading(doc, "3. Photo Appendix", level=2)
            for e, imgs in photo_entries:
                cap = f"{e.get('log_date','')} · {_entry_location(e)} · " \
                      f"{_cat_label(e.get('category',''))}"
                dr.add_heading(doc, cap, level=3)
                paths = [p for p in (_resolve_photo(i) for i in imgs) if p]
                if not paths:
                    continue
                # 3-per-row table of pictures
                ncol = 3
                nrow = (len(paths) + ncol - 1) // ncol
                table = doc.add_table(rows=nrow, cols=ncol)
                for idx, path in enumerate(paths):
                    cell = table.cell(idx // ncol, idx % ncol)
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    try:
                        p.add_run().add_picture(path, width=Inches(2.0))
                    except Exception:
                        p.add_run("[photo unavailable]")
                doc.add_paragraph()

    dr.save_doc(doc, output_path)
    return output_path

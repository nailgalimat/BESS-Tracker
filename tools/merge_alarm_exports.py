"""tools/merge_alarm_exports.py — put daily SCADA alarm exports back together.

    python tools\\merge_alarm_exports.py "C:\\path\\to\\folder"
    python tools\\merge_alarm_exports.py "C:\\path\\to\\folder" --out merged.xlsx

The SCADA alarm export truncates at roughly 4,000 rows, so a month has to be
taken a day at a time. This reads every workbook in the folder, keeps every
sheet (Communications / Production / Warning / …), drops rows that appear in
more than one file, and writes one workbook in the same shape — which is what
the report importer and the analysis scripts expect.

Reports what it found per day, so a missing day is obvious before the merged
file is used for anything.
"""
import argparse
import glob
import os
import sys
from collections import Counter, defaultdict

import openpyxl

ID, ACTIVATED = 0, 2          # column positions in the SCADA export


def read_book(path):
    """{sheet: [row, ...]} plus the header of each sheet."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out, heads = {}, {}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        heads[name] = rows[0]
        out[name] = [r for r in rows[1:] if r and any(v is not None for v in r)]
    wb.close()
    return out, heads


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('folder', help='folder holding the daily exports')
    ap.add_argument('--out', default='', help='output file (default: <folder>\\Alarms merged.xlsx)')
    a = ap.parse_args()

    files = sorted(f for f in glob.glob(os.path.join(a.folder, '*.xls*'))
                   if not os.path.basename(f).startswith('~$')
                   and 'merged' not in os.path.basename(f).lower())
    if not files:
        sys.exit('no workbooks in {}'.format(a.folder))
    print('{} file(s)'.format(len(files)))

    sheets = defaultdict(list)
    heads = {}
    seen = defaultdict(set)
    dupes = 0
    per_file = []

    for path in files:
        try:
            book, hd = read_book(path)
        except Exception as ex:                                  # noqa: BLE001
            print('  SKIPPED {} — {}'.format(os.path.basename(path), ex))
            continue
        heads.update({k: v for k, v in hd.items() if k not in heads})
        kept = 0
        days = Counter()
        for name, rows in book.items():
            for r in rows:
                key = r[ID] if r[ID] is not None else r
                if key in seen[name]:
                    dupes += 1
                    continue
                seen[name].add(key)
                sheets[name].append(r)
                kept += 1
                if len(r) > ACTIVATED and r[ACTIVATED]:
                    days[str(r[ACTIVATED])[:10]] += 1
        per_file.append((os.path.basename(path), kept, sorted(days)))

    print('\nper file:')
    for name, kept, days in per_file:
        span = '{} .. {}'.format(days[0], days[-1]) if days else 'no dates'
        print('  {:<44} {:5d} row(s)  {}'.format(name[:44], kept, span))

    all_days = sorted({d for _, _, ds in per_file for d in ds})
    print('\ndays covered: {}'.format(len(all_days)))
    if all_days:
        print('  {} .. {}'.format(all_days[0], all_days[-1]))
        # name the gaps: a day with no alarms at all is rare enough to check
        import datetime
        d0 = datetime.date.fromisoformat(all_days[0])
        d1 = datetime.date.fromisoformat(all_days[-1])
        have = set(all_days)
        missing = []
        d = d0
        while d <= d1:
            if d.isoformat() not in have:
                missing.append(d.isoformat())
            d += datetime.timedelta(days=1)
        if missing:
            print('  MISSING: {}'.format(', '.join(missing)))
        else:
            print('  no gaps')
    if dupes:
        print('duplicate rows dropped: {}'.format(dupes))

    out = a.out or os.path.join(a.folder, 'Alarms merged.xlsx')
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name in sorted(sheets):
        ws = wb.create_sheet(name[:31])
        if name in heads:
            ws.append(list(heads[name]))
        rows = sheets[name]
        # by time where there is one, so the merged file reads like one export
        try:
            rows.sort(key=lambda r: str(r[ACTIVATED] or ''))
        except Exception:                                        # noqa: BLE001
            pass
        for r in rows:
            ws.append(list(r))
        print('  sheet {:<18} {} row(s)'.format(name, len(rows)))
    wb.save(out)
    print('\nwrote {}'.format(out))


if __name__ == '__main__':
    main()

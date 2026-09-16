"""Report versions: every generation a new vN file that is never written over,
a manifest that says what the numbers were made from, "mark as sent" with the
Word-edited file, and a sent month locked until it is unlocked with a reason.

QA S1 gate G9 (vN, SHA of v1 unchanged), G11 (manifest), G13 (locked month) and
WF-REP-08/09, DATA-REPRO-01/03. The generators are replaced by stand-ins that
write a small DOCX and report key numbers, so this runs in seconds; the real
August report through the same path is checked in test_real_tashkent_august.
"""
import _harness as H            # must be first: isolates DB, sync, network, history
import functools
import hashlib
import json
import os
import stat
import warnings
warnings.filterwarnings('ignore')

import database.db_manager as dbm
H.fresh_db('versions.db')
import _synthetic_scada as S
import services.parse_cache as pc
import services.report_workflow_service as rw
import services.availability_service as av
import services.availability_inputs_service as avi
import services.month_dataset_service as mds
import services.report_versions_service as rvs
import services.tashkent_report_service as T
import services.bukhara_report_service as B

c = dbm.get_connection()
c.execute("INSERT INTO projects (name, num_zones, num_blocks, num_containers) VALUES ('ACWA TK', 1, 70, 1)")
c.execute("INSERT INTO projects (name, num_zones, num_blocks, num_containers) VALUES ('BK', 1, 15, 1)")
PID = c.execute("SELECT id FROM projects WHERE name='ACWA TK'").fetchone()[0]
BID = c.execute("SELECT id FROM projects WHERE name='BK'").fetchone()[0]
c.commit(); c.close()
rw.save_project_config(PID, site_type='tashkent', contractual_plant_capacity_mw=441)
rw.save_project_config(BID, site_type='bukhara')
Y, M = 2026, 9

# ── stand-in generators with the real signatures ───────────────────────────
STATE = {'avail': 99.80, 'rows441': 18}
CALLS = []


def _fake(real, site):
    @functools.wraps(real)
    def gen(**kw):
        CALLS.append((site, kw))
        from docx import Document
        d = Document()
        d.add_paragraph(f"{site} availability for the period was {STATE['avail']:.2f}%")
        d.save(kw['output_path'])
        kw['summary_out'].update({
            'availability_pct': STATE['avail'], 'rte_pct': 88.12, 'cycles_month': 26.6,
            'cycles_in_year': 249.6, 'cycles_lifetime': 256.4, 'avg_soc_pct': 44.1,
            'avg_soh_pct': 98.9,
            'rows': {'3.1': len(kw.get('pm_activities') or []), '3.2': 3,
                     '4.4.1': STATE['rows441'], '5.2': 2}})
        return kw['output_path']
    return gen


T.generate_tashkent_report = _fake(T.generate_tashkent_report, 'tashkent')
B.generate_bukhara_report = _fake(B.generate_bukhara_report, 'bukhara')

files = S.tashkent_month(os.path.join(H.WORK, 'exports'), Y, M)
mds.add_files(PID, Y, M, list(files.values()))
rw.record_pm(PID, '41', '2026-09-10', None, 4, 'PM as per checklist')
EXCL = av.add_exclusion('Grid Outage', '2026-09-11', '2026-09-12', '23:48', '04:25',
                        affected_blocks='1-70', project_id=PID, year=Y, month=M)


def params(pid=PID):
    return dict(rw.report_inputs(pid, Y, M), site_name='ACWA TK', project_details={'Customer': 'X'},
                report_number='APRS-BESS-2026-09', cm_activities=['Block 33: LCU VFD1 board — replaced'],
                safety_incidents=None, recommendations=['Keep coolant topped up'])


def sha(p):
    return hashlib.sha256(open(p, 'rb').read()).hexdigest()


print('=== v1 ===')
CONF = [{'kind': 'gap', 'severity': 'warn', 'title': 'Data gap', 'detail': 'PCS fault status · no data 04–30.09'}]
v1 = rvs.generate_version(PID, Y, M, params(), CONF, 'month still running')
H.check(v1['version'] == 1 and os.path.basename(v1['docx_abs']) == 'ACWA_TK_LTSA_2026-09_v1.docx',
        'first version file: {}'.format(os.path.basename(v1['docx_abs'])))
H.check(os.path.isfile(v1['docx_abs']) and sha(v1['docx_abs']) == v1['docx_sha256'], 'DOCX hash stored')
H.check(not os.access(v1['docx_abs'], os.W_OK) and not os.access(v1['manifest_abs'], os.W_OK),
        'the DOCX and its manifest are read-only')
site, kw = CALLS[-1]
H.check(site == 'tashkent' and kw['working_status_path'] == mds.active_by_type(PID, Y, M)['working_status']['abs_path']
        and kw['project_id'] == PID and kw['output_format'] == 'docx',
        'the generator got the data set copies, project_id and docx — no file picked')

print('\n=== the manifest ===')
man = rvs.load_manifest(v1['manifest_abs'])
for key in ('input_files', 'report_inputs', 'pm_records', 'text', 'history_records',
            'alarm_classifications', 'app', 'confirmations', 'key_numbers', 'docx'):
    H.check(key in man, 'manifest has {}'.format(key))
ds = {f['sha256'] for f in mds.list_files(PID, Y, M)}
H.check({f['sha256'] for f in man['input_files']} == ds and len(ds) == len(files) == 9,
        'input file hashes = the data set ({} files)'.format(len(ds)))
H.check(man['report_inputs'] == json.loads(json.dumps(rw.report_inputs(PID, Y, M), default=str)),
        'report_inputs exactly as the database gives them')
H.check(any(e['id'] == EXCL for e in man['report_inputs']['exclusions'])
        and any(r.get('pm_id') for r in man['report_inputs']['manual_unavailability']),
        'exclusions and PM-as-downtime are in it')
H.check(man['pm_records'] and man['pm_records'][0]['affected_blocks'] == '41', 'PM records as stored')
H.check(man['history_records'] == json.loads(json.dumps(B.load_history_records(B.DEFAULT_HISTORY_PATH),
                                                        default=str)),
        'the KPI history the cycles came from ({} records)'.format(len(man['history_records'])))
H.check(man['alarm_classifications']['sha256'] == pc.file_sha256(B.DEFAULT_ALARM_CLASSIFICATIONS_CSV),
        'alarm classification CSV hash')
H.check(man['app'].get('code_fingerprint') and (man['app'].get('commit') or man['app'].get('exe')),
        'app: code fingerprint {} and commit {}'.format(man['app'].get('code_fingerprint'),
                                                        (man['app'].get('commit') or '')[:8]))
H.check(man['confirmations'] == CONF and man['confirmation_note'] == 'month still running',
        'the confirmed question and the note are recorded')
H.check(man['key_numbers']['availability_pct'] == 99.80 and man['key_numbers']['rows']['4.4.1'] == 18,
        'key numbers and row counts')
H.check(man['text']['report_number'] == 'APRS-BESS-2026-09' and man['text']['cm_activities'],
        'customer text')
H.check('fault_notes' not in open(v1['manifest_abs'], encoding='utf-8').read(), 'no internal notes')

print('\n=== v2 after an input change; v1 untouched ===')
v1_sha = sha(v1['docx_abs'])
av.update_exclusion(EXCL, 'Grid Outage', '2026-09-11', '2026-09-12', '23:30', '04:25',
                    affected_blocks='1-70', year=Y, month=M)
STATE.update(avail=99.77, rows441=17)
v2 = rvs.generate_version(PID, Y, M, params())
H.check(v2['version'] == 2 and v2['docx_abs'] != v1['docx_abs'], 'v2 is a new file')
H.check(sha(v1['docx_abs']) == v1_sha and v1_sha == rvs.get_version(v1['id'])['docx_sha256'],
        'v1 SHA-256 unchanged')
vers = rvs.list_versions(PID, Y, M)
H.check([v['version'] for v in vers] == [2, 1], 'listed newest first')
ch = vers[0]['changes']
print('   v2:', ch)
H.check('Availability 99.80 % → 99.77 %' in ch and '4.4.1 rows 18 → 17' in ch
        and 'availability inputs changed' in ch, 'what changed: numbers, 4.4.1 rows, inputs')
H.check(vers[1]['changes'] == 'first version', 'v1: first version')

stray = os.path.join(os.path.dirname(v1['docx_abs']), 'ACWA_TK_LTSA_2026-09_v3.docx')
with open(stray, 'wb') as fh:
    fh.write(b'not ours')
v3 = rvs.generate_version(PID, Y, M, params())
H.check(v3['version'] == 4 and open(stray, 'rb').read() == b'not ours',
        'a file already named v3 is not written over: the next version is v{}'.format(v3['version']))
H.check(rvs.list_versions(PID, Y, M)[0]['changes'].startswith('No number changes vs v2'),
        'unchanged inputs: "{}"'.format(rvs.list_versions(PID, Y, M)[0]['changes']))

print('\n=== nothing is written when it cannot be generated ===')
n_before = len(rvs.list_versions(PID, Y, M))
soc = mds.active_by_type(PID, Y, M)['soc']
os.chmod(soc['abs_path'], stat.S_IWRITE | stat.S_IREAD)
with open(soc['abs_path'], 'ab') as fh:
    fh.write(b'tampered')
try:
    rvs.generate_version(PID, Y, M, params()); H.check(False, 'refused')
except RuntimeError as e:
    H.check('changed on disk' in str(e), 'a stored copy that no longer matches its hash stops it: {}'.format(e))
H.check(mds.add_files(PID, Y, M, [files['soc']])[0]['status'] == 'restored',
        'adding the same file again restores the stored copy')
mds.remove_file(mds.active_by_type(PID, Y, M)['hv_meter']['id'])
try:
    rvs.generate_version(PID, Y, M, params()); H.check(False, 'refused')
except ValueError as e:
    H.check('HV meter daily' in str(e), 'a missing required file stops it: {}'.format(e))
H.check(len(rvs.list_versions(PID, Y, M)) == n_before, 'no version row added by the failures')
mds.add_files(PID, Y, M, [files['hv_meter']])

print('\n=== mark the Word-edited file as sent: the month locks ===')
edited = os.path.join(H.WORK, 'APRS BESS September monthly report.docx')
from docx import Document
d = Document(); d.add_paragraph('edited in Word'); d.save(edited)
row = rvs.mark_sent(v2['id'], edited)
H.check(row['sent_at'] and row['sent_docx_sha256'] == sha(edited) and os.path.isfile(row['sent_abs'])
        and sha(row['sent_abs']) == sha(edited), 'copy and SHA-256 stored with v2')
st = rvs.month_status(PID, Y, M)
H.check(st['locked_at'] and st['sent']['version'] == 2 and st['log'][-1]['action'] == 'sent',
        'month locked, sent = v2, logged')


def digest():
    h = hashlib.sha256()
    conn = dbm.get_connection()
    try:
        for t in ('pm_activities', 'manual_unavailability', 'availability_exclusions',
                  'balancing_periods', 'report_months', 'month_files'):
            h.update(repr([tuple(r) for r in conn.execute(f'SELECT * FROM {t} ORDER BY rowid')]).encode())
    finally:
        conn.close()
    return h.hexdigest()


before = digest()
pm_id = rw.get_pm_activities(PID, Y, M)[0]['id']
aug_excl = av.add_exclusion('Grid Outage', '2026-08-20', '2026-08-20', '10:00', '11:00',
                            affected_blocks='3', project_id=PID, year=2026, month=8)
before = digest()
refused = {
    'add exclusion': lambda: av.add_exclusion('Grid Outage', '2026-09-20', '2026-09-20', '10:00', '11:00',
                                              affected_blocks='3', project_id=PID, year=Y, month=M),
    'add exclusion for every project (NULL project)': lambda: av.add_exclusion(
        'Grid Outage', '2026-09-20', '2026-09-20', '10:00', '11:00', affected_blocks='3'),
    'edit exclusion': lambda: av.update_exclusion(EXCL, 'Grid Outage', '2026-09-11', '2026-09-12',
                                                  '23:00', '04:25', affected_blocks='1-70'),
    'move an August window into September': lambda: av.update_exclusion(
        aug_excl, 'Grid Outage', '2026-09-20', '2026-09-20', '10:00', '11:00', affected_blocks='3',
        year=Y, month=M),
    'delete exclusion': lambda: av.delete_exclusion(EXCL),
    'add PM': lambda: rw.record_pm(PID, '42', '2026-09-11', None, 4, 'PM'),
    'edit PM': lambda: rw.update_pm_record(pm_id, '41', '2026-09-10', None, 3, 'PM'),
    'delete PM': lambda: rw.delete_pm_activity(pm_id),
    'add manual downtime': lambda: avi.record_manual(PID, '5', '2026-09-12', None, 2),
    'add balancing': lambda: av.add_balancing_period('2026-09-01', '2026-09-30', '17', project_id=PID,
                                                     year=Y, month=M),
    'save narrative': lambda: rw.save_report_month(PID, Y, M, recommendations='changed'),
    'add a data file': lambda: mds.add_files(PID, Y, M, [files['soc']]),
}
for what, fn in refused.items():
    try:
        fn()
        H.check(False, '{} refused'.format(what))
    except rw.MonthLockedError:
        H.check(True, '{} -> MonthLockedError'.format(what))
H.check(digest() == before, 'the locked month\'s inputs are unchanged')
H.check(avi.route_field_event({'id': 'pm-late-1', 'kind': 'pm', 'project_id': PID, 'blocks': '43',
                               'date_from': '2026-09-13', 'hours': 4}) == 'queued'
        and avi.get_queued_event('pm-late-1')['reason'] == 'locked' and digest() == before,
        'a phone PM for the locked month waits ("locked") instead of being applied or retried forever')
oct_id = av.add_exclusion('Grid Outage', '2026-10-02', '2026-10-02', '10:00', '11:00',
                          affected_blocks='3', project_id=PID, year=2026, month=10)
H.check(bool(oct_id), 'another month is not affected')

print('\n=== unlock needs a reason; lock again ===')
for bad in ('', '  ', 'x'):
    try:
        rvs.unlock_month(PID, Y, M, bad); H.check(False, 'refused')
    except ValueError:
        pass
H.check(rw.month_locked_at(PID, Y, M), 'no reason -> still locked')
rvs.unlock_month(PID, Y, M, 'customer asked to correct the 12.09 outage end time')
st = rvs.month_status(PID, Y, M)
H.check(not st['locked_at'] and st['log'][-1]['action'] == 'unlocked'
        and st['log'][-1]['reason'].startswith('customer asked'), 'unlocked, the reason is logged')
av.update_exclusion(EXCL, 'Grid Outage', '2026-09-11', '2026-09-12', '23:48', '04:40',
                    affected_blocks='1-70', year=Y, month=M)
H.check(True, 'inputs editable again after the unlock')
rvs.lock_month(PID, Y, M)
H.check(rw.month_locked_at(PID, Y, M) and rvs.month_status(PID, Y, M)['log'][-1]['action'] == 'locked',
        'locked again, logged')

print('\n=== Bukhara goes through the same versions ===')
bfolder = os.path.join(H.WORK, 'bk')
os.makedirs(bfolder, exist_ok=True)
for key in ('bk_site_kpi', 'bk_overall_lc', 'bk_battery_unit', 'bk_meter_daily', 'bk_alarms'):
    p = os.path.join(bfolder, key + '.xlsx')
    S.wide(p, S.stamps('2026-09-01', '2026-09-02', minutes=1440), ['x'], lambda t, c: 1)
    mds.store_file(BID, Y, M, p, key, 'chosen')
bv = rvs.generate_version(BID, Y, M, dict(rw.report_inputs(BID, Y, M), site_name='BK', project_details={}))
site, kw = CALLS[-1]
H.check(site == 'bukhara' and bv['version'] == 1 and 'site_kpi_path' in kw
        and 'manual_unavailability' not in kw and 'project_id' not in kw,
        'Bukhara v1: only arguments its generator takes')

print('\n=== the Release step ===')
from PyQt5.QtWidgets import QApplication
app = QApplication.instance() or QApplication([])
import ui.monthly_reports_page as mrp
page = mrp.MonthlyReportsPage()
page.proj_combo.setCurrentIndex(page.proj_combo.findData(PID))
page.year_spin.setValue(Y); page.month_combo.setCurrentIndex(M - 1)
page._open_month()
if page._an_worker:
    page._an_worker.wait(60000)
H.check(page.versions_table.rowCount() == 3 and page.versions_table.item(0, 1).text() == 'v4',
        'versions table: v4, v2, v1')
H.check(page.gen_btn.text().endswith('Generate v5'), 'button: "{}"'.format(page.gen_btn.text()))
H.check('sent v2' in page.tabs.tabText(page.tabs.indexOf(page._release_tab)) and 'locked' in
        page.tabs.tabText(page.tabs.indexOf(page._release_tab)), 'tab: "{}"'.format(
            page.tabs.tabText(page.tabs.indexOf(page._release_tab))))
H.check(page.unlock_btn.isEnabled() and not page.relock_btn.isEnabled()
        and not page.add_files_btn.isEnabled() and page.recommendations.isReadOnly()
        and page._lock_banners[0].text().startswith('🔒'),
        'locked: Unlock on, adding files off, narrative read-only, banner shown')

print(chr(10) + '=== generating from the Release step (page, worker, dialogs) ===')
from PyQt5.QtWidgets import QDialog, QMessageBox
OCT = 10
for path in files.values():
    mds.add_files(PID, Y, OCT, [path])


class _AcceptConfirm:                     # the "open questions" dialog, answered "generate anyway"
    def __init__(self, parent, questions):
        self.asked = questions
        self.note = type('N', (), {'text': staticmethod(lambda: 'checked with the SCADA team')})()
        SEEN.append(questions)

    def exec_(self):
        return QDialog.Accepted


SEEN = []
mrp.ConfirmGenerateDialog = _AcceptConfirm
mrp.QMessageBox.information = staticmethod(lambda *a, **k: None)
mrp.QMessageBox.warning = staticmethod(lambda *a, **k: None)
mrp.QMessageBox.critical = staticmethod(lambda *a, **k: None)
page.year_spin.setValue(Y); page.month_combo.setCurrentIndex(OCT - 1)
page._open_month()
if page._an_worker:
    page._an_worker.wait(120000)
page.report_number.setText('APRS-BESS-2026-10')
page._generate()
H.check(page._worker is not None and page._worker.wait(180000), 'the generation worker finished')
app.processEvents()
vers = rvs.list_versions(PID, Y, OCT)
H.check(len(vers) == 1 and vers[0]['version'] == 1 and os.path.isfile(vers[0]['docx_abs']),
        'October v1 written from the page')
H.check(SEEN and any(q['kind'] == 'month_end' for q in SEEN[0]),
        'the open questions were shown before generating ({} of them)'.format(len(SEEN[0]) if SEEN else 0))
man = rvs.load_manifest(vers[0]['manifest_abs'])
H.check(len(man['confirmations']) == len(SEEN[0])
        and man['confirmation_note'] == 'checked with the SCADA team',
        'the manifest records what was confirmed and the note')
H.check(man['text']['report_number'] == 'APRS-BESS-2026-10', 'the customer text from the page')
H.check(page.versions_table.rowCount() == 1 and 'v1' in page.log.toPlainText(),
        'the Release step lists it: "{}"'.format(page.log.toPlainText()[-60:].replace(chr(10), ' ')))

H.finish()

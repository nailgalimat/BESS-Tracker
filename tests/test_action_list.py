"""The action list: the office's organisational items, from the Excel sheet to
the Plan page and the Today screen.

What this pins down:

  * the column guess — the user's real sheet has "No. · Topic · Description ·
    Remarks / To do · Target Date", multi-line cells and a Target Date that
    arrives as a datetime;
  * a **re-import** of the same sheet with one row changed and one row added:
    nothing is duplicated and nothing already worked on is lost — the sheet is
    a living document and it is imported again every few weeks;
  * the line that matters most: an action item has no plant block, and it
    **never** reaches the customer's monthly report — not section 3.2, not the
    PM hours. It is a table of its own precisely so that it cannot;
  * the Plan page's Action list tab and the Today panel, driven headless, with
    the Today rule: a count is the length of the very list its link opens.

The user's own Book1.xlsx is used where it helps and skipped when it is not on
this computer; the synthetic sheet keeps the test running anywhere.
"""
import datetime
import os

import _harness as H            # must be first
import openpyxl

import database.db_manager as dbm
import services.action_list_service as als

TODAY = datetime.date(2026, 9, 23)
REAL = os.environ.get('BESS_ACTION_XLSX',
                      r'C:\Users\user1\Desktop\ACWA BESS Tashkent\Book1.xlsx')

H.fresh_db('actions.db')
c = dbm.get_connection()
c.execute("INSERT INTO projects (name, num_zones, num_blocks, num_containers, "
          "project_type) VALUES ('TK', 2, 16, 4, 'BESS')")
c.execute("INSERT INTO projects (name, num_zones, num_blocks, num_containers, "
          "project_type) VALUES ('BK', 2, 15, 4, 'BESS')")
PID, OTHER = [r[0] for r in c.execute("SELECT id FROM projects ORDER BY id")]
for z in (1, 2):
    for b in range(1, 9):
        c.execute("INSERT INTO containers (project_id, zone_number, block_number, "
                  "container_index, container_type, serial_number) "
                  "VALUES (?,?,?,1,'LC Cabinet',?)", (PID, z, b, 'S%d%d' % (z, b)))
c.commit(); c.close()
import services.project_service as psvc
psvc.get_block_map(PID, refresh=True)


# ── a sheet shaped exactly like the user's ───────────────────────────────
def sheet(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws.append(['No.', 'Topic', 'Description', 'Remarks / To do', 'Target Date'])
    for r in rows:
        ws.append(list(r))
    wb.save(path)
    return path


V1 = sheet(os.path.join(H.WORK, 'actions_v1.xlsx'), [
    (1, 'PM activities',
     'Checklists\nPhotos\nReport with conclusions',
     'Complete CL\nPhotos\nReport', datetime.datetime(2026, 10, 5)),
    (2, 'Spare parts of EPC', '26 battery packs\nany others spare parts?',
     'Confirm EPC purchased, compare with Annex 13',
     datetime.datetime(2026, 9, 28)),
    (3, 'Gas sensors', 'How often need to replace?',
     'Confirm frequency of replacement', datetime.datetime(2026, 9, 28)),
    # the user's real file ends with numbered but empty spacer rows
    (4, None, None, None, None),
    (5, None, None, None, None),
])

pv = als.preview_excel(V1)
H.check(pv['mapping'].get('No.') == 'seq' and pv['mapping'].get('Topic') == 'topic'
        and pv['mapping'].get('Description') == 'description'
        and pv['mapping'].get('Remarks / To do') == 'todo'
        and pv['mapping'].get('Target Date') == 'due_date',
        'the column guess reads all five columns: {}'.format(pv['mapping']))

res = als.import_from_excel(PID, V1, pv['mapping'], log=lambda *_a: None)
H.check(res['added'] == 3 and res['updated'] == 0,
        'three items imported, the numbered blank rows skipped: {}'.format(res))
items = als.items(PID)
gas = [i for i in items if i['topic'] == 'Gas sensors'][0]
pm = [i for i in items if i['topic'] == 'PM activities'][0]
H.check(pm['description'] == 'Checklists\nPhotos\nReport with conclusions'
        and pm['todo'] == 'Complete CL\nPhotos\nReport',
        'a multi-line cell stays multi-line: {!r}'.format(pm['todo']))
H.check(gas['due_date'] == '2026-09-28' and gas['seq'] == 3
        and gas['status'] == 'open',
        "Excel's datetime becomes a plain date: {}".format(gas['due_date']))
H.check(all(i['uuid'] for i in items), 'every item has a uuid, so a phone can name it')

# ── someone works on one of them ─────────────────────────────────────────
als.assign(gas['uuid'], 'u-tech1', 'tech1')
als.save(PID, pm['uuid'], status='in progress')
worked = als.get(gas['uuid'])
H.check(worked['assigned_to'] == 'u-tech1' and worked['assigned_name'] == 'tech1',
        'an item can be given to a person')

# ── the same sheet again: one row changed, one row added ─────────────────
V2 = sheet(os.path.join(H.WORK, 'actions_v2.xlsx'), [
    (1, 'PM activities', 'Checklists\nPhotos\nReport with conclusions',
     'Complete CL\nPhotos\nReport', datetime.datetime(2026, 10, 5)),
    (2, 'Spare parts of EPC', '26 battery packs\nany others spare parts?',
     'Confirm EPC purchased, compare with Annex 13',
     datetime.datetime(2026, 9, 28)),
    (3, 'Gas sensors', 'How often need to replace? Every 2 years as per UM?',
     'Confirm frequency of replacement\nNeed code for sensor alone',
     datetime.datetime(2026, 10, 12)),                    # ← changed
    (4, 'HVAC', 'Need BOM and order critical spares',
     'Get the BOM\nPlace order', datetime.datetime(2026, 9, 30)),   # ← new
])
res2 = als.import_from_excel(PID, V2, pv['mapping'], log=lambda *_a: None)
after = als.items(PID)
H.check(res2['added'] == 1 and res2['updated'] == 1 and res2['unchanged'] == 2,
        'the re-import adds one, updates one, leaves two alone: {}'.format(res2))
H.check(len(after) == 4 and len({i['topic'] for i in after}) == 4,
        'four items, nothing duplicated: {}'.format([i['topic'] for i in after]))
gas2 = als.get(gas['uuid'])
H.check(gas2['due_date'] == '2026-10-12'
        and 'Need code for sensor alone' in gas2['todo'],
        'the row that changed in the file is updated: {}'.format(gas2['due_date']))
H.check(gas2['assigned_to'] == 'u-tech1'
        and als.get(pm['uuid'])['status'] == 'in progress',
        'and the work already done on it survives the re-import — the '
        'assignment and the status are ours, not the spreadsheet\'s')

# an item that has fallen out of the file is KEPT, and said so
V3 = sheet(os.path.join(H.WORK, 'actions_v3.xlsx'),
           [(1, 'PM activities', 'Checklists\nPhotos\nReport with conclusions',
             'Complete CL\nPhotos\nReport', datetime.datetime(2026, 10, 5))])
res3 = als.import_from_excel(PID, V3, pv['mapping'], log=lambda *_a: None)
H.check(len(als.items(PID)) == 4 and len(res3['untouched']) == 3,
        'an item no longer in the file is kept, not silently dropped: {}'.format(
            res3['untouched']))

# ── the whole point: none of this is plant work ──────────────────────────
import services.report_workflow_service as rw

cm = rw.corrective_rows(PID, 2026, 9)
H.check(cm == [], 'no action item reaches section 3.2: {} row(s)'.format(len(cm)))
H.check(rw.get_pm_activities(PID, 2026, 9) == [],
        'and none of them becomes PM hours')
H.check(rw.get_pm_activities(PID, 2026, 10) == [],
        'not in the month they are due either')
conn = dbm.get_connection()
n_journal = conn.execute("SELECT COUNT(*) FROM work_log_entries").fetchone()[0]
n_items = conn.execute("SELECT COUNT(*) FROM action_items").fetchone()[0]
conn.close()
H.check(n_journal == 0 and n_items == 4,
        'nothing was written into the work journal: {} journal row(s), {} action '
        'item(s)'.format(n_journal, n_items))

# and one plant's list is not another's
H.check(als.items(OTHER) == [], 'the other project has an empty list of its own')

# ── done, and the export the office sends on ─────────────────────────────
als.mark_done(gas['uuid'], note='Honeywell: every 2 years, code 30110',
              by='tech1', source='phone')
done = als.get(gas['uuid'])
H.check(done['status'] == 'done' and done['done_at'] and done['done_by'] == 'tech1',
        'marking it done records who and when')
H.check(len(als.items(PID, include_done=False)) == 3,
        'and it leaves the open list')

out = os.path.join(H.WORK, 'action list out.xlsx')
exp = als.export_excel(PID, out)
wb = openpyxl.load_workbook(out)
ws = wb.active
head = [c.value for c in ws[1]]
row = {head[i]: ws.cell(row=2, column=i + 1).value for i in range(len(head))}
H.check(exp['rows'] == 4 and 'Status' in head and 'Done by' in head
        and 'Done on' in head,
        'the export carries the status, who did it and when: {}'.format(head))
gas_row = [r for r in ws.iter_rows(min_row=2, values_only=True)
           if r[1] == 'Gas sensors'][0]
H.check(gas_row[5] == 'Done' and gas_row[7] == 'tech1'
        and 'code 30110' in (gas_row[9] or ''),
        'the finished item says so in the sheet: {}'.format(gas_row[5:]))
pm_row = [r for r in ws.iter_rows(min_row=2, values_only=True)
          if r[1] == 'PM activities'][0]
H.check('\n' in (pm_row[3] or ''),
        'and a multi-line cell goes back out as multi-line: {!r}'.format(pm_row[3]))
wb.close()

# ── due soon: one query behind the count and the list ────────────────────
soon = als.due_soon(PID, today=TODAY)
H.check(all(i['due_date'] <= '2026-09-30' for i in soon) and len(soon) == 2,
        'overdue or due within seven days of 23 Sep: {}'.format(
            [(i['topic'], i['due_date']) for i in soon]))
als.save(PID, None, topic='Electrical safety certificate',
         todo='Updated certificates', due_date='2026-09-01')
H.check([i['overdue'] for i in als.due_soon(PID, today=TODAY)].count(True) == 1,
        'an item past its date is flagged overdue')

# ── the user's own sheet ─────────────────────────────────────────────────
if os.path.isfile(REAL):
    conn = dbm.get_connection()
    conn.execute("INSERT INTO projects (name, num_zones, num_blocks, "
                 "num_containers, project_type) VALUES ('Real', 2, 16, 4, 'BESS')")
    conn.commit()
    RID = conn.execute("SELECT id FROM projects ORDER BY id DESC").fetchone()[0]
    conn.close()
    rpv = als.preview_excel(REAL, sheet='Sheet1')
    H.check(set(rpv['mapping'].values()) >= {'seq', 'topic', 'description',
                                             'todo', 'due_date'},
            'the real sheet is read without help: {}'.format(rpv['mapping']))
    r = als.import_from_excel(RID, REAL, rpv['mapping'], sheet='Sheet1',
                              log=lambda *_a: None)
    real_items = als.items(RID)
    H.check(r['added'] == 15 and len(real_items) == 15,
            '15 real items, the three numbered blanks skipped: {}'.format(r))
    gs = [i for i in real_items if i['topic'] == 'Gas sensors']
    H.check(gs and gs[0]['todo'] == ('Confirm frequency of replacement\n'
                                     'Need code for sensor alone'),
            'with its two-line "to do" intact')
    r2 = als.import_from_excel(RID, REAL, rpv['mapping'], sheet='Sheet1',
                               log=lambda *_a: None)
    H.check(r2['added'] == 0 and r2['updated'] == 0 and r2['unchanged'] == 15
            and len(als.items(RID)) == 15,
            'importing the real sheet twice changes nothing: {}'.format(r2))
else:
    print('   note   the real Book1.xlsx is not on this computer — skipped')

# ── the Plan page's Action list tab ──────────────────────────────────────
from PyQt5.QtCore import Qt                                        # noqa: E402
from PyQt5.QtWidgets import (QApplication, QDialog, QFileDialog,   # noqa: E402
                             QInputDialog, QMessageBox)

app = QApplication.instance() or QApplication([])
import ui.planner_page as pp                                       # noqa: E402

QMessageBox.information = staticmethod(lambda *a, **k: QMessageBox.Ok)
QMessageBox.warning = staticmethod(lambda *a, **k: QMessageBox.Ok)
QMessageBox.question = staticmethod(lambda *a, **k: QMessageBox.Yes)

page = pp.PlannerPage()
page.set_current_project(PID)
tabs = [page.tabs.tabText(i) for i in range(page.tabs.count())]
H.check(tabs[-1] == 'Action list' and tabs[:4] == ['Schedule', 'PM due',
                                                   'Campaigns', 'Checklists'],
        'the Action list is a tab on the Plan page, appended: {}'.format(tabs))
page.tabs.setCurrentIndex(4)
H.check(page.act_tbl.columnCount() == 6
        and [page.act_tbl.horizontalHeaderItem(i).text() for i in range(6)]
        == ['No.', 'Topic', 'To do', 'Due', 'Assigned to', 'Status'],
        'with the columns the user asked for')
H.check(page.act_tbl.rowCount() == len(als.items(PID, include_done=False)),
        'and the open items in it: {} row(s)'.format(page.act_tbl.rowCount()))

# overdue in red
red = [r for r in range(page.act_tbl.rowCount())
       if page.act_tbl.item(r, 1).foreground().color().name() == '#b4232a']
H.check(len(red) == 1, 'the overdue item is red: {} row(s)'.format(len(red)))

# a multi-line "to do" is one line in the table, whole in the tooltip
multi = [r for r in range(page.act_tbl.rowCount())
         if page.act_tbl.item(r, 1).text() == 'PM activities'][0]
H.check('\n' not in page.act_tbl.item(multi, 2).text()
        and '\n' in page.act_tbl.item(multi, 2).toolTip(),
        'the table shows one line per item and keeps the rest in the tooltip')

# add one by hand
pp._ActionItemDialog.exec_ = lambda self: (
    self.topic.setText('Response time'),
    self.todo.setPlainText('Submit list for internal review first'),
    self.has_due.setChecked(False),
    QDialog.Accepted)[-1]
page._add_action()
H.check(any(i['topic'] == 'Response time' for i in als.items(PID)),
        'an item can be added by hand')

# mark one done through the page
QInputDialog.getMultiLineText = staticmethod(
    lambda *a, **k: ('BOM received from HQ', True))
row = [r for r in range(page.act_tbl.rowCount())
       if page.act_tbl.item(r, 1).text() == 'Spare parts of EPC'][0]
page.act_tbl.selectRow(row)
page._done_action()
sp = [i for i in als.items(PID) if i['topic'] == 'Spare parts of EPC'][0]
H.check(sp['status'] == 'done' and sp['done_note'] == 'BOM received from HQ',
        'and marked done with its note: {}'.format(sp['status']))

# export through the page
out2 = os.path.join(H.WORK, 'page export.xlsx')
QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (out2, ''))
page._export_actions()
H.check(os.path.isfile(out2), 'the page exports the list to Excel')

# import through the page, with the mapping dialog answered for us
pp._ActionImportDialog.exec_ = lambda self: (
    self.path.setText(V2),
    self.sheet.addItem('Sheet1', 'Sheet1'),
    self._reread(),
    QDialog.Accepted)[-1]
before_n = len(als.items(PID))
page._import_actions()
H.check(len(als.items(PID)) == before_n,
        'importing the same sheet again through the page adds nothing: '
        '{} → {}'.format(before_n, len(als.items(PID))))
H.check(page.tabs.currentIndex() == 4, 'and it lands on the Action list tab')

# ── the Today panel, and the rule its count lives by ─────────────────────
import services.today_service as ts                                # noqa: E402
import ui.today_page as tp                                         # noqa: E402

rows = ts.actions(PID, today=TODAY)
H.check(rows == als.due_soon(PID, today=TODAY),
        'Today asks the service that owns the table, it does not count twice')

today_page = tp.TodayPage()
today_page.set_current_project(PID, 'TK')
today_page.set_month(2026, 9)
H.check(today_page._grid.count() == 8,
        'the board has eight panels now: {}'.format(today_page._grid.count()))
H.check(today_page.p_actions.count_lbl.text() == str(len(rows)),
        'the Action list panel prints {} for {} row(s)'.format(
            today_page.p_actions.count_lbl.text(), len(rows)))

asked = []
today_page.open_filtered.connect(lambda label, flt: asked.append((label, dict(flt))))
today_page.p_actions.link_btn.click()
H.check(asked and asked[0][0] == 'Plan'
        and asked[0][1] == {'tab': 'actions', 'due': 'soon'},
        'its link asks the shell for the Plan page, filtered: {}'.format(asked))

page.apply_filter(asked[0][1])
H.check(page.tabs.currentIndex() == 4
        and page.act_tbl.rowCount() == len(rows),
        'and that filter opens exactly {} row(s) for a count of {}'.format(
            page.act_tbl.rowCount(), len(rows)))

# the board still fits a narrow window
today_page.show()
for width, want in ((1400, 3), (1000, 2), (700, 1)):
    today_page.resize(width, 800)
    app.processEvents()
    cols = {today_page._grid.getItemPosition(i)[1]
            + today_page._grid.getItemPosition(i)[3]
            for i in range(today_page._grid.count())}
    H.check(today_page._cols == want and max(cols) <= want,
            '{} px wide: {} column(s), nothing past the edge'.format(
                width, today_page._cols))

H.finish()

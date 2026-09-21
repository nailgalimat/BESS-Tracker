"""Two customers, one checklist name — and the rows that reach their files.

Four defects a QA pass found in the new PM-checklist work. Each is reproduced
here the way it happens in the office:

  H4  importing "01) PCS Checklist" for a second project took over the FIRST
      project's template. checklist_templates.name was globally UNIQUE, so the
      row's project, its copy of the workbook and its items all moved across;
      the Tashkent checklists already filled in pointed at Bukhara's workbook,
      and the export wrote into the other customer's file.
  M3  an item added after a campaign was planned never reached the phones: the
      template's updated_at did not move, so the sync had no reason to upload
      the items again, and the runs already planned were not queued either.
  M5  two items added to the same group came out in the customer's file in
      reverse order.
  L1  a comment longer than 32767 characters makes Excel call the file damaged
      and offer to repair it — for the customer, the checklist did not arrive.

Everything here builds its own workbook, so it runs on a computer that does
not have the customer's files.
"""
import _harness as H            # must be first
import os
import sqlite3
import time

import openpyxl

import database.db_manager as dbm
import services.checklist_excel as cx
import services.checklist_pm_service as cs

NAME = '01) PCS Checklist'


def make_workbook(path):
    """A checklist in the layout the customer's own files use: header fields
    on rows 4-6, headings on row 7, items from row 8, then Progress."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = 'Preventive Maintenance Checklist'
    for ref in ('C4', 'E4', 'C5', 'E5', 'C6', 'E6'):
        ws[ref] = '-'
    for col, head in zip('ABCDEFG', ('S.No', 'Equipment', 'Activity',
                                     'Description', 'Status', 'Verified / Done',
                                     'Comments')):
        ws[col + '7'] = head
    for n, (no, eq, act, text) in enumerate((
            ('1', 'PCS', 'Visual', 'Check the cabinet door seals'),
            ('2', 'PCS', 'Visual', 'Check the cooling fans'),
            ('3', 'RMU', 'Visual', 'Check the RMU'))):
        r = 8 + n
        ws['A%d' % r], ws['B%d' % r] = no, eq
        ws['C%d' % r], ws['D%d' % r] = act, text
    ws['B11'] = 'Progress'
    wb.save(path)
    return path


SRC = make_workbook(os.path.join(H.WORK, NAME + '.xlsx'))


def template_row(tid):
    conn = dbm.get_connection()
    try:
        r = conn.execute("SELECT * FROM checklist_templates WHERE id=?",
                         (tid,)).fetchone()
        return dict(r) if r else None
    finally:
        conn.close()


# ── H4: a template belongs to a project ──────────────────────────────────
H.fresh_db('cl_scope.db')
conn = dbm.get_connection()
for nm in ('Tashkent', 'Bukhara'):
    conn.execute("INSERT INTO projects (name, num_zones, num_blocks, "
                 "num_containers, project_type) VALUES (?,1,8,4,'BESS')", (nm,))
TK, BK = [r[0] for r in conn.execute("SELECT id FROM projects ORDER BY id")]
conn.commit(); conn.close()

tk_tid = cs.import_template(TK, SRC)
run_uuid = cs.plan_runs(TK, tk_tid, [3], '2026-09-21', campaign='PM Sep 2026')[0]
tk_items = cs.items(tk_tid)
cs.save_results(run_uuid, {tk_items[1]['id']: {'result': cs.NOK,
                                               'comment': 'fan noisy'}})

bk_tid = cs.import_template(BK, SRC)
H.check(bk_tid != tk_tid,
        'the second plant gets its own template row ({} / {})'.format(tk_tid, bk_tid))
H.check(template_row(tk_tid)['project_id'] == TK,
        "Tashkent's template still belongs to Tashkent "
        "(it used to be taken over by whoever imported last)")
H.check(template_row(bk_tid)['project_id'] == BK,
        "and Bukhara's to Bukhara")
H.check([t['id'] for t in cs.templates(TK)] == [tk_tid]
        and [t['id'] for t in cs.templates(BK)] == [bk_tid],
        'each plant sees one checklist, its own')
detail = cs.run_detail(run_uuid)
H.check(detail['template_id'] == tk_tid
        and [i['item_id'] for i in detail['items']] == [i['id'] for i in tk_items],
        'the Tashkent checklist already filled in still points at its own items')
nok = [i for i in detail['items'] if i['result'] == cs.NOK]
H.check(len(nok) == 1 and 'fan noisy' in nok[0]['comment'],
        'and what was written on it is still there')

conn = dbm.get_connection()
try:
    conn.execute("INSERT INTO checklist_templates (name, project_id) VALUES (?,?)",
                 (NAME, TK))
    conn.commit()
    dup = 'accepted'
except sqlite3.IntegrityError:
    dup = 'refused'
finally:
    conn.close()
H.check(dup == 'refused',
        'the database itself refuses a second "{}" on one plant'.format(NAME))

# ── H4: an existing database is migrated, not wiped ──────────────────────
# The old table carried `name TEXT NOT NULL UNIQUE`. SQLite cannot drop a
# column constraint, so initialize_database rebuilds the table once — and the
# rows that are already in it have to come through unchanged.
legacy = os.path.join(H.WORK, 'legacy.db')
for s in ('', '-wal', '-shm'):
    try:
        os.remove(legacy + s)
    except OSError:
        pass
lc = sqlite3.connect(legacy)
lc.executescript("""
    CREATE TABLE checklist_templates (
        id             INTEGER PRIMARY KEY AUTOINCREMENT,
        name           TEXT    NOT NULL UNIQUE,
        description    TEXT    DEFAULT '',
        container_type TEXT    DEFAULT '',
        scope          TEXT    DEFAULT 'container',
        created_at     TEXT    DEFAULT (datetime('now'))
    );
    CREATE TABLE checklist_items (
        id          INTEGER PRIMARY KEY AUTOINCREMENT,
        template_id INTEGER NOT NULL,
        order_num   INTEGER NOT NULL DEFAULT 0,
        category    TEXT    DEFAULT '',
        description TEXT    NOT NULL,
        expected    TEXT    DEFAULT '',
        FOREIGN KEY (template_id) REFERENCES checklist_templates(id) ON DELETE CASCADE
    );
    INSERT INTO checklist_templates (name, description, scope)
        VALUES ('01) PCS Checklist', 'the one that was already here', 'block');
    INSERT INTO checklist_items (template_id, order_num, description)
        VALUES (1, 0, 'Check the cabinet door seals');
""")
lc.commit(); lc.close()

dbm.DB_PATH = legacy
dbm.initialize_database()
conn = dbm.get_connection()
kept = conn.execute("SELECT * FROM checklist_templates").fetchall()
items_kept = conn.execute("SELECT * FROM checklist_items").fetchall()
uniq = []
for idx in conn.execute("PRAGMA index_list(checklist_templates)").fetchall():
    if idx['origin'] == 'u':
        uniq.append([r[2] for r in conn.execute(
            "PRAGMA index_info(%r)" % idx['name']).fetchall()])
H.check(len(kept) == 1 and kept[0]['id'] == 1
        and kept[0]['description'] == 'the one that was already here'
        and len(items_kept) == 1,
        'the template and its item survive the rebuild, under the same ids')
H.check(['name'] not in uniq,
        'and the global UNIQUE on the name is gone: {}'.format(uniq))
try:
    conn.execute("INSERT INTO checklist_templates (name, project_id) VALUES (?,2)",
                 (NAME,))
    conn.commit()
    second = 'accepted'
except sqlite3.IntegrityError:
    second = 'refused'
H.check(second == 'accepted',
        'a second plant may now hold a checklist with the same name')
conn.close()
dbm.initialize_database()          # forward-only: running it again is a no-op
conn = dbm.get_connection()
H.check(conn.execute("SELECT COUNT(*) FROM checklist_templates").fetchone()[0] == 2,
        'and the migration does not run a second time')
conn.close()

# ── M3: an item added after the campaign was planned reaches the phones ──
dbm.DB_PATH = os.path.join(H.WORK, 'cl_scope.db')
uuids = cs.plan_runs(TK, tk_tid, [5, 6], '2026-09-22', campaign='PM Sep 2026')
for u in uuids:
    cs.mark_run_synced(u)


def run_state(u):
    conn = dbm.get_connection()
    try:
        r = conn.execute("SELECT sync_status, updated_at FROM checklist_runs "
                         "WHERE uuid=?", (u,)).fetchone()
        return dict(r)
    finally:
        conn.close()


before_tpl = template_row(tk_tid)['updated_at']
before_runs = {u: run_state(u) for u in uuids}
time.sleep(1.1)                      # the stamps are per second
added_id = cs.add_item(tk_tid, 'Extra check the customer asked for',
                       equipment='PCS')
after_tpl = template_row(tk_tid)['updated_at']
after_runs = {u: run_state(u) for u in uuids}
H.check(after_tpl > before_tpl,
        'adding an item moves the template stamp, so the sync uploads the '
        'items again ({} -> {})'.format(before_tpl, after_tpl))
H.check(all(after_runs[u]['sync_status'] == 'pending' for u in uuids),
        'and the checklists already planned are queued again: {}'.format(
            [after_runs[u]['sync_status'] for u in uuids]))
H.check(all(after_runs[u]['updated_at'] == before_runs[u]['updated_at']
            for u in uuids),
        'but their own stamp does not move — a phone that has already filled '
        'one in must not lose its answers to an added line')
H.check({r['uuid'] for r in cs.runs_for_sync(TK)} >= set(uuids),
        'the sync picks them up')

for u in uuids:
    cs.mark_run_synced(u)
time.sleep(1.1)
cs.delete_item(added_id)
H.check(template_row(tk_tid)['updated_at'] > after_tpl
        and all(run_state(u)['sync_status'] == 'pending' for u in uuids),
        'taking the item away again says so too')

# ── M5 / L1: what goes into the customer's file ──────────────────────────
out = os.path.join(H.WORK, 'filled.xlsx')
header = {'plant': 'ACWA Tashkent BESS', 'date': '21.09.2026'}
added = [{'after_row': 10, 'no': '3a', 'equipment': 'RMU',
          'text': 'First added item', 'result': cx.OK, 'comment': ''},
         {'after_row': 10, 'no': '3b', 'equipment': 'RMU',
          'text': 'Second added item', 'result': cx.OK, 'comment': ''}]
cx.write_filled(SRC, out, header, {8: {'result': cx.OK, 'comment': 'fine'}}, added)
wb = openpyxl.load_workbook(out)
ws = wb.active
H.check(ws['D11'].value == 'First added item'
        and ws['D12'].value == 'Second added item',
        'two items added to the same group keep their order in the customer\'s '
        'file: {} then {}'.format(ws['D11'].value, ws['D12'].value))
H.check(ws['C4'].value == 'ACWA Tashkent BESS' and ws['E8'].value is True,
        'the header and the ticked box are still written')
wb.close()

out2 = os.path.join(H.WORK, 'filled-long.xlsx')
cx.write_filled(SRC, out2, header,
                {8: {'result': cx.NOK, 'comment': 'x' * 40000}})
wb = openpyxl.load_workbook(out2)
ws = wb.active
val = ws['G8'].value or ''
H.check(len(val) <= 32767 and val.startswith('x'),
        'a 40 000-character comment is cut to what Excel can hold ({} chars) '
        '— a longer one makes Excel call the file damaged'.format(len(val)))
wb.close()

H.finish()

"""Excel stock import: the SAP file the owner is sent, and the round trip.

The owner receives stock snapshots as SAP MB52 exports — every cell padded with
spaces, four junk rows above the header — and used to type them in by hand.

The regression that matters most is in here: `min_quantity` is the low-stock
alert threshold the owner set by hand, and `set_stock_level`'s `min_quantity=0`
default would zero every alert in the warehouse on the first import. A
re-import must leave those thresholds exactly as they were.

Workbooks are built here with openpyxl. The real file is customer data and is
never copied into the repository, so nothing in this test reads it.
"""
import os
import _harness as H            # must be first

import openpyxl

import services.stock_excel_service as sx
import services.stock_service as ss
import services.material_service as ms

H.fresh_db('stock_excel.db')
WH = ss.get_main_warehouse()["id"]


def sap_file(name, rows, header=('Plnt', 'Location', 'BUn', 'Material',
                                 'Material Description', 'Description',
                                 'Unrestr.', 'Trans./Tfr')):
    """A workbook shaped like the owner's SAP export.

    Column A empty, four junk rows with a date stamp, the header on row 5, row 6
    blank, data from row 7 — and every cell left-padded with spaces the way the
    real dump lines its columns up visually.
    """
    path = os.path.join(H.WORK, name)
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Sheet1'
    sheet.cell(row=2, column=1,
               value='2026.09.22         Dynamic List Display            1')
    for col, text in enumerate(header, start=2):          # B..I — A stays empty
        sheet.cell(row=5, column=col, value=' ' * (col * 5) + str(text))
    for offset, row in enumerate(rows):
        for col, value in enumerate(row, start=2):
            if value is None:
                continue
            sheet.cell(row=7 + offset, column=col,
                       value=' ' * (col * 7) + str(value))
    book.save(path)
    return path


def template_file(name, rows, headers=('Material #', 'Description', 'Unit', 'Qty')):
    """A workbook in the app's own layout — header on row 1, no padding."""
    path = os.path.join(H.WORK, name)
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Current Stock'
    for col, text in enumerate(headers, start=1):
        sheet.cell(row=1, column=col, value=text)
    for offset, row in enumerate(rows):
        for col, value in enumerate(row, start=1):
            sheet.cell(row=2 + offset, column=col, value=value)
    book.save(path)
    return path


# ── 1. the SAP layout, padding and junk rows and all ─────────────────────────
sap = sap_file('sap.xlsx', [
    ('2209', 'UZ21', 'PCS', '6SW00366', 'Widget A', 'UZ ACWA-Tash Ava', '4', '0'),
    ('2209', 'UZ21', 'PCS', '6SW02777', 'Widget B', 'UZ ACWA-Tash Ava', '1', '3'),
    ('2209', 'UZ21', 'KG', 'ZE000135', 'Grease C', 'UZ ACWA-Tash Ava', '200', '0'),
])
p = sx.parse(sap)
H.check(p['layout'] == 'sap', "the SAP layout is recognised: {}".format(p['layout']))
H.check(p['header_row'] == 5,
        'the header row is found at 5, under the junk rows: {}'.format(p['header_row']))
H.check(len(p['rows']) == 3, 'three data rows read: {}'.format(len(p['rows'])))
H.check([r['material_number'] for r in p['rows']]
        == ['6SW00366', '6SW02777', 'ZE000135'],
        'material numbers come back with the padding stripped: {}'.format(
            [r['material_number'] for r in p['rows']]))
H.check([r['quantity'] for r in p['rows']] == [4.0, 1.0, 200.0],
        'padded text quantities are parsed: {}'.format(
            [r['quantity'] for r in p['rows']]))
H.check(p['rows'][0]['description'] == 'Widget A' and p['rows'][2]['unit'] == 'KG',
        'the material description (column F, not the location one in G) and the '
        'unit are stripped: {!r} / {!r}'.format(p['rows'][0]['description'],
                                                p['rows'][2]['unit']))
H.check(p['plants'] == ['2209'] and p['locations'] == ['UZ21'],
        'the plant and location the file names are reported for the preview: '
        '{} / {}'.format(p['plants'], p['locations']))
H.check(p['rows'][1]['in_transfer'] == 3.0,
        'stock in transfer is carried but is not the counted quantity')
H.check(not p['problems'], 'a clean SAP file has no problems: {}'.format(p['problems']))

# ── 2. plan classifies against what is on the shelf ──────────────────────────
ss.set_stock_level(WH, '6SW00366', quantity=4, min_quantity=2, unit='PCS')
ss.set_stock_level(WH, '6SW02777', quantity=9, min_quantity=5, unit='PCS')
planned = sx.plan(WH, p['rows'])
by_material = {r['material_number']: r for r in planned}
H.check(by_material['6SW00366']['status'] == 'unchanged',
        'a material already at the counted quantity is unchanged')
H.check(by_material['6SW02777']['status'] == 'changed'
        and by_material['6SW02777']['before'] == 9.0
        and by_material['6SW02777']['after'] == 1.0
        and by_material['6SW02777']['delta'] == -8.0,
        'a different quantity is changed, 9 → 1 (delta {})'.format(
            by_material['6SW02777']['delta']))
H.check(by_material['ZE000135']['status'] == 'new',
        'a material not stocked here yet is new')
H.check(sx.summarise(planned) == {'new': 1, 'changed': 1, 'unchanged': 1,
                                 'thresholds': 0, 'total': 3},
        'the summary counts them: {}'.format(sx.summarise(planned)))
H.check(p['has_min_column'] is False
        and all(r['min_quantity'] is None for r in p['rows']),
        'a SAP file carries no Min Qty column, so it has no opinion on any '
        'threshold')
H.check(ss.get_stock_quantity(WH, '6SW02777')['quantity'] == 9.0,
        'plan() wrote nothing — the shelf still says 9')

# ── 3. THE regression: min_quantity survives a re-import ─────────────────────
before_thresholds = {r['material_number']: r['min_quantity']
                     for r in ss.get_stock(WH)}
counts = sx.apply(WH, planned, reference='IMPORT sap.xlsx',
                  transaction_date='2026-09-22')
after_thresholds = {r['material_number']: r['min_quantity']
                    for r in ss.get_stock(WH)}
H.check(after_thresholds['6SW00366'] == 2.0 and after_thresholds['6SW02777'] == 5.0,
        'the hand-set thresholds are exactly as they were: 6SW00366={}, '
        '6SW02777={}'.format(after_thresholds['6SW00366'],
                             after_thresholds['6SW02777']))
H.check(all(after_thresholds[m] == v for m, v in before_thresholds.items()),
        'no threshold in the warehouse moved at all')
H.check(after_thresholds['ZE000135'] == 0.0,
        'a brand-new item starts with no threshold, for the owner to set')

# and the quantities did land
H.check(ss.get_stock_quantity(WH, '6SW02777')['quantity'] == 1.0,
        'the counted quantity is on the shelf: {}'.format(
            ss.get_stock_quantity(WH, '6SW02777')['quantity']))
H.check(ss.get_stock_quantity(WH, 'ZE000135')['quantity'] == 200.0,
        'the new material is stocked at 200')
H.check(ss.get_stock_quantity(WH, 'ZE000135')['unit'] == 'KG',
        'and carries its unit of measure')
H.check(counts['created'] == 1 and counts['updated'] == 1
        and counts['unchanged'] == 1,
        'apply reports 1 created / 1 updated / 1 unchanged: {}'.format(
            {k: counts[k] for k in ('created', 'updated', 'unchanged')}))

# ── 4. the delta is an IN/OUT transaction that reconciles ────────────────────
tx = ss.get_transactions(warehouse_id=WH, limit=50)
imported = [t for t in tx if t['reference'] == 'IMPORT sap.xlsx']
H.check(len(imported) == 2,
        'two transactions — the unchanged row moves nothing: {}'.format(len(imported)))
H.check(all(t['transaction_type'] in ('IN', 'OUT') for t in imported),
        'only the existing IN/OUT types are used, no fourth type: {}'.format(
            sorted({t['transaction_type'] for t in imported})))
out = [t for t in imported if t['material_number'] == '6SW02777'][0]
H.check(out['transaction_type'] == 'OUT' and out['quantity'] == 8.0,
        'the drop from 9 to 1 is an OUT of 8: {} {}'.format(
            out['transaction_type'], out['quantity']))
into = [t for t in imported if t['material_number'] == 'ZE000135'][0]
H.check(into['transaction_type'] == 'IN' and into['quantity'] == 200.0,
        'the new material arrives as an IN of 200')
H.check(9.0 - out['quantity'] == ss.get_stock_quantity(WH, '6SW02777')['quantity'],
        'the audit trail reconciles: 9 - 8 = the quantity now on the shelf')

# ── 5. a description already in the catalogue is not blanked ────────────────
ms.save_material('6SW00366', 'Hand-written description', 'PCS', 'keep my notes')
blank_desc = sap_file('sap_blank_desc.xlsx', [
    ('2209', 'UZ21', 'PCS', '6SW00366', None, 'UZ ACWA-Tash Ava', '4', '0'),
])
sx.apply(WH, sx.plan(WH, sx.parse(blank_desc)['rows']), reference='IMPORT blank')
kept = ms.get_material('6SW00366')
H.check(kept['description'] == 'Hand-written description',
        'an empty description cell did not blank the catalogue: {!r}'.format(
            kept['description']))
H.check(kept['notes'] == 'keep my notes', 'and the notes are untouched')
H.check((ms.get_material('ZE000135') or {}).get('description') == 'Grease C',
        "the SAP file's description filled a material the app had none for")

# ── 6. an unreadable quantity is a problem, not a zero ──────────────────────
bad = sap_file('sap_bad_qty.xlsx', [
    ('2209', 'UZ21', 'PCS', 'BAD-QTY-1', 'Widget D', 'UZ', 'twelve', '0'),
    ('2209', 'UZ21', 'PCS', 'BAD-QTY-2', 'Widget E', 'UZ', None, '0'),
    ('2209', 'UZ21', 'PCS', 'GOOD-1', 'Widget F', 'UZ', '1 500', '0'),
    ('2209', 'UZ21', 'PCS', 'GOOD-1', 'Widget F again', 'UZ', '7', '0'),
    ('2209', 'UZ21', 'PCS', None, 'Row with no material', 'UZ', '5', '0'),
])
pb = sx.parse(bad)
kinds = sorted({x['kind'] for x in pb['problems']})
H.check([r['material_number'] for r in pb['rows']] == ['GOOD-1'],
        'only the readable row is offered for import: {}'.format(
            [r['material_number'] for r in pb['rows']]))
H.check(pb['rows'][0]['quantity'] == 1500.0,
        'a padded thousands separator still reads as 1500: {}'.format(
            pb['rows'][0]['quantity']))
H.check(kinds == ['duplicate', 'no_material', 'quantity'],
        'unreadable quantity, duplicate and missing material are all reported: '
        '{}'.format(kinds))
H.check(len([x for x in pb['problems'] if x['kind'] == 'quantity']) == 2,
        'both the word and the empty cell are quantity problems, not zeros')
H.check(ss.get_stock_quantity(WH, 'BAD-QTY-1') is None,
        'nothing about the bad rows reached the database')

# ── 7. no header row at all is a clear error ────────────────────────────────
junk = os.path.join(H.WORK, 'junk.xlsx')
_book = openpyxl.Workbook()
_book.active['A1'] = 'this is not a stock list'
_book.active['A2'] = 'holiday roster'
_book.save(junk)
try:
    sx.parse(junk)
    H.check(False, 'a file with no header row must be rejected')
except ValueError as e:
    H.check('Material' in str(e),
            'the error says what was expected: {!r}'.format(str(e)[:70]))

# ── 8. round trip: the app's own export re-imports ──────────────────────────
exported = os.path.join(H.WORK, 'export.xlsx')
ss.export_stock_excel(exported)
pe = sx.parse(exported)
H.check(pe['layout'] == 'template' and pe['sheet'] == 'Current Stock',
        "export_stock_excel's own sheet is read, not the Transactions one: "
        '{} / {}'.format(pe['layout'], pe['sheet']))
live = {r['material_number']: r['quantity'] for r in ss.get_stock(WH)}
round_tripped = {r['material_number']: r['quantity'] for r in pe['rows']}
H.check(all(round_tripped.get(m) == q for m, q in live.items()),
        'every stocked material comes back at the same quantity: {} of {}'.format(
            len(round_tripped), len(live)))
H.check(all(r['status'] == 'unchanged' for r in sx.plan(WH, pe['rows'])),
        'so re-importing the export changes nothing')
H.check(pe['has_min_column'] and pe['ignored_columns'] == ['Low Stock Alert'],
        "the export's own Min Qty column is an input; only the computed Low "
        'Stock Alert flag is ignored: {}'.format(pe['ignored_columns']))
H.check(all(r['min_quantity'] == 2.0 for r in pe['rows']
            if r['material_number'] == '6SW00366'),
        'and the export carries the thresholds it printed, so a round trip '
        'restates them rather than zeroing them')

# ── 8b. Min Qty: absent column, empty cell, filled cell, unreadable ────────
# The rule this codebase states in sync_client._ENTRY_OPT_COLS: an absent field
# means "no opinion", a present one means "apply it".
MINQ = ('Material #', 'Description', 'Unit', 'Qty', 'Min Qty')
ss.set_stock_level(WH, 'MIN-1', quantity=10, min_quantity=5, unit='PCS')

# (a) column present, cell empty → the owner's threshold is left alone
empty_cell = template_file('min_empty.xlsx',
                           [('MIN-1', 'Widget', 'PCS', 7, None)], headers=MINQ)
pm = sx.parse(empty_cell)
H.check(pm['has_min_column'] and pm['rows'][0]['min_quantity'] is None,
        'an empty Min Qty cell is "no opinion", not zero: {!r}'.format(
            pm['rows'][0]['min_quantity']))
plan_empty = sx.plan(WH, pm['rows'])
H.check(plan_empty[0]['min_changed'] is False
        and plan_empty[0]['min_after'] == 5.0,
        'so the plan shows the threshold staying at 5')
counts_empty = sx.apply(WH, plan_empty, reference='IMPORT min_empty')
kept = {r['material_number']: r for r in ss.get_stock(WH)}['MIN-1']
H.check(kept['min_quantity'] == 5.0 and kept['quantity'] == 7.0,
        'the quantity is applied (7) and the hand-set threshold is untouched '
        '(5): qty={} min={}'.format(kept['quantity'], kept['min_quantity']))
H.check(counts_empty['thresholds_set'] == 0,
        'and no threshold write is reported: {}'.format(
            counts_empty['thresholds_set']))

# (b) column present, cell filled → apply it
filled = template_file('min_filled.xlsx',
                       [('MIN-1', 'Widget', 'PCS', 7, 12)], headers=MINQ)
plan_filled = sx.plan(WH, sx.parse(filled)['rows'])
H.check(plan_filled[0]['min_changed'] is True
        and plan_filled[0]['min_before'] == 5.0
        and plan_filled[0]['min_after'] == 12.0,
        'a filled Min Qty cell plans a threshold change, 5 → 12')
H.check(plan_filled[0]['status'] == 'changed'
        and sx.summarise(plan_filled)['thresholds'] == 1,
        'a file that moves only the threshold is a real change, not an '
        '"all unchanged" no-op: status={}'.format(plan_filled[0]['status']))
counts_filled = sx.apply(WH, plan_filled, reference='IMPORT min_filled')
now = {r['material_number']: r for r in ss.get_stock(WH)}['MIN-1']
H.check(now['min_quantity'] == 12.0,
        'and the threshold really is 12 now: {}'.format(now['min_quantity']))
H.check(counts_filled['thresholds_set'] == 1
        and counts_filled['updated'] == 1
        and counts_filled['transactions'] == 0,
        'counted as one threshold set and one row updated, with no stock '
        'transaction — nothing moved on the shelf: {}'.format(
            {k: counts_filled[k] for k in
             ('thresholds_set', 'updated', 'transactions')}))
H.check(sx.apply(WH, sx.plan(WH, sx.parse(filled)['rows']),
                 reference='IMPORT again')['thresholds_set'] == 0,
        'importing the same file twice reports no second threshold change')

# (c) a negative or unreadable Min Qty is a problem, and costs the row nothing
bad_min = template_file('min_bad.xlsx', [
    ('MIN-1', 'Widget', 'PCS', 9, -3),
    ('MIN-2', 'Other widget', 'PCS', 4, 'soon'),
], headers=MINQ)
pbm = sx.parse(bad_min)
H.check([x['kind'] for x in pbm['problems']] == ['min_quantity', 'min_quantity'],
        'both a negative and an unreadable threshold are reported: {}'.format(
            [x['kind'] for x in pbm['problems']]))
H.check(len(pbm['rows']) == 2
        and all(r['min_quantity'] is None for r in pbm['rows']),
        'the rows are still importable — a bad threshold does not cost a row '
        'its quantity: {} row(s)'.format(len(pbm['rows'])))
sx.apply(WH, sx.plan(WH, pbm['rows']), reference='IMPORT min_bad')
after_bad = {r['material_number']: r for r in ss.get_stock(WH)}
H.check(after_bad['MIN-1']['min_quantity'] == 12.0
        and after_bad['MIN-1']['quantity'] == 9.0,
        'the quantity moved to 9 and the threshold stayed at 12: {}'.format(
            after_bad['MIN-1']['min_quantity']))
H.check(after_bad['MIN-2']['quantity'] == 4.0
        and after_bad['MIN-2']['min_quantity'] == 0.0,
        'and the unreadable one imported with no threshold at all')

# (d) the stale-plan guard: apply must use what the FILE said, not what the
# plan saw, or it would write a threshold back over one edited since.
stale_plan = sx.plan(WH, sx.parse(empty_cell)['rows'])      # file: no opinion
ss.set_stock_level(WH, 'MIN-1', quantity=9, min_quantity=44, unit='PCS')
sx.apply(WH, stale_plan, reference='IMPORT stale')
H.check({r['material_number']: r for r in ss.get_stock(WH)}
        ['MIN-1']['min_quantity'] == 44.0,
        'a threshold edited after the preview was built survives the import')

# ── 9. round trip: the downloaded template imports ──────────────────────────
tpl = os.path.join(H.WORK, 'template.xlsx')
sx.write_template(tpl)
pt = sx.parse(tpl)
H.check(pt['layout'] == 'template' and len(pt['rows']) == 2,
        'the template parses, instruction row and all, with its two example '
        'rows: header {} / {} row(s)'.format(pt['header_row'], len(pt['rows'])))
H.check(all(r['material_number'].startswith('EXAMPLE-') for r in pt['rows']),
        'and its example rows are obviously fake: {}'.format(
            [r['material_number'] for r in pt['rows']]))
H.check(pt['has_min_column']
        and [r['min_quantity'] for r in pt['rows']] == [2.0, None],
        'the template offers Min Qty and shows both ways to use it — one row '
        'sets a threshold, one leaves it to the app: {}'.format(
            [r['min_quantity'] for r in pt['rows']]))

# ── 10. the preview dialog: nothing is written until it is confirmed ────────
from PyQt5.QtWidgets import QApplication, QDialog, QFileDialog, QMessageBox
from PyQt5.QtCore import Qt
app = QApplication.instance() or QApplication([])
import ui.stock_page as sp

count_file = template_file('count.xlsx', [
    ('DLG-NEW-1', 'Dialog new part', 'PCS', 6),
    ('DLG-NEW-2', 'Dialog second part', 'PCS', 2),
    ('6SW00366', 'Widget A', 'PCS', 4),          # unchanged
])
parsed = sx.parse(count_file)
dlg = sx.plan(WH, parsed['rows'])
preview = sp.StockImportDialog('Main Warehouse', parsed, dlg)
H.check(preview.table.rowCount() == 3,
        'the preview lists all three rows: {}'.format(preview.table.rowCount()))
H.check(len(preview.selected_rows()) == 3, 'all rows start ticked')
preview._select_changes()
H.check([r['material_number'] for r in preview.selected_rows()]
        == ['DLG-NEW-1', 'DLG-NEW-2'],
        '"Only new / changed" drops the unchanged row: {}'.format(
            [r['material_number'] for r in preview.selected_rows()]))
preview.table.item(1, 0).setCheckState(Qt.Unchecked)
H.check([r['material_number'] for r in preview.selected_rows()] == ['DLG-NEW-1'],
        'a row can be excluded on its own')
H.check('Plant' not in preview.source_label.text(),
        'an app-layout file names no plant, so the preview says so')

# the page's own path: cancelled preview writes nothing, confirmed one writes
page = sp.StockViewTab()
page.wh_combo.setCurrentIndex(page.wh_combo.findData(WH))
QFileDialog.getOpenFileName = staticmethod(lambda *a, **k: (count_file, ''))
QMessageBox.information = staticmethod(lambda *a, **k: QMessageBox.Ok)
QMessageBox.critical = staticmethod(lambda *a, **k: QMessageBox.Ok)

sp.StockImportDialog.exec_ = lambda self: QDialog.Rejected
page._import_excel()
H.check(ss.get_stock_quantity(WH, 'DLG-NEW-1') is None,
        'a cancelled preview wrote nothing at all')

sp.StockImportDialog.exec_ = lambda self: QDialog.Accepted
page._import_excel()
H.check(ss.get_stock_quantity(WH, 'DLG-NEW-1')['quantity'] == 6.0
        and ss.get_stock_quantity(WH, 'DLG-NEW-2')['quantity'] == 2.0,
        'confirming writes the expected rows: {} / {}'.format(
            ss.get_stock_quantity(WH, 'DLG-NEW-1'),
            ss.get_stock_quantity(WH, 'DLG-NEW-2')))
H.check({r['material_number']: r['min_quantity'] for r in ss.get_stock(WH)}
        ['6SW00366'] == 2.0,
        'and the threshold on the unchanged row is still the owner’s 2')
H.check(any(t['reference'].startswith('IMPORT count.xlsx')
            for t in ss.get_transactions(warehouse_id=WH, limit=50)),
        'the transactions name the import file they came from')

# ── 11. one material number, one row, whatever case it was typed in ────────
# stock_items is unique on (warehouse_id, material_number), so a hand-typed
# lower-case number would sit beside the imported upper-case one as a second
# row for the same part, with nothing to catch it.
ss.set_stock_level(WH, 'case-test-1', quantity=3, min_quantity=1, unit='PCS')
rows = [r for r in ss.get_stock(WH)
        if r['material_number'].upper() == 'CASE-TEST-1']
H.check(len(rows) == 1 and rows[0]['material_number'] == 'CASE-TEST-1',
        'set_stock_level stores the material number upper-cased, like the rest '
        'of the app: {}'.format([r['material_number'] for r in rows]))
ss.set_stock_level(WH, 'CASE-TEST-1', quantity=8, min_quantity=1, unit='PCS')
rows = [r for r in ss.get_stock(WH)
        if r['material_number'].upper() == 'CASE-TEST-1']
H.check(len(rows) == 1 and rows[0]['quantity'] == 8.0,
        'so typing it in either case updates the one row instead of making a '
        'second: {} row(s)'.format(len(rows)))
case_file = template_file('case.xlsx', [('case-test-1', 'Widget', 'PCS', 2)])
sx.apply(WH, sx.plan(WH, sx.parse(case_file)['rows']), reference='IMPORT case')
rows = [r for r in ss.get_stock(WH)
        if r['material_number'].upper() == 'CASE-TEST-1']
H.check(len(rows) == 1 and rows[0]['quantity'] == 2.0
        and rows[0]['min_quantity'] == 1.0,
        'and an import of the lower-case spelling lands on that same row, '
        'threshold intact: {} row(s), min={}'.format(
            len(rows), rows[0]['min_quantity']))

H.finish()

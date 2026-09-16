"""Generate the real Tashkent August 2026 report from the SCADA exports and read
it back — once straight from the export folder, then again through the month
data set as report versions v1 and v2, with their manifests.

Frame-level tests cannot catch a crash in the real pipeline (one shipped on
2026-09-10: a frame without Deactivation reached build_event_type_table). Only
generating the report does. Several minutes; run with `tests\\run.py --real`.

GOLDEN holds customer-visible numbers that have been accepted. When one changes
on purpose, update it here in the same commit — that is the approval.

Frozen inputs (QA G12). The golden is computed from August's availability
inputs as they were frozen — exclusions, manual rows, PM, balancing, and the KPI
history the cycles-in-year figure uses — not from whatever the live database
holds that day. They are site operational data, so they are never in the repo:
the file is `frozen_report_inputs_2026-08.json` in the August SCADA folder
(customer data, not committed), or the file named by BESS_FROZEN_INPUTS — which
may also be the manifest of the version sent to the customer, it carries the
same `report_inputs` and `history_records`. When the file does not exist the
test writes it once from the snapshot's report_inputs(1, 2026, 8). Every run
also checks that the snapshot still gives the same inputs: after August is
marked sent they can only change through "Unlock month".
"""
import _harness as H            # must be first: isolates DB, sync, network, history
import copy
import datetime
import hashlib
import json
import os
import re
import time
import warnings
warnings.filterwarnings('ignore')

GOLDEN = {
    # contractual method, with every availability input the Monthly Reports page
    # passes. Was 99.54 until 2026-09-11/12, when August's exclusion windows were
    # corrected against the SCADA data — starts moved to when the plant actually
    # went down, restoration of the blocks that did not come back added, and the
    # KKS windows stretched to the stop they covered. Both rounds user-approved;
    # exclusion_edge_report is what found them. 99.81 -> 99.85 on 2026-09-14: KKS
    # windows widened from the morning isolation to the restart, and 04.08
    # 11:09-17:20 blocks 14-16 recorded as the planned switching it was.
    # 99.85 -> 99.81 the same day: PM counts as unavailability (4 planned hours
    # per block, the rest of the stop under the KKS window — user-confirmed), and
    # the six PM rows dated 1st-12th had been read as January-December and
    # dropped. 24 PM x 4 h x 4 PCS = 384 unit-h + 18.4 fault = 402.4.
    # 99.81 -> 99.80 the same day: whole-container BESS faults now count
    # (bess_container_fault_episodes) — block 28 LC2, 30.08 07:30-14:20, one PCS
    # unit idle while the block charged, LC status RUNNING throughout.
    'availability_pct': 99.80,
    # calendar year to date: counter at 31 Aug minus counter at 1 Mar, plus the
    # January and February records (snapshots begin in March 2026)
    'cycles_in_year': 249.6,
    'cycles_month': 26.6,          # CMU counter delta, August
    'cycles_lifetime': 256.4,      # CMU counter, 31 August
}

SRC = H.require_scada('August', 'SCADA Raw data')


def f(name):
    return os.path.join(SRC, name)


def _digest(p):
    return hashlib.sha256(open(p, 'rb').read()).hexdigest() if os.path.exists(p) else None


real_hist = H.REAL_HISTORY['DEFAULT_HISTORY_PATH']
hist_before = _digest(real_hist)

import services.report_workflow_service as rw                      # noqa: E402
from services.tashkent_report_service import generate_tashkent_report  # noqa: E402
from services.asset_tree_service import UMBRELLA_FOOTNOTE             # noqa: E402

import services.bukhara_report_service as B                        # noqa: E402

# Exactly what the Monthly Reports page passes — capacities, exclusions, manual
# downtime, PM as downtime, balancing — so the figures checked are the ones the
# customer receives; frozen (see the docstring).
FROZEN = os.environ.get('BESS_FROZEN_INPUTS') or os.path.join(
    H.SCADA_DIR, 'August', 'frozen_report_inputs_2026-08.json')
LIVE = json.loads(json.dumps(rw.report_inputs(1, 2026, 8), default=str))
if not os.path.exists(FROZEN):
    with open(FROZEN, 'w', encoding='utf-8') as fh:
        json.dump({'about': 'Availability inputs of Tashkent August 2026 as the accepted golden '
                            'numbers were made — site data, never commit.',
                   'project_id': 1, 'year': 2026, 'month': 8,
                   'frozen_at': datetime.datetime.now().isoformat(timespec='seconds'),
                   'report_inputs': LIVE,
                   'history_records': B.load_history_records()}, fh, indent=1, ensure_ascii=False)
    print('frozen inputs written (first run):', FROZEN)
with open(FROZEN, encoding='utf-8') as fh:
    _frozen = json.load(fh)
INPUTS, HISTORY = _frozen['report_inputs'], _frozen['history_records']
print('frozen inputs:', FROZEN)
for k in ('exclusions', 'manual_unavailability', 'balancing_periods'):
    print('{:<22} {}'.format(k, len(INPUTS.get(k) or [])))
_diff = sorted(k for k in set(INPUTS) | set(LIVE) if INPUTS.get(k) != LIVE.get(k))
H.check(not _diff, 'the snapshot\'s report_inputs(1, 2026, 8) equal the frozen inputs{}'.format(
    '' if not _diff else ' — they differ in: ' + ', '.join(_diff)))
H.check(any(r.get('year') == 2026 and r.get('month_num') in (1, 2) for r in HISTORY),
        'frozen KPI history holds the January/February 2026 records cycles-in-year needs')

ARGS = dict(
    INPUTS,
    working_status_path=f('LC working status.xlsx'),
    pcs_cd_path=f('PCS_charge_discharge_status.xlsx'),
    soc_path=f('SOC august.xlsx'),
    soh_snapshot_path=f('Soh of last of the month.xlsx'),
    lc_charge_path=f('LC daily charge.xlsx'),
    lc_discharge_path=f('LC daily discharge.xlsx'),
    hv_meter_daily_path=f('HV meter daily export and import.xlsx'),
    alarm_path=f('Alarms report.XLSX'),
    cycles_first_day_path=f('Cycles first day of the month.xlsx'),
    cycles_last_day_path=f('Cycles last day of the month.xlsx'),
    lc_total_charge_path=f('LC total charge.xlsx'),
    lc_total_discharge_path=f('LC total discharge.xlsx'),
    pcs_fault_path=f('PCS fault status.xlsx'),
    site_name='ACWA BESS Tashkent',
    project_details={'Plant': 'Tashkent BESS'},
    progress_callback=None,          # set below: the log is checked too
    history_records=copy.deepcopy(HISTORY),
)
LOG = []
ARGS['progress_callback'] = LOG.append

# Word only: that is what the report is delivered as. The PDF renderer still
# lays the report out (its story is built either way, so a crash in it still
# shows up here) but writes no pages unless a PDF is asked for.
docx = os.path.join(H.WORK, 'August_2026.docx')
_t0 = time.time()
generate_tashkent_report(output_path=docx, output_format='docx', **ARGS)
T_XLSX = time.time() - _t0
H.check(os.path.exists(docx) and os.path.getsize(docx) > 50000,
        'DOCX written ({:,} bytes) — a DOCX failure is otherwise only logged'.format(
            os.path.getsize(docx) if os.path.exists(docx) else 0))
H.check(not os.path.exists(docx.rsplit('.', 1)[0] + '.pdf'),
        'no PDF left beside it')

from docx import Document                                           # noqa: E402

_doc = Document(docx)
txt = '\n'.join([p.text for p in _doc.paragraphs]
                + [c.text for t in _doc.tables for r in t.rows for c in r.cells])

print('\n=== general alarms ===')
H.check(txt.count(UMBRELLA_FOOTNOTE[:60]) >= 1, 'the general-alarm note is printed')
H.check('Input dry node fault' in txt, 'the held-back trigger is named in a note')
for gone in ('Dry Node Input Fault', 'LC System Alarm State', 'LCU Other Alarm'):
    H.check(gone not in txt, '"{}" names no table row'.format(gone))
H.check('Unclassified' not in txt, 'no blank-resolution "Unclassified" row')

print('\n=== availability ===')
m = re.search(r'availability for the period was ([\d.]+)%', txt)
avail = float(m.group(1)) if m else None
H.check(avail is not None and abs(avail - GOLDEN['availability_pct']) < 0.005,
        'BESS availability {}% (accepted {}%)'.format(avail, GOLDEN['availability_pct']))

print('\n=== fault names are printed whole ===')
# The tables used to slice names at 38-50 characters, so the customer read
# "PCS - Converter Unit 1 Fault Status 1: AC" and never learned it was an AC
# over-voltage. Both renderers wrap cell text, so nothing needs cutting.
import sqlite3                                                     # noqa: E402
_c = sqlite3.connect('file:{}?mode=ro'.format(H.DB.replace(os.sep, '/')), uri=True)
triggers = [t for (t,) in _c.execute(
    "SELECT DISTINCT trigger_name FROM alarm_events WHERE project_id=1 "
    "AND year=2026 AND month=8") if t and len(t) > 40]
_c.close()
cells = [c.text.strip() for t in Document(docx).tables for r in t.rows for c in r.cells]
# a cell that is the *start* of a known fault name, but not the whole name
cut = sorted({c for c in cells if len(c) >= 25
              and any(t.startswith(c) and t != c for t in triggers)})
H.check(not cut, 'no cell holds a half-written name ({} long names in the month){}'.format(
    len(triggers), '' if not cut else ': ' + repr(cut[:2])))
docx_text = '\n'.join(cells)
shown = [t for t in triggers if t in docx_text]
H.check(bool(shown), 'long names do appear in full, e.g. "{}"'.format(
    max(shown, key=len) if shown else '-'))

print('\n=== power-loss alarms from planned stops stay out of the fault tables ===')
# Islanding protection and SCU-DSP comm exceptions are what a PCS raises when its
# supply goes. In August 2026 386 h of them sat in the Faults/Warnings tables —
# raised a few minutes before each grid-outage window, and on every KKS day from
# the morning isolation (~09:30) while the windows had the afternoon work times.
pl_hours, pl_rows = 0.0, []
for t in Document(docx).tables:
    hdr = [c.text.strip() for c in t.rows[0].cells]
    if len(hdr) > 5 and hdr[1] in ('Fault Name', 'Warning') and hdr[5] == 'Total h':
        for row in t.rows[1:]:
            cells = [c.text.strip() for c in row.cells]
            if any(k in cells[1] for k in ('Islanding protection', 'SCU-DSP comm exception')):
                h = float(cells[5].replace(',', '') or 0)
                pl_hours += h
                pl_rows.append('{} {} h'.format(cells[1].replace('\n', ' ')[-40:], cells[5]))
for r in pl_rows:
    print('   ' + r)
H.check(pl_hours < 25, 'islanding / SCU-DSP left in the tables: {:.1f} h (was 386 h)'.format(pl_hours))

print('\n=== 4.4.1 unavailability reasons: no grid-outage edges on Aug 7-12 ===')
edge = []
for t in Document(docx).tables:
    hdr = [c.text.strip() for c in t.rows[0].cells]
    if 'Dominant Cause' in hdr:
        for row in t.rows[1:]:
            cells = [c.text.strip() for c in row.cells]
            print('   ' + ' | '.join(cells))
            if any(d in cells[1] for d in ('08-Aug', '09-Aug', '10-Aug', '11-Aug', '12-Aug')):
                edge.append(cells[0] + ' ' + cells[1])
H.check(not edge, 'no incident spans the corrected outage nights: {}'.format(edge or 'none'))

print('\n=== cycles ===')


def docx_value(path, label):
    for t in Document(path).tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) > 2 and cells[1] == label:
                return float(cells[2])
    return None


for key, label in (('cycles_month', 'Number of Cycles in Reported Month'),
                   ('cycles_in_year', 'Accumulative Number of Cycles in one Year'),
                   ('cycles_lifetime', 'Accumulative Number of Cycles')):
    v = docx_value(docx, label)
    H.check(v is not None and abs(v - GOLDEN[key]) < 0.05,
            '{}: {} (accepted {})'.format(label, v, GOLDEN[key]))

print('\n=== the month data set: the same exports, added once ===')
import database.db_manager as dbm                                  # noqa: E402
import services.month_dataset_service as mds                       # noqa: E402
import services.report_versions_service as rvs                     # noqa: E402
import services.report_precheck_service as rpc                     # noqa: E402
import services.parse_cache as pc                                  # noqa: E402

# The snapshot copies the live database: if August is already locked there, or
# holds files kept next to the live database, set them aside in the copy.
_c = dbm.get_connection()
_c.execute("UPDATE report_months SET locked_at=NULL WHERE project_id=1 AND year=2026 AND month=8")
_c.execute("UPDATE month_files SET status='removed' WHERE project_id=1 AND year=2026 AND month=8")
_c.commit(); _c.close()
TYPES13 = {'working_status', 'pcs_cd', 'pcs_fault', 'soc', 'alarms', 'lc_charge', 'lc_discharge',
           'hv_meter', 'cycles_first_day', 'cycles_last_day', 'soh_snapshot', 'lc_total_charge',
           'lc_total_discharge'}
_t0 = time.time()
added = mds.add_files(1, 2026, 8, mds.folder_files(SRC))
T_ADD = time.time() - _t0
H.check(len(added) == 13 and {r['type'] for r in added} == TYPES13
        and all(r['status'] == 'added' for r in added),
        '13 exports recognised and copied ({:.1f} s)'.format(T_ADD))
_t0 = time.time()
mds.analyse_pending(1, 2026, 8)
T_READ = time.time() - _t0
for g in mds.coverage_grid(1, 2026, 8):
    none = [d for d in mds.expected_days(2026, 8) if d not in g['days']]
    print('   {:<32} days {:>2}  part {:>2}  missing {:>2}  gaps>1h {}'.format(
        g['label'], len(g['days']), sum(v == 'partial' for v in g['days'].values()),
        len(none), len(g['gaps'])))
QS = rpc.run_precheck(1, 2026, 8)
for q in QS:
    print('   check: {:<5} {:<9} {} — {}'.format(q['severity'], q['kind'], q['title'], q['detail'][:90]))
H.check(not [q for q in QS if q['kind'] == 'missing'], 'the check finds every required file')

print('\n=== v1 from the data set: no file picked, no xlsx read ===')
SAFETY = [{'incident': 'TEST-MARKER near miss at block 12', 'equipment_loss': '-',
           'weight': 'minor', 'countermeasure': 'toolbox talk'}]
PARAMS = dict(INPUTS, site_name='ACWA BESS Tashkent', project_details={'Plant': 'Tashkent BESS'},
              safety_incidents=SAFETY)          # the re-run carries a safety incident (QA H-8)
_misses = pc.STATS['misses']
LOG2 = []
_t0 = time.time()
v1 = rvs.generate_version(1, 2026, 8, PARAMS,
                          confirmations=[{'kind': 'test', 'title': 'real-report test run'}],
                          progress=LOG2.append)
T_V1 = time.time() - _t0
H.check(pc.STATS['misses'] == _misses, 'every workbook came from the parse cache ({} hits)'.format(
    pc.STATS['hits']))
docx2 = v1['docx_abs']
H.check(docx_value(docx2, 'Accumulative Number of Cycles in one Year')
        == docx_value(docx, 'Accumulative Number of Cycles in one Year'),
        'Cycles in one Year stable on re-run (it used to grow by a month per run)')
_txt2 = '\n'.join([p.text for p in Document(docx2).paragraphs]
                  + [c.text for t in Document(docx2).tables for r in t.rows for c in r.cells])
m2 = re.search(r'availability for the period was ([\d.]+)%', _txt2)
H.check(m2 and abs(float(m2.group(1)) - GOLDEN['availability_pct']) < 0.005,
        'v1 availability {}% — the same as the export-folder run'.format(m2.group(1) if m2 else None))

print('\n=== safety incidents reach section 6 ===')
H.check('No safety incidents in the reporting period.' in txt,
        'nothing entered -> "No safety incidents in the reporting period."')
H.check('TEST-MARKER near miss at block 12' in _txt2
        and 'No safety incidents in the reporting period.' not in _txt2,
        'an entered incident is printed instead (was always "No safety incidents")')

print('\n=== the v1 manifest (QA G11) ===')
man = rvs.load_manifest(v1['manifest_abs'])
H.check(len(man['input_files']) == 13 and all(
    f['sha256'] == pc.file_sha256(os.path.join(SRC, f['original_name'])) for f in man['input_files']),
    '13 input files, each with the SHA-256 of the export it was copied from')
H.check(man['report_inputs'] == INPUTS, 'report_inputs used = the frozen inputs')
H.check(any(r.get('year') == 2026 and r.get('month_num') in (1, 2) for r in man['history_records']),
        'the KPI history used for cycles-in-year ({} records)'.format(len(man['history_records'])))
H.check(bool(man['alarm_classifications']['sha256']) and bool(man['app'].get('code_fingerprint')),
        'classification CSV hash {}…, report code {}'.format(man['alarm_classifications']['sha256'][:10],
                                                            man['app'].get('code_fingerprint')))
kn = man['key_numbers']
print('   key numbers:', {k: v for k, v in kn.items() if k != 'rows'}, 'rows:', kn['rows'])
H.check(abs(kn['availability_pct'] - GOLDEN['availability_pct']) < 0.005
        and abs(kn['cycles_in_year'] - GOLDEN['cycles_in_year']) < 0.05
        and abs(kn['cycles_month'] - GOLDEN['cycles_month']) < 0.05
        and abs(kn['cycles_lifetime'] - GOLDEN['cycles_lifetime']) < 0.05,
        'key numbers = the golden: {availability_pct:.2f} % · {cycles_month:.1f} / '
        '{cycles_in_year:.1f} / {cycles_lifetime:.1f}'.format(**kn))
H.check(all(kn.get(k) for k in ('rte_pct', 'avg_soc_pct', 'avg_soh_pct'))
        and all(isinstance(kn['rows'].get(k), int) for k in ('3.1', '3.2', '4.4.1', '5.2')),
        'RTE, SOC, SOH and the 3.1 / 3.2 / 4.4.1 / 5.2 row counts recorded')
H.check(man['confirmations'] and man['confirmations'][0]['title'] == 'real-report test run',
        'confirmations recorded')

print('\n=== v2: a new file; v1 untouched ===')
_t0 = time.time()
v2 = rvs.generate_version(1, 2026, 8, PARAMS)
T_V2 = time.time() - _t0
H.check(_digest(v1['docx_abs']) == v1['docx_sha256'] and v2['docx_abs'] != v1['docx_abs'],
        'v{} written beside v{}, whose SHA-256 is unchanged'.format(v2['version'], v1['version']))
H.check(rvs.list_versions(1, 2026, 8)[0]['changes'].startswith('No number changes'),
        'v2 against v1: "{}"'.format(rvs.list_versions(1, 2026, 8)[0]['changes']))

print('\n=== generation time ===')
print('   straight from the export folder (reads xlsx)   {:6.1f} s'.format(T_XLSX))
print('   add the 13 files to the month                  {:6.1f} s'.format(T_ADD))
print('   read them once for coverage (fills the cache)  {:6.1f} s'.format(T_READ))
print('   v1 from the data set (cache)                   {:6.1f} s'.format(T_V1))
print('   v2 from the data set (cache)                   {:6.1f} s'.format(T_V2))
H.check(T_V2 < T_XLSX, 'a re-generation from the data set is faster than reading the exports')

print('\n=== a PM record covers its own stop: nothing to do in August ===')
# August's PM was entered as manual rows inside Scheduled Maintenance windows;
# it has no PM records, so the rule must not touch it (golden above unchanged).
H.check(not any('PM stop' in str(m) or 'PM-covered' in str(m) or 'PM record #' in str(m) for m in LOG),
        'no PM stop attributed in the August log')

H.check(_digest(real_hist) == hist_before, 'the real KPI history file was not touched')

H.finish()

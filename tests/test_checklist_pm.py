"""PM checklists end to end on the desktop side: import the customer's
workbook, plan one checklist per block for a campaign, leave items out, add
one, fill it in the way a phone would, and export it back into the customer's
own file.
"""
import os

import _harness as H            # must be first
import openpyxl

import database.db_manager as dbm
import services.checklist_pm_service as cs

PM_DIR = os.environ.get('BESS_PM_DIR',
                        r'C:\Users\user1\Desktop\ACWA BESS Tashkent\PM')
PCS = os.path.join(PM_DIR, '01) PCS Checklist.xlsx')
if not os.path.isfile(PCS):
    H.skip('the PM checklist workbooks are not on this computer')

H.fresh_db('checklists.db')
c = dbm.get_connection()
c.execute("INSERT INTO projects (name, num_zones, num_blocks, num_containers, "
          "project_type) VALUES ('TK', 2, 16, 4, 'BESS')")
PID = c.execute("SELECT id FROM projects").fetchone()[0]
# The zone map comes from the containers, not from num_zones/num_blocks: a
# project with no containers has no map at all (plant_block_to_zone returns
# None, None), which is why the run keeps plant_block as the real number.
for z in (1, 2):
    for b in range(1, 9):
        c.execute("INSERT INTO containers (project_id, zone_number, block_number, "
                  "container_index, container_type, serial_number) "
                  "VALUES (?,?,?,1,'LC Cabinet',?)", (PID, z, b, 'S%d%d' % (z, b)))
c.commit(); c.close()
import services.project_service as psvc
psvc.get_block_map(PID, refresh=True)

# ── import ───────────────────────────────────────────────────────────────
tid = cs.import_template(PID, PCS)
its = cs.items(tid)
H.check(len(its) == 12 and its[0]['excel_row'] == 8,
        'the customer\'s workbook is imported as {} items'.format(len(its)))
tpl = cs.templates(PID)[0]
H.check(tpl['kind'] == 'PCS' and os.path.isfile(tpl['source_file']),
        'the workbook itself is kept beside the database ({})'.format(
            os.path.basename(tpl['source_file'])))
H.check(cs.import_template(PID, PCS) == tid and len(cs.items(tid)) == 12,
        'importing the same file again keeps the template and its 12 items')
H.check([i['id'] for i in cs.items(tid)] == [i['id'] for i in its],
        'and every item keeps its id — results of planned checklists point at it')

# ── plan a campaign: two blocks, one item out of scope ───────────────────
skip_item = its[3]['id']
uuids = cs.plan_runs(PID, tid, [12, 13], '2026-09-21', campaign='PM Sep 2026',
                     excluded_item_ids=[skip_item])
H.check(len(uuids) == 2, 'one checklist per block')
H.check(cs.plan_runs(PID, tid, [12], '2026-09-21', campaign='PM Sep 2026') == [uuids[0]],
        'planning the same block twice does not make a second checklist')
run = cs.run_detail(uuids[0])
H.check(run['plant_block'] == 12 and run['zone_number'] == 2 and run['block_number'] == 4,
        'the run knows its plant block and its zone/block')
H.check([i for i in run['items'] if i['item_id'] == skip_item][0]['result'] == cs.EXCLUDED,
        'the item left out of the campaign is marked Excluded, not deleted')

# an item this campaign adds
extra = cs.add_item(tid, 'Check the anti-condensation heater', equipment='PCS')
H.check(len(cs.items(tid)) == 13 and cs.items(tid)[-1]['added'] == 1,
        'the desktop can add an item; it goes at the end of its group')
try:
    cs.delete_item(its[0]['id'])
    H.check(False, "deleting the customer's own item must be refused")
except ValueError as e:
    H.check('exclude it instead' in str(e),
            "the customer's own item cannot be deleted: {}".format(str(e)[:40]))

# ── fill it the way the phone does ───────────────────────────────────────
run = cs.run_detail(uuids[0])
fill = {}
for i, it in enumerate(run['items']):
    if it['result'] == cs.EXCLUDED:
        fill[it['item_id']] = {'result': cs.EXCLUDED, 'comment': 'RMU not in this PM'}
    elif i == 1:
        fill[it['item_id']] = {'result': cs.NOK, 'comment': 'PCS 2: fan noisy, replaced'}
    elif i == 2:
        fill[it['item_id']] = {'result': cs.NA, 'comment': ''}
    else:
        fill[it['item_id']] = {'result': cs.OK, 'comment': ''}
p = cs.save_results(uuids[0], fill, status='Done', signed_by='Field team',
                    ptw_no='PTW-2609-140', serial='A2561724523', source='phone')
H.check(p['items'] == 13 and p['done'] == 12 and p['nok'] == 1 and p['excluded'] == 1,
        'progress counts what was ticked: {}'.format(p))
try:
    cs.save_results(uuids[0], {run['items'][0]['item_id']: {'result': 'Maybe'}})
    H.check(False, 'an unknown result must be refused')
except ValueError as e:
    H.check('Unknown result' in str(e), 'an unknown result is refused')

half = cs.run_detail(uuids[1])
cs.save_results(uuids[1], {half['items'][0]['item_id']: {'result': cs.OK}}, source='phone')
p2 = cs.progress(uuids[1])
H.check(p2['done'] == 1, 'a half-filled checklist leaves the rest pending ({})'.format(p2))
lst = cs.runs(PID, campaign='PM Sep 2026')
H.check(len(lst) == 2 and {r['n_done'] for r in lst} == {1, 12},
        'the campaign lists both blocks with how far each one is')

# ── export into the customer's own file ──────────────────────────────────
out = cs.export_run(uuids[0], H.WORK, plant_name='ACWA Riverside BESS')
H.check(os.path.basename(out).startswith('PCS Block 12'),
        'the file is named after the checklist and the block: {}'.format(
            os.path.basename(out)))
wb = openpyxl.load_workbook(out)
ws = wb.active
H.check(ws['C4'].value == 'ACWA Riverside BESS' and ws['C5'].value == 'Block 12'
        and ws['E4'].value == '21.09.2026' and ws['E5'].value == 'A2561724523',
        'the header is filled from the run')
H.check(ws['E8'].value is True and ws['E9'].value is False
        and 'fan noisy' in (ws['G9'].value or ''),
        'OK ticks, NOK unticks and carries the comment')
H.check(ws['E10'].value in (None, ''), 'N/A leaves the box empty')
H.check(ws['E11'].value in (None, '') and 'RMU not in this PM' in (ws['G11'].value or ''),
        'the excluded item keeps its row with our note')
H.check('heater' in (ws['D20'].value or '') and ws['E20'].value is True,
        'the added item is written at the end of its group')
H.check((ws['B21'].value or '').lower() == 'progress',
        'and the Progress row moved down with it')
wb.close()
H.check(cs.export_campaign(PID, 'PM Sep 2026', H.WORK, 'ACWA Riverside BESS')
        and len([f for f in os.listdir(H.WORK) if f.endswith('.xlsx')]) == 2,
        'a campaign exports one file per block')

# ── the same work through the Plan page ──────────────────────────────────
from PyQt5.QtCore import Qt                                        # noqa: E402
from PyQt5.QtWidgets import (QApplication, QDialog, QFileDialog,   # noqa: E402
                             QMessageBox)

app = QApplication.instance() or QApplication([])
import ui.planner_page as pp                                       # noqa: E402

QMessageBox.information = staticmethod(lambda *a, **k: QMessageBox.Ok)
QMessageBox.warning = staticmethod(lambda *a, **k: QMessageBox.Ok)
QMessageBox.question = staticmethod(lambda *a, **k: QMessageBox.Yes)

page = pp.PlannerPage()
page.set_current_project(PID)
# a tab of the Plan page, not a menu item of its own. Its index is pinned
# because the page loads a tab's contents by index (_on_tab); tabs are
# appended after it, never inserted before it.
_tabs = [page.tabs.tabText(i) for i in range(page.tabs.count())]
H.check(_tabs[3] == 'Checklists',
        'the Plan page has a Checklists tab (not a menu item of its own): {}'
        .format(_tabs))
page.tabs.setCurrentIndex(3)
H.check(page.cl_tbl.rowCount() == 2
        and page.cl_tbl.item(0, 4).text() in ('1/13', '12/13'),
        'it lists the planned checklists with how far each one is')


def plan_exec(self):
    self.kind.setCurrentIndex(1)                   # only the blocks I list
    self.blocks.setText('14')
    self.campaign.setText('PM UI test')
    self.tbl.item(3, 0).setCheckState(Qt.Unchecked)
    self.tbl.item(3, 4).setText('RMU not in this PM')
    return QDialog.Accepted


pp._ChecklistPlanDialog.exec_ = plan_exec
page._plan_checklists()
ui_run = [r for r in cs.runs(PID, campaign='PM UI test')]
H.check(len(ui_run) == 1 and ui_run[0]['plant_block'] == 14,
        'planning through the page makes one checklist for the block listed')
det = cs.run_detail(ui_run[0]['uuid'])
left_out = [i for i in det['items'] if i['result'] == cs.EXCLUDED]
H.check(len(left_out) == 1 and left_out[0]['comment'] == 'RMU not in this PM',
        "an item left out keeps the user's own note, not invented text")


def fill_exec(self):
    for r in range(self.tbl.rowCount()):
        w = self.tbl.cellWidget(r, 4)
        if w:
            w.setCurrentIndex(1)                   # OK
    self.ptw.setText('PTW-UI-1')
    return QDialog.Accepted


pp._ChecklistFillDialog.exec_ = fill_exec
page.cl_campaign.setCurrentIndex(page.cl_campaign.findData('PM UI test'))
page.cl_tbl.selectRow(0)
page._open_checklist()
done = cs.run_detail(ui_run[0]['uuid'])
H.check(done['status'] == 'Done' and done['ptw_no'] == 'PTW-UI-1',
        'filling it on the desktop marks it done and keeps the PTW')
H.check([i for i in done['items'] if i['result'] == cs.EXCLUDED][0]['comment']
        == 'RMU not in this PM',
        'and the excluded item is still out of scope, with its note')

out_dir = os.path.join(H.WORK, 'ui_export')
QFileDialog.getExistingDirectory = staticmethod(lambda *a, **k: out_dir)
page.cl_tbl.selectRow(0)
page._export_checklist()
H.check(os.path.isdir(out_dir) and any(f.endswith('.xlsx') for f in os.listdir(out_dir)),
        'and the page exports it into the customer\'s workbook')

# ── the customer re-issues the workbook with one row more ────────────────
# The dangerous case: every item below the new row moves down, so an item
# recognised by its Excel row alone would take over the results of the one
# above it.
import services.checklist_excel as cx                              # noqa: E402

tpl_name = cs.templates(PID)[0]['name']
before = {i['id']: i['description'] for i in cs.items(tid)}
nok_item = [i for i in cs.run_detail(uuids[0])['items'] if i['result'] == cs.NOK][0]
modified = cx.write_filled(
    PCS, os.path.join(H.WORK, 'PCS re-issued.xlsx'), {}, {},
    added=[{'after_row': its[0]['excel_row'], 'no': '1a', 'equipment': 'PCS',
            'text': 'Extra check the customer added', 'result': '', 'comment': ''}])
H.check(cs.import_template(PID, modified, name=tpl_name) == tid,
        'the re-issued workbook lands on the same template')
after = {i['id']: i['description'] for i in cs.items(tid)}
H.check(all(after.get(k) == v for k, v in before.items()),
        'every item that already existed still reads the same, under the same id')
H.check(len(after) == len(before) + 1
        and any(v == 'Extra check the customer added' for v in after.values()),
        'and the new row is added, not swapped in: {} items'.format(len(after)))
still = [i for i in cs.run_detail(uuids[0])['items'] if i['item_id'] == nok_item['item_id']][0]
H.check(still['result'] == cs.NOK and still['text'] == nok_item['text']
        and 'fan noisy' in still['comment'],
        'the NOK stays on the item it was written against')
H.check(cs.progress(uuids[0])['items'] == len(after),
        'and the new item counts in that checklist too')

# the other way round: the customer rewords an item. The row it sits on is
# what recognises it, so it keeps its id instead of arriving as a new item.
reworded = cs.items(tid)[5]
c = dbm.get_connection()
c.execute("UPDATE checklist_items SET description='Wording from the old issue' "
          "WHERE id=?", (reworded['id'],))
c.commit(); c.close()
cs.import_template(PID, modified, name=tpl_name)
again = {i['id']: i['description'] for i in cs.items(tid)}
H.check(len(again) == len(after) and again.get(reworded['id']) == reworded['description'],
        'a reworded item is recognised by its row, not added a second time')

H.finish()
